"""Lyra Slack bot — relay Slack <-> Toqan + rotinas de gestão das EPSs (v5.11).

v5.11 — DM aberta pelo e-mail do Slack: domínio iFood (DM_ALLOWED_DOMAINS) vê tudo, e-mail de EPS vê só a própria EPS
        (precisa do escopo users:read.email). DM com memória da conversa (últimas mensagens / thread), "ontem" completa
        o pedido anterior. Na DM a resposta sai na thread da pergunta, como no canal. A conversa recebe a data de hoje,
        e data digitada no futuro (29/11 pedido em 30/10) vira a data passada mais provável, avisando. Sem pergunta de volta: falta indicador = CSAT, falta período = D-1, falta fila = todas.
        "Resultado por operador" / "ofensores": quem tem visão completa (e-mail iFood ou canal interno) recebe direto
        no chat (até 15, do pior para o melhor; ofensores = Q4), com a planilha completa se passar disso. Nos canais
        das EPSs segue em planilha. Período pedido: ontem, mês ou semana; fila citada filtra.

v5.10 — PRINT_READING=0 desliga a leitura dos prints; falha técnica na leitura (download, Toqan, resposta fora do
        formato) vira alerta para o Felipe e não aparece para a EPS. Só print ilegível de verdade é devolvido.
        STATUS_SHADOW aceita lista de EPS em teste (ex.: AeC) e STATUS_TEST_CHANNELS guarda o canal de testes de cada
        EPS, para o modo teste seguir no canal de testes depois que EPS_CHANNELS virar o canal oficial.
        Acompanhamento de Fila da AeC (print hora a hora, só o CONSOLIDADO): cobra HC abaixo do planejado com volume
        no previsto, TMA > 110% do plan e SLA 1ª resposta abaixo da meta (SLA_METAS ou o P SLA do print). Print
        ilegível é pedido uma vez por thread; depois só alerta.

v5.9 — TMA_IN_REPORT=0 tira o TMA do report e da cobrança (CSAT segue normal).

v5.8 — EPS_EXTRA_CHANNELS: canais extras de conversa com recorte de uma EPS (sem report).

v5.7 — controles em print (Volume, Produtividade): a Lyra lê a imagem (upload no Toqan), confere pelos totais do
       próprio print e cobra (fila acumulando, gente sobrando/faltando, logado sem produção, equipe abaixo da célula);
       se não conseguir ler, diz que o print não está legível. Review: CSAT informado x promotoras/detratoras.

v5.6 — status de hora em hora em modo teste (STATUS_SHADOW=1): lê o canal real e responde no canal de testes da EPS;
       cobra quando o CSAT piora com o mesmo próximo passo; NRT sem dados não é cobrado.

v5.5 — sugestão de plano: Lyra consulta a Conhecimento (política/matriz), causa só a do dado, sem régua inventada;
       o app tira da resposta promessa de planilha/arquivo que ele não gerou.

v5.4 — fila com ação em andamento para o mesmo problema: cobra status e efeito, não ação nova; prazo vencido: resultado
       ou novo plano. Problema diferente da ação: pede plano para o problema novo.

v5.3 — mensagens montadas pelo app (thread do report, variação, esteira) saem com divisória entre as seções.

v5.2 — lê o status de hora em hora das EPSs (STATUS_CHANNELS) e cobra: CSAT/SLA fora da meta da fila, justificativa
       repetida, próximo passo genérico, previsão sem prazo. Causa do iFood fica só no registro (não cobra SLA).

v5.1 — dúvida de processo/política/agente: o app manda a Lyra chamar a Lyra Conhecimento (sem busca externa nem chute).

v5.0 — pergunta de causa: análise de variação primeiro; se a queda vem do efeito mix (mix ≥ 80% da queda e
       performance estável), responde só a leitura do mix, sem cobrar plano e sem esteira. Filas sem meta: "sem meta".

v4.7 — thread do report em seções D-1 x Mês + Leitura; conversa na thread leva o contexto da fila.

v4.6 — TMA em segundos; dispersão pela régua do Cockpit (CSAT Q4 ÷ Q1); esteira de impacto do mês (Lyra Qualidade
       v2.0) na thread das filas cobradas; owners por nome; Review e Canarinhos fora; chamados de moedas
       (Fluxo Owners CX -> #rewards-engine-support-cx).

v4.5 — TMA oficial no report (tempo de atribuição por interação × interações por ticket, vs plan da célula),
       decomposto em "mais interações" x "tempo por interação"; fila fora do alvo de CSAT OU TMA é cobrada.

v4.4 — report de analista: problema do D-1 x do mês, quadrantes mês x D-1 com dispersão, fila sem CSAT
       no D-1 fica fora; Lyra lê a thread ao vivo e comenta o plano (feedback de analista) além de registrar.

v4.3 — alarmística dos agentes: teste chega em cada subagente, reaviso enquanto continuar parado,
       teste na hora quando uma conversa volta com erro de acesso, aviso de agente esquecido na troca do token.

v4.2 — o APP monta o layout do report: título "Governança · <Fila> — dd/mm" no canal e detalhe
       + cobrança na thread. O modelo só devolve dados (JSON por fila); texto livre também é lido.

v4.1 — bloco por fila enxuto (3–4 linhas: status, D-1 vs meta, ofensor, cobrança 👉 com owner).

v4.0 — report em BLOCOS POR FILA no canal da EPS (um post por fila, cada um com a sua thread);
cobrança pelo D-1; owner marcado só quando o D-1 ficou abaixo da meta; compromissos e escalation por fila.

v3.9 — sem nome de operador no Slack; owners das filas marcados nas cobranças; escalation 5h
depois do report para quem não respondeu; Excel de operadores (CSAT + quadrante) sob pedido.

v3.8 — "roda o report [da aec|da csu]" na DM: dispara o report na hora (teste manual).

v3.7 — monitoramento dos agentes (Lyra + MONITORED_AGENTS: Nexus, Vera, ClubIA…) 4x ao dia, DM só quando
um agente para ou volta; lembrete do token do Databricks no 6º dia (validade de 7 dias a partir de
"troquei o token" na DM, ou inferida quando um agente volta depois de erro de token).

v3.6 — MOP conversacional: span por supervisor/coordenador, busca das pessoas citadas na
pergunta (só da EPS do canal) e envio do Excel do MOP quando pedem ("me manda o MOP").

v3.5 — duas planilhas: MOP oficial em MOP_SPREADSHEET_ID (aba MOP Consolidado);
Compromissos, Estado Lyra e MOP backup em LYRA_SPREADSHEET_ID.

v3.4 — MOP pelo fluxo "Reportar Oráculo": o Excel postado na thread do chamado de MOP
(MOP_FLOW_CHANNEL) é lido sozinho; sem cobrança de MOP nos canais; aviso só por DM
se o MOP não for atualizado até 18h do 1º dia útil / dia 15.

v3.3 — conversa: a Lyra fala como gestora; o app sempre informa a situação do MOP
(números, "ainda não enviado" ou "não deu para consultar") em vez de omitir.

v3.2 — segregação por EPS:
- EPS_CHANNELS mapeia cada EPS ao seu canal. Report, leitura de thread,
  compromissos, cobrança de MOP e conversa ficam isolados por EPS.
- Toda mensagem ao Toqan leva um bloco ESCOPO; em canal de EPS os dados são
  filtrados pelo e-mail do operador (aec / csu) e só as células da EPS aparecem.
- Upload de MOP no canal de uma EPS troca só as linhas daquela EPS.
- DM: DM_ALLOWED_USERS e e-mails iFood (visão completa) ou de EPS (só a EPS); canal não configurado não recebe resposta.

O que mudou na v3.1 (sustentação):
- Agendador único, resistente a restart: o estado de cada rotina fica na aba
  "Estado Lyra" da planilha, então um redeploy não perde leituras de thread
  nem duplica o report.
- Report dividido em mensagem principal + thread (delimitador ===THREAD===).
  Falha de consulta não é publicada como report; vira alerta.
- Ciclo de compromissos com Fila, Responsável e Thread; CSAT pré/pós por fila
  (lido das linhas "• *Fila:* Mês X% · D-1 Y%" do report).
- Respostas tardias atualizam o compromisso existente (ex.: prazo informado
  depois) em vez de duplicar; "sem resposta" é encerrado quando a liderança
  responde.
- Upload de MOP: lê valores (não fórmulas), escolhe a aba certa, valida
  colunas, faz backup antes de substituir e pode ser restrito a canais.
- Endpoints de debug protegidos por token.
- Feriados nacionais respeitados na cobrança do MOP.
"""
from __future__ import annotations

import base64
import io
import json
import logging
import os
import re
import threading
import time
import unicodedata
import urllib.error
import urllib.request
from datetime import date, datetime, time as datetime_time, timedelta, timezone
from typing import Any, Callable
from zoneinfo import ZoneInfo

from dotenv import load_dotenv
from flask import Flask, jsonify, request
from slack_bolt import App
from slack_bolt.adapter.flask import SlackRequestHandler

load_dotenv()

logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO").upper())
logger = logging.getLogger("lyra-slack-bot")

# ---------------------------------------------------------------------------
# Configuração
# ---------------------------------------------------------------------------
BRAZIL_TIMEZONE = ZoneInfo("America/Sao_Paulo")

TOQAN_BASE_URL = "https://api.toqan.ai/api"
TOQAN_TIMEOUT_SECONDS = 60
POLL_INTERVAL_SECONDS = 10
MAX_POLL_ATTEMPTS = 180  # 180 x 10s = 30 minutos
PROCESSING_REACTION = "eyes"
DONE_REACTION = "white_check_mark"
ERROR_MESSAGE = "Não consegui processar agora, tente novamente"
TIMEOUT_ERROR_MESSAGE = "Não consegui processar a tempo. Tente refazer a pergunta de forma mais específica."
EMPTY_ANSWER_MESSAGE = "Não consegui obter uma resposta agora."

DAILY_REPORT_CHANNEL = os.getenv("DAILY_REPORT_CHANNEL", "C0BF6JVFG7N")
ALERT_CHANNEL = os.getenv("ALERT_CHANNEL", "").strip()
DEBUG_TOKEN = os.getenv("DEBUG_TOKEN", "").strip()


def _parse_channel_map(raw: str) -> dict[str, str]:
    """'AeC:C111,CSU:C222' -> {'AeC': 'C111', 'CSU': 'C222'}"""
    result: dict[str, str] = {}
    for part in raw.split(","):
        if ":" in part:
            name, channel = part.split(":", 1)
            if name.strip() and channel.strip():
                result[name.strip()] = channel.strip()
    return result


def _parse_id_list(raw: str) -> set[str]:
    return {item.strip() for item in raw.split(",") if item.strip()}


# --- Segregação por EPS -------------------------------------------------------
# Cada EPS tem o próprio canal; nada de uma EPS aparece no canal da outra.
EPS_CHANNELS = _parse_channel_map(os.getenv("EPS_CHANNELS", ""))
# Canal interno iFood com report consolidado (todas as EPS). Vazio = desligado.
CONSOLIDATED_CHANNEL = os.getenv("CONSOLIDATED_CHANNEL", "").strip()
# Outros canais internos onde a Lyra responde com visão completa.
INTERNAL_CHANNELS = _parse_id_list(os.getenv("INTERNAL_CHANNELS", "")) | (
    {CONSOLIDATED_CHANNEL} if CONSOLIDATED_CHANNEL else set()
)
# Quem pode falar com a Lyra por DM (visão completa). Padrão: o ALERT_CHANNEL, se for um usuário.
DM_ALLOWED_USERS = _parse_id_list(os.getenv("DM_ALLOWED_USERS", "")) or (
    {t for t in _parse_id_list(ALERT_CHANNEL) if t.startswith("U")}
)
# DM liberada para e-mails desses domínios (visão completa). E-mail de EPS (aec/csu) vê só a própria EPS. Vazio = só a lista.
DM_ALLOWED_DOMAINS = {d.strip().lower().lstrip("@") for d in os.getenv("DM_ALLOWED_DOMAINS", "ifood.com.br").split(",")
                      if d.strip()}
CONSOLIDATED = "Consolidado"

# Células de cada EPS (fonte: Felipe, out/2026). O nome da EPS também é o trecho
# procurado no e-mail do operador ("aec" em @if-aec..., "csu" em ...csu...).
# Owners de cada fila (IDs do Slack), marcados nas cobranças do report.
# Pode ser sobrescrito por FILA_OWNERS="CX Suporte:U1|U2;Madrugada:U3" (aceita também "@Nome Sobrenome").
_DEFAULT_FILA_OWNERS: dict[str, list[str]] = {
    "CX Suporte": ["U09GNBTALNN", "U094HUNCPFS"],
    "Madrugada": ["U094HUNCPFS"],
    "CX Review - AeC": ["U076W2160TF"],
    "CX Review - CSU": ["U05EVUPKF0R"],
    "CX Super Cliente - AeC": ["U09GNBTALNN"],
    "CX Super Cliente - CSU": ["U05EVUPKF0R"],
    # por nome: o app procura o usuário no Slack (escopo users:read) e marca o ID certo
    "PF Cancelamento": ["@Adriana Ventura da Silva"],
    "PF Super Clientes": ["@Adriana Ventura da Silva"],
    "PF Revisão Manual": ["@Adriana Ventura da Silva"],
    "CX Denúncia": ["U02CL56C46N"],
    "Fraudes DX": ["U02CL56C46N"],
}


def _parse_owners(raw: str) -> dict[str, list[str]]:
    owners: dict[str, list[str]] = {}
    for part in raw.split(";"):
        if ":" in part:
            fila, ids = part.split(":", 1)
            ids_list = [i.strip() for i in ids.split("|") if i.strip()]
            if fila.strip() and ids_list:
                owners[fila.strip()] = ids_list
    return owners


FILA_OWNERS: dict[str, list[str]] = {**_DEFAULT_FILA_OWNERS, **_parse_owners(os.getenv("FILA_OWNERS", ""))}  # variável complementa/substitui por fila

# TMA planejado por célula (aba PLAN do Excel mensal): tempo único (s) e média de interações.
# TMA = tempo de atribuição por interação × média de interações por ticket (coluna X da CAPA).
# Filas que não têm TMA (não aparece no report nem é cobrado)
NO_TMA_FILAS = {"DX - PF Contestação", "DX - PF Contestação Bloqueio temporário"}
# TMA_IN_REPORT=0 tira o TMA do report e da cobrança (ex.: enquanto a query não bate com a planilha de eficiência)
TMA_IN_REPORT = os.getenv("TMA_IN_REPORT", "1").strip().lower() not in {"0", "false", "nao", "não", "off"}
_DEFAULT_TMA_PLAN: dict[str, tuple[float, float]] = {
    "CX Suporte": (520, 1.8), "Madrugada": (562, 1.89),
    "CX Review - AeC": (620, 2.05), "CX Review - CSU": (620, 2.05),
    "CX Super Cliente - AeC": (725, 1.73), "CX Super Cliente - CSU": (725, 1.73),
    "PF Cancelamento": (450, 1.23), "PF Super Clientes": (400, 1.23), "PF Revisão Manual": (410, 1.27),
    "CX Denúncia": (560, 1.0), "Fraudes DX": (585, 1.0),
}


def _parse_tma_plan(raw: str) -> dict[str, tuple[float, float]]:
    """'CX Suporte:520/1.8;Madrugada:562/1.89' -> {fila: (tempo_unico_s, interacoes)}"""
    plan: dict[str, tuple[float, float]] = {}
    for part in raw.split(";"):
        if ":" not in part or "/" not in part:
            continue
        fila, values = part.rsplit(":", 1)
        tempo, inter = values.split("/", 1)
        try:
            plan[fila.strip()] = (float(tempo.replace(",", ".")), float(inter.replace(",", ".")))
        except ValueError:
            continue
    return plan


TMA_PLAN: dict[str, tuple[float, float]] = {**_DEFAULT_TMA_PLAN, **_parse_tma_plan(os.getenv("TMA_PLAN", ""))}
ESCALATION_DELAY = timedelta(hours=5)

EPS_CELULAS: dict[str, list[str]] = {
    "AeC": [
        "CX Suporte", "Madrugada", "CX Review - AeC", "CX Super Cliente - AeC",
    ],
    "CSU": [
        "CX Review - CSU", "CX Super Cliente - CSU", "PF Cancelamento",
        "PF Super Clientes", "PF Revisão Manual", "CX Denúncia", "Fraudes DX",
        "DX - PF Contestação", "DX - PF Contestação Bloqueio temporário",
    ],
}


class Tenant:
    """Um destino de report: uma EPS (eps='AeC') ou a visão consolidada (eps=None)."""

    def __init__(self, key: str, channel: str, eps: str | None) -> None:
        self.key, self.channel, self.eps = key, channel, eps

    def __repr__(self) -> str:
        return f"Tenant({self.key}, {self.channel})"


def _build_tenants() -> list[Tenant]:
    tenants = [Tenant(eps, channel, eps) for eps, channel in EPS_CHANNELS.items()]
    if CONSOLIDATED_CHANNEL:
        tenants.append(Tenant(CONSOLIDATED, CONSOLIDATED_CHANNEL, None))
    if not tenants:  # legado: um canal só, sem recorte de EPS
        tenants.append(Tenant(CONSOLIDATED, DAILY_REPORT_CHANNEL, None))
    return tenants


TENANTS = _build_tenants()
# canal -> EPS do canal (None = visão completa)
CHANNEL_SCOPE: dict[str, str | None] = {t.channel: t.eps for t in TENANTS}
CHANNEL_SCOPE.update({channel: None for channel in INTERNAL_CHANNELS})
# Canais extras de conversa com recorte de uma EPS (ex.: teste com alguém da EPS): a Lyra responde quando marcada,
# só com dados daquela EPS; report e rotinas continuam no canal principal (EPS_CHANNELS).
# Formato: "AeC:C111,AeC:C222,CSU:C333" (pode repetir a EPS).
EPS_EXTRA_CHANNELS: list[tuple[str, str]] = [
    (part.split(":", 1)[0].strip(), part.split(":", 1)[1].strip())
    for part in os.getenv("EPS_EXTRA_CHANNELS", "").split(",") if ":" in part and part.split(":", 1)[1].strip()
]
for _eps, _channel in EPS_EXTRA_CHANNELS:
    _known = next((key for key in EPS_CELULAS if key.casefold() == _eps.casefold()), _eps)
    CHANNEL_SCOPE.setdefault(_channel, _known)

DM_NOT_ALLOWED_MESSAGE = "Por aqui eu não consigo responder. Me chama no canal da sua EPS que eu te ajudo por lá. 🙂"
CHANNEL_NOT_CONFIGURED_MESSAGE = "Ainda não estou configurada para responder neste canal."

# --- MOP ---------------------------------------------------------------------
# Planilha da Lyra: Compromissos, Estado Lyra e MOP backup
LYRA_SPREADSHEET_ID = os.getenv("LYRA_SPREADSHEET_ID", "1Up11UQ1j0h9t8KW276AxJPGfEsShs9AsUztF5PAmyyE")
# Planilha do MOP oficial (aba "MOP Consolidado")
MOP_SPREADSHEET_ID = os.getenv("MOP_SPREADSHEET_ID", "1cy23m4iN0D7vEiJbbAw9hRpUtaHgGuRew_fgAUz_l14")
MOP_WORKSHEET_GID = int(os.getenv("MOP_WORKSHEET_GID", "464304384"))
MOP_SHEET_URL = (
    f"https://docs.google.com/spreadsheets/d/{MOP_SPREADSHEET_ID}/edit?gid={MOP_WORKSHEET_GID}#gid={MOP_WORKSHEET_GID}"
)
MOP_BACKUP_WORKSHEET_TITLE = "MOP backup"
MOP_ALLOWED_CHANNELS = {c.strip() for c in os.getenv("MOP_ALLOWED_CHANNELS", "").split(",") if c.strip()}
MOP_REQUIRED_COLUMNS = ("Nome", "Email")
MOP_EXPECTED_COLUMNS = (
    "Nome", "Email", "Pilar", "Célula", "EPS", "Supervisor", "Coordenador",
    "Entrada", "Saída", "Status", "Versão do Oráculo", "Formato de Trabalho",
)
MOP_HEADER_SCAN_ROWS = 10

MOP_CREDENTIALS_ERROR_MESSAGE = "Não consigo atualizar o MOP agora — credenciais do Google não configuradas."
MOP_DOWNLOAD_ERROR_MESSAGE = "Não consegui baixar o arquivo. Tente enviar novamente."
MOP_READ_ERROR_MESSAGE = "Não consegui ler o arquivo Excel. Verifique se o formato está correto."
MOP_SHEETS_ERROR_MESSAGE = "Erro ao atualizar a planilha do MOP. A versão anterior foi mantida. Tente novamente."
MOP_CHANNEL_NOT_ALLOWED_MESSAGE = (
    "Esse canal não está liberado para atualizar o MOP. Envie o arquivo no canal oficial do MOP marcando a Lyra."
)

MOP_INTENT_KEYWORDS = (
    "mop", "hierarquia", "analista", "anlista", "equipe", "pessoas", "efetivo",
    "headcount", "supervisor", "coordenador", "célula", "celula", "pilar",
    "oráculo", "oraculo", "variante", "eps", "span", "spam", "gestor", "gestão", "gestao",
    "time do", "time da", "turno", "escala", "férias", "ferias", "licença", "licenca",
    "entrada", "saída", "saida", "formato de trabalho", "remoto", "presencial", "quem é", "quem e",
)

# Canal do fluxo "Reportar Oráculo": o Excel do MOP postado na thread de um chamado
# "MOP - Atualização Mensal & Quinzenal" é lido automaticamente. Vazio = desligado.
MOP_FLOW_CHANNEL = os.getenv("MOP_FLOW_CHANNEL", "").strip()

# --- Compromissos ------------------------------------------------------------
COMPROMISSOS_WORKSHEET_TITLE = "Compromissos"
# As 8 primeiras colunas são as originais (linhas antigas continuam válidas);
# Fila, Responsável e Thread foram adicionadas ao final.
COMPROMISSOS_HEADERS = [
    "Problema Identificado", "Ação Tomada", "Data Inicio", "Data Fim",
    "Duração", "Resultado Pré", "Resultado Pós", "Sucesso",
    "Fila", "Responsável", "Thread", "EPS",
]
COL_ACAO = COMPROMISSOS_HEADERS.index("Ação Tomada") + 1
COL_DATA_FIM = COMPROMISSOS_HEADERS.index("Data Fim") + 1
COL_DURACAO = COMPROMISSOS_HEADERS.index("Duração") + 1
COL_POS = COMPROMISSOS_HEADERS.index("Resultado Pós") + 1
COL_SUCESSO = COMPROMISSOS_HEADERS.index("Sucesso") + 1
COL_RESPONSAVEL = COMPROMISSOS_HEADERS.index("Responsável") + 1
SEM_RESPOSTA = "sem resposta"
FILA_GERAL = "Geral"
OPEN_COMMITMENTS_LIMIT = 15
SEM_RESPOSTA_WINDOW_DAYS = 3
EXPIRED_AFTER_DAYS = 7

# --- Estado persistido --------------------------------------------------------
STATE_WORKSHEET_TITLE = "Estado Lyra"
STATE_HEADERS = ["Chave", "Valor", "Atualizado em"]

# --- Agenda (horário de Brasília) ------------------------------------------------
DAILY_REPORT_TIME = datetime_time(9, 0)
REPORT_CATCHUP_UNTIL = datetime_time(13, 0)   # restart até 13h ainda publica o report do dia
REPORT_MAX_ATTEMPTS = 3
REPORT_RETRY_MINUTES = 10
PREVIOUS_THREAD_READ_TIME = datetime_time(8, 45)  # relê a thread de ontem antes do report
THREAD_READ_FIRST_TIME = datetime_time(11, 0)
THREAD_READ_FIRST_MIN_DELAY = timedelta(hours=2)
THREAD_READ_LATE_TIME = datetime_time(18, 0)
THREAD_READ_LATE_MIN_DELAY = timedelta(hours=4)
# Aviso silencioso (só para o ALERT_CHANNEL) se o MOP não foi atualizado no dia previsto
MOP_MISSING_ALERT_TIME = datetime_time(18, 0)
MOP_UPDATE_GRACE_DAYS = 2  # atualização feita até 2 dias antes também conta
SCHEDULER_TICK_SECONDS = 60

# --- Saúde dos agentes e do token do Databricks ---------------------------------------
# Agentes monitorados além da Lyra: "Nexus:<api_key>,Vera:<api_key>,ClubIA:<api_key>"
MONITORED_AGENTS = _parse_channel_map(os.getenv("MONITORED_AGENTS", ""))
AGENT_CHECK_TIMES = [
    datetime_time(int(h), 30) for h in os.getenv("AGENT_CHECK_HOURS", "7,11,15,19").split(",") if h.strip().isdigit()
]
AGENT_CHECK_MAX_POLLS = 30  # 30 x 10s = 5 min por agente
# O token do Databricks dura DATABRICKS_TOKEN_VALIDITY_DAYS a partir da troca.
DATABRICKS_TOKEN_VALIDITY_DAYS = int(os.getenv("DATABRICKS_TOKEN_VALIDITY_DAYS", "7"))
# Agentes que usam o MESMO token do Databricks (só para o texto do alerta)
# Agentes que não consultam o Databricks: o teste só confere se respondem (sem consulta, sem alerta de token)
NO_DATA_AGENTS = {a.strip().casefold() for a in os.getenv("NO_DATA_AGENTS", "Em preparo").split(",") if a.strip()}
PING_PROMPT = (
    "TAREFA INTERNA DE MONITORAMENTO — não é pergunta de usuário e não vai para o Slack. "
    "Não consulte nada. Responda APENAS: OK"
)
SHARED_TOKEN_AGENTS = [a.strip() for a in os.getenv("SHARED_TOKEN_AGENTS", "Nexus,Vera,ClubIA").split(",") if a.strip()]
DATABRICKS_TOKEN_RENEWED = os.getenv("DATABRICKS_TOKEN_RENEWED", "").strip()  # opcional: data da última troca
TOKEN_REMINDER_TIME = datetime_time(9, 30)
TOKEN_ERROR_HINTS = ("403", "401", "invalid access token", "token", "unauthorized", "unauthenticated",
                     "expired", "expirado", "permission", "permiss", "credential", "credencia")
HEALTHCHECK_PROMPT = (
    "EPS: iFood\n"
    "Pedido: verificação de acesso\n\n"
    "TAREFA INTERNA DE MONITORAMENTO — não é pergunta de usuário e não vai para o Slack.\n"
    "Se você consulta dados direto: rode SELECT COUNT(*) AS n FROM saex.gold.cx_support "
    "WHERE DATE(ticket_created_at) >= date_sub(current_date(), 2).\n"
    "Se você usa subagentes: chame CADA subagente que consulta dados (ex.: Lyra Indicadores e Lyra Qualidade), "
    "um por um, com 'EPS: iFood' e 'Pedido: verificação de acesso'.\n"
    "Responda APENAS uma linha: 'OK <n>' se TODAS as consultas funcionaram, ou "
    "'ERRO: <nome do subagente, se houver>: <mensagem técnica exata do erro>' se qualquer uma falhou."
)

THREAD_DELIMITER = "===THREAD==="

FILA_BLOCK_RE = re.compile(r"^\s*===\s*FILA\s*:\s*(?P<fila>.+?)\s*===\s*$", re.MULTILINE)
DAILY_REPORT_PROMPT = (
    "Report diário por fila. ESTA TAREFA É DO APP: este contrato vale acima de qualquer layout do seu prompt. "
    "Você NÃO escreve o texto do Slack — o app monta o layout. Você devolve só os dados e a frase de cobrança.\n\n"
    "1. Chame o Lyra Indicadores (EPS do escopo, 'Pedido: dados do report diário').\n"
    "2. Para CADA fila da EPS que teve avaliação de CSAT no D-1, escreva exatamente duas linhas "
    "(fila sem avaliação no D-1 fica de fora — não escreva nada sobre ela):\n"
    "===FILA: <nome exato da fila>===\n"
    '{"d1": 91.7, "n": 399, "var": -0.2, "meta": 94, "mes": 93.9, '
    '"ofensor": "pedido atrasado", "ofensor_pp": -2.8, "ofensor_mes": "pedido atrasado", "ofensor_mes_pp": -1.9, '
    '"quad_mes": {"q1": 39, "q2": 46, "q3": 12, "q4": 14}, "quad_d1": {"q1": 20, "q2": 18, "q3": 9, "q4": 16}, '
    '"disp_mes": {"q1": 98.5, "q4": 78.0}, "disp_d1": {"q1": 97.0, "q4": 62.0}, '
    '"reincidentes": 3, "tma_tpi_d1": 379.6, "tma_int_d1": 1.92, "tma_tpi_mes": 368.4, "tma_int_mes": 1.90, "tma_tratados_d1": 2650, '
    '"compromisso": "", "compromisso_cobre": null, '
    '"cobranca": "Pedido atrasado tira 2,8 p.p. no D-1 e já é o problema do mês, e a dispersão aumentou — o que muda hoje na tratativa de atraso?", '
    '"elogio": ""}\n'
    "3. Campos: d1 = CSAT D-1 em %; n = avaliações no D-1; var = variação do D-1 vs dia anterior em p.p. (negativo se caiu); "
    "meta em % (null se não definida); mes = CSAT do mês em %; ofensor / ofensor_pp = motivo de maior impacto negativo no D-1 "
    "e o impacto em p.p. (negativo); ofensor_mes / ofensor_mes_pp = o mesmo no mês; "
    "quad_mes / quad_d1 = quantidade de operadores em cada quadrante no mês e no D-1; "
    "disp_mes / disp_d1 = CSAT (%) dos operadores do Q1 e do Q4 no mês e no D-1 (dispersão = Q4 ÷ Q1, régua do Cockpit); "
    "reincidentes = operadores reincidentes no Q4; "
    "tma_tpi_d1 / tma_tpi_mes = tempo de atribuição por interação em SEGUNDOS; tma_int_d1 / tma_int_mes = média de interações por ticket; "
    "tma_tratados_d1 = interações tratadas no D-1 (copie do bloco TMA POR FILA do Indicadores; o app calcula o TMA e compara com o plan); "
    "compromisso = compromisso em aberto DESTA fila no contexto ('<ação> · prazo <dd/mm>') ou vazio; "
    "compromisso_cobre = true se essa ação ataca o problema de hoje (mesmo indicador/causa), false se o problema de hoje é outro, null se não há compromisso.\n"
    "4. cobranca (se CSAT do D-1 abaixo da meta OU TMA do D-1 acima do plan — o plan vem na lista PLAN DE TMA desta mensagem): "
    "UMA frase de analista de CX, assertiva, que conecta os dados — problema do D-1 vs do mês (estrutural ou pontual), "
    "a dispersão (índice Q4÷Q1 com duas casas, ex.: 0,86 → 0,76, nunca em %; caindo = dispersão aumentou = comportamento operacional) e o TMA (se subiu por mais interações por ticket = não resolveu "
    "no primeiro contato; se por tempo por interação = atendimento mais longo) — e termina numa pergunta direta sobre a ação. "
    "NÃO repita o valor do TMA nem do CSAT na cobrança (o app já mostra os números exatos nas seções acima; diga só "
    "'acima do plan' / 'abaixo da meta'). Interações por ticket são interações, não reaberturas. "
    "Se há compromisso em aberto que ataca o problema (compromisso_cobre true), NÃO peça ação nova: a pergunta é sobre o status "
    "e o efeito da ação que já está rodando (ex.: 'a ação X está segurando o problema?'). Se o problema de hoje é outro, peça a ação para ele. "
    "Sem 'Plano aqui na thread' (o app completa). elogio (CSAT e TMA no alvo): reconhecimento curto e específico.\n"
    "5. Só isso: nenhum título, texto, bullet, separador (---) ou comentário fora dessas linhas. Números com ponto decimal. "
    "Nunca nome, usuário ou e-mail de operador. Não escreva menções (@).\n"
    "6. Se qualquer consulta falhar ou voltar vazia: responda SOMENTE uma linha começando com ERRO: e o motivo."
)
REPORT_FAILED_CHANNEL_MESSAGE = (
    "Pessoal, o report de CSAT de hoje não saiu — tive um problema pra consultar os dados. "
    "Já sinalizei o time responsável e publico assim que normalizar."
)

try:  # feriados nacionais (lib opcional)
    import holidays as _holidays_lib

    _BR_HOLIDAYS = _holidays_lib.country_holidays("BR")
except Exception:  # pragma: no cover - lib ausente
    _BR_HOLIDAYS = None

# Slack Bolt sobe sem credenciais para smoke tests locais; em produção os
# valores reais vêm das variáveis de ambiente.
slack_app = App(
    token=os.getenv("SLACK_BOT_TOKEN") or "dev-token",
    signing_secret=os.getenv("SLACK_SIGNING_SECRET") or "dev-signing-secret",
    token_verification_enabled=bool(os.getenv("SLACK_BOT_TOKEN")),
)
slack_handler = SlackRequestHandler(slack_app)
flask_app = Flask(__name__)
app = flask_app
PORT = int(os.getenv("PORT", "3000"))

_MENTION_RE = re.compile(r"<@[A-Z0-9]+(?:\|[^>]+)?>")
_THINK_RE = re.compile(r"<think>.*?</think>", flags=re.DOTALL | re.IGNORECASE)
_IN_FLIGHT: set[str] = set()
_IN_FLIGHT_LOCK = threading.Lock()


# ---------------------------------------------------------------------------
# Helpers gerais
# ---------------------------------------------------------------------------
def _resp_get(response: Any, key: str, default: Any = None) -> Any:
    """Lê um campo de SlackResponse ou dict. SlackResponse NÃO é dict."""
    if response is None:
        return default
    getter = getattr(response, "get", None)
    if callable(getter):
        try:
            return getter(key, default)
        except Exception:
            pass
    try:
        return response[key]
    except Exception:
        return default


def _now_brazil() -> datetime:
    return datetime.now(BRAZIL_TIMEZONE)


def _today_brazil() -> datetime:
    return _now_brazil()


def _as_date(value: datetime | date) -> date:
    return value.date() if isinstance(value, datetime) else value


def _format_date(value: datetime | date) -> str:
    return value.strftime("%d/%m/%Y")


def _norm_text(value: Any) -> str:
    """Minúsculas, sem acento, hífens/espaços colapsados — para comparar nomes."""
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[\s\-–—_:]+", " ", text.casefold())
    return text.strip()


def _to_float(value: Any) -> float | None:
    match = re.search(r"\d+(?:[.,]\d+)?", str(value or ""))
    return float(match.group(0).replace(",", ".")) if match else None


def _fmt_pct(value: float) -> str:
    return f"{value:.1f}".replace(".", ",") + "%"


def _to_int(value: Any) -> int:
    try:
        return int(str(value).strip())
    except Exception:
        return 0


def _col_letter(n: int) -> str:
    letters = ""
    while n > 0:
        n, rem = divmod(n - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _row_value(row: dict[str, Any], *names: str) -> str:
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _alert(text: str) -> None:
    """Avisa o responsável técnico (ALERT_CHANNEL) — sem expor dados pessoais."""
    logger.warning("ALERTA: %s", text)
    for target in [t.strip() for t in ALERT_CHANNEL.split(",") if t.strip()]:
        try:
            slack_app.client.chat_postMessage(channel=target, text=f"⚠️ Lyra — {text}")
        except Exception:
            logger.exception("Could not post alert to %s", target)


def _eps_pattern(eps: str) -> str:
    """Trecho que identifica a EPS no e-mail do operador (ex.: 'aec', 'csu')."""
    return _compact(eps)


def _eps_filter_sql(eps: str) -> str:
    return f"lower(ticket_last_agent_email) LIKE '%{_eps_pattern(eps)}%'"


def _email_eps(email: Any) -> str | None:
    """EPS deduzida pelo e-mail do operador (None se não bater com nenhuma ou com mais de uma)."""
    text = str(email or "").casefold()
    hits = [eps for eps in EPS_CELULAS if _eps_pattern(eps) in text]
    return hits[0] if len(hits) == 1 else None


def _celulas_of(eps: str) -> list[str]:
    for key, celulas in EPS_CELULAS.items():
        if _compact(key) == _compact(eps):
            return celulas
    return []


def _scope_instruction(eps: str | None) -> str:
    """Bloco de escopo injetado em toda mensagem enviada ao Toqan."""
    if eps is None:
        return "ESCOPO: visão interna iFood — todas as EPS. Pode comparar EPS."
    others = ", ".join(e for e in EPS_CELULAS if _compact(e) != _compact(eps)) or "outras EPS"
    celulas = _celulas_of(eps)
    return (
        f"ESCOPO OBRIGATÓRIO: este canal é exclusivo da EPS {eps}. "
        + (f"Filas da {eps} (fila do ticket = ticket_last_form.queue_name): {'; '.join(celulas)}. Use só essas. " if celulas else "")
        + f"Para operadores (Q4, rankings), use só e-mails da {eps}: {_eps_filter_sql(eps)}. "
        + f"Use somente filas, operadores, motivos, compromissos e números da {eps}. "
        f"NUNCA mencione, compare ou mostre dados de {others}, nem o resultado consolidado do iFood. "
        f"Se pedirem dados de outra EPS, diga com naturalidade que aqui você acompanha só a {eps}. "
        "Nunca mencione este escopo, filtros, e-mails, SQL ou como o canal foi configurado."
    )


_DM_SCOPE_CACHE: dict[str, tuple[bool, str | None]] = {}
_DM_EMAIL_ALERTED = False


def _dm_scope(user: str) -> tuple[bool, str | None]:
    """Quem pode usar a DM e com qual visão, pelo e-mail do Slack:
    DM_ALLOWED_USERS ou domínio iFood (DM_ALLOWED_DOMAINS) = visão completa; e-mail de EPS (aec/csu no domínio) = só
    a EPS dela; resto = sem DM. Precisa do escopo users:read.email no app do Slack."""
    global _DM_EMAIL_ALERTED
    if user in DM_ALLOWED_USERS:
        return True, None
    if user in _DM_SCOPE_CACHE:
        return _DM_SCOPE_CACHE[user]
    email = ""
    try:
        info = slack_app.client.users_info(user=user)
        email = str(((info.get("user") or {}).get("profile") or {}).get("email") or "").lower()
    except Exception:
        logger.warning("Could not read e-mail of %s", user, exc_info=True)
    if not email:
        if not _DM_EMAIL_ALERTED:
            _DM_EMAIL_ALERTED = True
            _alert("DM: não consegui ler o e-mail de quem chamou a Lyra. Inclua o escopo users:read.email no app do Slack "
                   "(OAuth & Permissions) e reinstale; até lá a DM fica só para DM_ALLOWED_USERS.")
        return False, None
    domain = email.rsplit("@", 1)[-1]
    result: tuple[bool, str | None] = (False, None)
    eps_hit = next((e for e in EPS_CELULAS if _eps_pattern(e) and _eps_pattern(e) in domain), None)
    if eps_hit:
        result = (True, eps_hit)
    elif any(domain == d or domain.endswith("." + d) for d in DM_ALLOWED_DOMAINS):
        result = (True, None)
    _DM_SCOPE_CACHE[user] = result
    return result


def _event_scope(event: dict[str, Any]) -> tuple[bool, str | None]:
    """(permitido?, EPS do escopo). EPS None = visão completa."""
    if event.get("channel_type") == "im":
        return _dm_scope(str(event.get("user") or ""))
    channel = str(event.get("channel") or "")
    if channel in CHANNEL_SCOPE:
        return True, CHANNEL_SCOPE[channel]
    return False, None


# ---------------------------------------------------------------------------
# Toqan
# ---------------------------------------------------------------------------
class ToqanAPIError(RuntimeError):
    """Raised when the Toqan Agent API cannot complete a request."""


def _require_toqan_key() -> str:
    api_key = os.getenv("TOQAN_API_KEY")
    if not api_key:
        raise ToqanAPIError("TOQAN_API_KEY is not configured")
    return api_key


def _toqan_request(method: str, payload: dict[str, Any], api_key: str | None = None,
                   http_method: str | None = None) -> dict[str, Any]:
    """Make one authenticated JSON request to the Toqan Agent API (chave da Lyra por padrão)."""
    body = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        url=f"{TOQAN_BASE_URL}/{method}",
        data=body,
        headers={"X-Api-Key": api_key or _require_toqan_key(), "Content-Type": "application/json"},
        method=http_method or ("POST" if method in {"create_conversation", "continue_conversation"} else "GET"),
    )
    try:
        with urllib.request.urlopen(req, timeout=TOQAN_TIMEOUT_SECONDS) as response:
            raw = response.read().decode("utf-8")
    except urllib.error.HTTPError as exc:
        logger.error("Toqan API returned HTTP %s for %s", exc.code, method)
        raise ToqanAPIError(f"Toqan HTTP {exc.code}") from exc
    except (urllib.error.URLError, TimeoutError, OSError) as exc:
        logger.error("Toqan API request failed for %s: %s", method, exc)
        raise ToqanAPIError("Toqan request failed") from exc
    try:
        data = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ToqanAPIError("Toqan returned invalid JSON") from exc
    if not isinstance(data, dict):
        raise ToqanAPIError("Toqan returned an unexpected response")
    return data


def create_conversation(user_message: str, api_key: str | None = None) -> tuple[str, str]:
    data = _toqan_request("create_conversation", {"user_message": user_message}, api_key)
    conversation_id = data.get("conversation_id")
    request_id = data.get("request_id")
    if not isinstance(conversation_id, str) or not conversation_id:
        raise ToqanAPIError("Toqan response did not include conversation_id")
    if not isinstance(request_id, str) or not request_id:
        raise ToqanAPIError("Toqan response did not include request_id")
    return conversation_id, request_id


def get_answer(
    conversation_id: str,
    request_id: str,
    progress_callback: Callable[[str], Any] | None = None,
    api_key: str | None = None,
    max_attempts: int = MAX_POLL_ATTEMPTS,
) -> str:
    """Poll Toqan until the answer is finished or the poll limit is reached."""
    payload = {"conversation_id": conversation_id, "request_id": request_id}
    for attempt in range(max_attempts):
        data = _toqan_request("get_answer", payload, api_key)
        if data.get("status") == "finished":
            answer = data.get("answer", "")
            if not isinstance(answer, str):
                raise ToqanAPIError("Toqan returned a non-text answer")
            return _clean_answer(answer)

        progress_message = {
            18: "Ainda processando — essa é uma consulta complexa, só um momento...",
            42: "Ainda consultando o Databricks, quase lá...",
            72: "Demorando mais que o normal — a consulta é pesada, aguarde...",
            108: "Continuando a consulta, não desista — já estamos em 18 min...",
        }.get(attempt)
        if progress_message and progress_callback:
            try:
                progress_callback(progress_message)
            except Exception:
                logger.warning("Could not send Toqan progress message", exc_info=True)

        if attempt < max_attempts - 1:
            time.sleep(POLL_INTERVAL_SECONDS)
    raise ToqanAPIError("Toqan answer polling timed out")


TOQAN_UPLOAD_URL = os.getenv("TOQAN_UPLOAD_URL", f"{TOQAN_BASE_URL}/upload_file")


def toqan_upload_file(content: bytes, filename: str, mimetype: str = "application/octet-stream") -> str:
    """Sobe um arquivo (ex.: print do Slack) para o Toqan e devolve o file_id."""
    boundary = f"----lyra{int(time.time() * 1000)}"
    safe_name = re.sub(r"[^\w.\-]", "_", filename or "arquivo")
    body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"{safe_name}\"\r\n"
            f"Content-Type: {mimetype}\r\n\r\n").encode("utf-8") + content + f"\r\n--{boundary}--\r\n".encode("utf-8")
    req = urllib.request.Request(TOQAN_UPLOAD_URL, data=body, method="PUT", headers={
        "X-Api-Key": _require_toqan_key(), "Content-Type": f"multipart/form-data; boundary={boundary}"})
    try:
        with urllib.request.urlopen(req, timeout=TOQAN_TIMEOUT_SECONDS) as response:
            data = json.loads(response.read().decode("utf-8") or "{}")
    except urllib.error.HTTPError as exc:
        raise ToqanAPIError(f"Toqan upload HTTP {exc.code}") from exc
    except (urllib.error.URLError, TimeoutError, OSError, ValueError) as exc:
        raise ToqanAPIError(f"Toqan upload failed: {exc}") from exc
    file_id = data.get("file_id") or data.get("id") if isinstance(data, dict) else None
    if not file_id:
        raise ToqanAPIError("Toqan upload did not return file_id")
    return str(file_id)


def ask_toqan_with_files(prompt: str, file_ids: list[str], max_attempts: int = MAX_POLL_ATTEMPTS) -> str:
    """Pergunta com arquivos anexados (private_user_files). Se a criação com arquivo não for aceita,
    abre a conversa e manda o pedido com os arquivos no continue_conversation."""
    files = [{"id": fid} for fid in file_ids]
    try:
        data = _toqan_request("create_conversation", {"user_message": prompt, "private_user_files": files})
        conversation_id, request_id = data.get("conversation_id"), data.get("request_id")
        if conversation_id and request_id:
            return get_answer(str(conversation_id), str(request_id), max_attempts=max_attempts)
    except ToqanAPIError:
        logger.info("create_conversation with files not accepted; trying continue_conversation")
    conversation_id, request_id = create_conversation("Vou te mandar imagens para ler. Responda só: OK.")
    get_answer(conversation_id, request_id, max_attempts=30)
    data = _toqan_request("continue_conversation", {"conversation_id": conversation_id, "user_message": prompt,
                                                    "private_user_files": files})
    request_id = data.get("request_id")
    if not request_id:
        raise ToqanAPIError("continue_conversation did not return request_id")
    return get_answer(conversation_id, str(request_id), max_attempts=max_attempts)


def _clean_answer(answer: str) -> str:
    return _THINK_RE.sub("", answer).strip()


def ask_toqan(prompt: str, progress_callback: Callable[[str], Any] | None = None,
              max_attempts: int = MAX_POLL_ATTEMPTS) -> str:
    conversation_id, request_id = create_conversation(prompt)
    return get_answer(conversation_id, request_id, progress_callback=progress_callback, max_attempts=max_attempts)


def _strip_mentions(text: str) -> str:
    return re.sub(r"\s+", " ", _MENTION_RE.sub("", text)).strip()


def _event_key(event: dict[str, Any]) -> str:
    return str(event.get("client_msg_id") or event.get("ts") or event.get("text") or "unknown")


# ---------------------------------------------------------------------------
# Google Sheets
# ---------------------------------------------------------------------------
class MOPGoogleCredentialsError(RuntimeError):
    """Google service-account configuration is absent or invalid."""


class MOPDownloadError(RuntimeError):
    """An attached Slack file cannot be downloaded."""


class MOPReadError(RuntimeError):
    """The attached workbook cannot be read."""


class MOPValidationError(RuntimeError):
    """The workbook does not look like a MOP (missing columns or rows)."""


class MOPChannelNotAllowedError(RuntimeError):
    """MOP uploads are restricted to MOP_ALLOWED_CHANNELS."""


class MOPSheetsError(RuntimeError):
    """The destination Google Sheet cannot be updated."""


def _get_google_client():
    """Create the Google client from the service-account JSON (base64 or raw), if set."""
    raw = os.getenv("GOOGLE_CREDENTIALS_JSON")
    if not raw:
        return None
    try:
        try:
            credentials_json = json.loads(base64.b64decode(raw))
        except Exception:
            credentials_json = json.loads(raw)
        from google.oauth2.service_account import Credentials
        import gspread

        scopes = [
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ]
        credentials = Credentials.from_service_account_info(credentials_json, scopes=scopes)
        return gspread.authorize(credentials)
    except Exception as exc:
        logger.error("Could not configure Google Sheets client: %s", exc)
        raise MOPGoogleCredentialsError(str(exc)) from exc


def _safe_google_client():
    try:
        return _get_google_client()
    except Exception:
        logger.warning("Google Sheets unavailable", exc_info=True)
        return None


def _open_spreadsheet(google_client: Any) -> Any:
    """Planilha da Lyra (Compromissos, Estado Lyra, MOP backup)."""
    return google_client.open_by_key(LYRA_SPREADSHEET_ID)


def _open_mop_spreadsheet(google_client: Any) -> Any:
    """Planilha do MOP oficial."""
    return google_client.open_by_key(MOP_SPREADSHEET_ID)


def _is_not_found(exc: Exception) -> bool:
    return exc.__class__.__name__ == "WorksheetNotFound"


def _worksheet_by_gid(spreadsheet: Any, worksheet_gid: int) -> Any:
    get_by_id = getattr(spreadsheet, "get_worksheet_by_id", None)
    if callable(get_by_id):
        try:
            worksheet = get_by_id(worksheet_gid)
            if worksheet is not None:
                return worksheet
        except Exception as exc:
            if not _is_not_found(exc):
                raise
    for worksheet in spreadsheet.worksheets():
        if int(worksheet.id) == worksheet_gid:
            return worksheet
    raise MOPSheetsError(f"Worksheet gid {worksheet_gid} was not found")


def _get_or_create_worksheet(spreadsheet: Any, title: str, rows: int = 1000, cols: int = 26) -> Any:
    try:
        return spreadsheet.worksheet(title)
    except Exception as exc:
        if not _is_not_found(exc):
            raise
        return spreadsheet.add_worksheet(title, rows=rows, cols=cols)


def _ensure_worksheet(spreadsheet: Any, title: str, headers: list[str]) -> Any:
    """Garante a aba, o número de colunas e o cabeçalho."""
    worksheet = _get_or_create_worksheet(spreadsheet, title, rows=1000, cols=max(len(headers), 8))
    col_count = int(getattr(worksheet, "col_count", len(headers)) or len(headers))
    if col_count < len(headers):
        worksheet.add_cols(len(headers) - col_count)
    current = [str(v).strip() for v in worksheet.row_values(1)[: len(headers)]]
    if current != headers:
        worksheet.update(
            values=[headers],
            range_name=f"A1:{_col_letter(len(headers))}1",
            value_input_option="RAW",
        )
    return worksheet


def _worksheet_rows_with_header(worksheet: Any) -> list[tuple[int, dict[str, str]]]:
    values = worksheet.get_all_values()
    if not values:
        return []
    headers = [str(value).strip() for value in values[0]]
    return [
        (n, {h: str(row[i]) if i < len(row) else "" for i, h in enumerate(headers) if h})
        for n, row in enumerate(values[1:], start=2)
        if any(str(cell).strip() for cell in row)
    ]


# ---------------------------------------------------------------------------
# Estado persistido (aba "Estado Lyra")
# ---------------------------------------------------------------------------
class StateStore:
    """Chave-valor gravado na planilha, com cache em memória.

    Se o Google estiver indisponível, funciona só em memória (e avisa no log):
    nesse modo um restart pode repetir uma rotina.
    """

    def __init__(self) -> None:
        self._lock = threading.Lock()
        self._values: dict[str, str] = {}
        self._rows: dict[str, int] = {}
        self._worksheet: Any = None
        self._next_row = 2
        self._loaded = False

    def _ensure_loaded(self) -> None:
        if self._loaded:
            return
        google_client = _safe_google_client()
        if google_client is None:
            return
        try:
            worksheet = _ensure_worksheet(_open_spreadsheet(google_client), STATE_WORKSHEET_TITLE, STATE_HEADERS)
            values = worksheet.get_all_values()
        except Exception:
            logger.warning("State store: could not load worksheet; using memory only", exc_info=True)
            return
        for n, row in enumerate(values[1:], start=2):
            key = str(row[0]).strip() if row else ""
            if key:
                self._rows[key] = n
                # o que já está em memória (gravado antes do load) prevalece
                self._values.setdefault(key, str(row[1]) if len(row) > 1 else "")
        self._next_row = len(values) + 1
        self._worksheet = worksheet
        self._loaded = True

    def get(self, key: str) -> str:
        with self._lock:
            self._ensure_loaded()
            return self._values.get(key, "")

    def set(self, key: str, value: str) -> None:
        with self._lock:
            self._ensure_loaded()
            self._values[key] = value
            if self._worksheet is None:
                return
            stamp = _now_brazil().strftime("%d/%m/%Y %H:%M")
            try:
                if key in self._rows:
                    n = self._rows[key]
                    self._worksheet.update(
                        values=[[key, value, stamp]], range_name=f"A{n}:C{n}", value_input_option="RAW"
                    )
                else:
                    self._worksheet.append_row([key, value, stamp], value_input_option="RAW")
                    self._rows[key] = self._next_row
                    self._next_row += 1
            except Exception:
                logger.warning("State store: could not persist %s", key, exc_info=True)

    def snapshot(self, prefix: str = "") -> dict[str, str]:
        with self._lock:
            self._ensure_loaded()
            return {k: v for k, v in self._values.items() if k.startswith(prefix)}

    @property
    def persistent(self) -> bool:
        return self._worksheet is not None


STATE = StateStore()


# ---------------------------------------------------------------------------
# MOP
# ---------------------------------------------------------------------------
def _is_xlsx_attachment(file_info: dict[str, Any]) -> bool:
    name = str(file_info.get("name") or "").lower()
    mimetype = str(file_info.get("mimetype") or "").lower()
    filetype = str(file_info.get("filetype") or "").lower()
    if name.endswith(".xlsx"):
        return True
    if "spreadsheet" in mimetype or "excel" in mimetype:
        return True
    return filetype in ("xlsx", "excel")


def _download_slack_file(file_info: dict[str, Any], client: Any) -> bytes:
    details = dict(file_info)
    file_id = details.get("id")
    if file_id and client is not None and not (details.get("url_private_download") or details.get("url_private")):
        try:
            remote_file = _resp_get(client.files_info(file=file_id), "file")
            if isinstance(remote_file, dict):
                details.update(remote_file)
        except Exception as exc:
            logger.warning("Slack files.info failed for attached file: %s", exc)
            raise MOPDownloadError("files.info falhou (verifique o scope files:read)") from exc

    file_url = details.get("url_private_download") or details.get("url_private")
    if not file_url:
        raise MOPDownloadError("o arquivo não tem link de download")

    headers = {}
    slack_token = os.getenv("SLACK_BOT_TOKEN")
    if slack_token:
        headers["Authorization"] = f"Bearer {slack_token}"
    req = urllib.request.Request(str(file_url), headers=headers)
    try:
        with urllib.request.urlopen(req, timeout=TOQAN_TIMEOUT_SECONDS) as response:
            content = response.read()
    except (urllib.error.URLError, TimeoutError, OSError) as exc:
        logger.error("Slack file download failed")
        raise MOPDownloadError("falha no download (verifique o scope files:read)") from exc
    if not content:
        raise MOPDownloadError("o download veio vazio")
    if content[:2] != b"PK":  # xlsx é um zip; HTML aqui = falta de permissão
        raise MOPDownloadError("o Slack não liberou o arquivo (verifique o scope files:read)")
    return content


def _cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):  # antes de date: datetime é subclasse de date
        if value.time() == datetime_time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, datetime_time):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _compact(value: Any) -> str:
    return _norm_text(value).replace(" ", "")


def _find_col(header: list[Any], column: str) -> int | None:
    """Índice da coluna (exato, depois prefixo), ignorando acento, caixa, espaço e hífen."""
    target = _compact(column)
    cells = [_compact(h) for h in header]
    for i, cell in enumerate(cells):
        if cell and cell == target:
            return i
    for i, cell in enumerate(cells):
        if cell and cell.startswith(target):
            return i
    return None


def _header_matches(header_cells: list[str], column: str) -> bool:
    return _find_col(header_cells, column) is not None


def _eps_matches(value: Any, eps: str) -> bool:
    v, e = _compact(value), _compact(eps)
    return bool(v) and (v == e or e in v)


def _read_mop_workbook(workbook_bytes: bytes) -> tuple[list[list[str]], str, list[str]]:
    """Encontra a aba/linha de cabeçalho do MOP e devolve (linhas, aba, colunas ausentes).

    Lê VALORES (data_only=True), nunca fórmulas.
    """
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=True, read_only=True)
    except Exception as exc:
        logger.error("Excel workbook could not be opened: %s", exc)
        raise MOPReadError(str(exc)) from exc

    try:
        chosen: tuple[Any, int] | None = None
        for worksheet in workbook.worksheets:
            for idx, row in enumerate(
                worksheet.iter_rows(min_row=1, max_row=MOP_HEADER_SCAN_ROWS, values_only=True), start=1
            ):
                cells = [_norm_text(_cell_to_str(c)) for c in row]
                if all(_header_matches(cells, col) for col in MOP_REQUIRED_COLUMNS):
                    chosen = (worksheet, idx)
                    break
            if chosen:
                break
        if chosen is None:
            raise MOPValidationError(
                "não encontrei uma aba com as colunas Nome e Email nas primeiras linhas"
            )
        worksheet, header_row = chosen
        rows: list[list[str]] = []
        for row in worksheet.iter_rows(min_row=header_row, values_only=True):
            converted = [_cell_to_str(v) for v in row]
            while converted and converted[-1] == "":
                converted.pop()
            rows.append(converted)
        title = str(worksheet.title)
    except MOPValidationError:
        raise
    except Exception as exc:
        logger.error("Excel workbook could not be read: %s", exc)
        raise MOPReadError(str(exc)) from exc
    finally:
        try:
            workbook.close()
        except Exception:
            pass

    while rows and not any(rows[-1]):
        rows.pop()
    data_rows = [r for r in rows[1:] if any(r)]
    if not data_rows:
        raise MOPValidationError("a aba do MOP não tem nenhuma pessoa abaixo do cabeçalho")
    header_cells = [_norm_text(c) for c in rows[0]]
    missing = [col for col in MOP_EXPECTED_COLUMNS if not _header_matches(header_cells, col)]
    return rows, title, missing


def _count_people(values: list[list[Any]]) -> int:
    return len([r for r in values[1:] if any(str(c).strip() for c in r)])


def _row_eps(row: list[Any], eps_col: int | None, email_col: int | None) -> str:
    """EPS da linha do MOP: coluna EPS; se vazia, deduz pelo e-mail."""
    value = str(row[eps_col]).strip() if eps_col is not None and eps_col < len(row) else ""
    if value:
        return value
    email = row[email_col] if email_col is not None and email_col < len(row) else ""
    return _email_eps(email) or ""


def _count_eps_people(values: list[list[Any]], eps: str) -> int:
    if not values:
        return 0
    eps_col, email_col = _find_col(values[0], "EPS"), _find_col(values[0], "Email")
    return len([
        r for r in values[1:]
        if any(str(c).strip() for c in r) and _eps_matches(_row_eps(r, eps_col, email_col), eps)
    ])


def _merge_mop_rows(current: list[list[str]], new_rows: list[list[str]], eps: str) -> tuple[list[list[str]], int, int]:
    """Troca só as linhas da EPS e mantém as das outras.

    Devolve (planilha final, linhas de outra EPS ignoradas no arquivo, pessoas da EPS no arquivo).
    """
    new_header = new_rows[0]
    if current and _find_col(current[0], "Email") is not None:
        target_header = list(current[0])
    else:
        target_header = list(new_header)
    if _find_col(target_header, "EPS") is None:
        target_header.append("EPS")
    eps_col = _find_col(target_header, "EPS")
    email_col = _find_col(target_header, "Email")
    mapping = [_find_col(new_header, h) if str(h).strip() else None for h in target_header]
    new_eps_col, new_email_col = _find_col(new_header, "EPS"), _find_col(new_header, "Email")

    ignored = 0
    incoming: list[list[str]] = []
    for row in new_rows[1:]:
        if not any(str(c).strip() for c in row):
            continue
        row_eps = _row_eps(row, new_eps_col, new_email_col)
        if row_eps and not _eps_matches(row_eps, eps):
            ignored += 1
            continue
        out = [row[j] if j is not None and j < len(row) else "" for j in mapping]
        out[eps_col] = row_eps or eps
        incoming.append(out)

    kept = [
        list(r) for r in current[1:]
        if any(str(c).strip() for c in r) and not _eps_matches(_row_eps(r, eps_col, email_col), eps)
    ] if current and _find_col(current[0], "Email") is not None else []
    return [target_header] + kept + incoming, ignored, len(incoming)


def _update_mop_sheet(
    rows: list[list[str]], google_client: Any, eps: str | None = None
) -> tuple[int, int, int]:
    """Atualiza a aba do MOP com backup e rollback. Devolve (antes, depois, ignoradas).

    Com eps: substitui só as linhas daquela EPS (antes/depois contam só a EPS).
    Sem eps: substitui a aba inteira.
    """
    try:
        spreadsheet = _open_spreadsheet(google_client)  # onde fica o backup
        worksheet = _worksheet_by_gid(_open_mop_spreadsheet(google_client), MOP_WORKSHEET_GID)
        current = worksheet.get_all_values()
    except MOPSheetsError:
        raise
    except Exception as exc:
        raise MOPSheetsError(str(exc)) from exc

    ignored = 0
    if eps:
        before = _count_eps_people(current, eps)
        rows, ignored, after_eps = _merge_mop_rows(current, rows, eps)
        if after_eps == 0:
            raise MOPValidationError(f"o arquivo não tem nenhuma pessoa da {eps}")
    width = max([len(r) for r in rows] + [len(r) for r in current] + [1])

    # 1) backup da versão atual
    try:
        backup = _get_or_create_worksheet(
            spreadsheet, MOP_BACKUP_WORKSHEET_TITLE, rows=max(len(current) + 10, 1000), cols=max(width, 26)
        )
        backup.resize(rows=max(len(current) + 10, int(backup.row_count)), cols=max(width, int(backup.col_count)))
        backup.clear()
        if current:
            backup.update(values=current, range_name="A1", value_input_option="RAW")
    except Exception as exc:
        logger.error("MOP backup failed; aborting update: %s", exc)
        raise MOPSheetsError("não consegui salvar o backup da versão atual") from exc

    # 2) substitui; se falhar, restaura
    try:
        worksheet.resize(rows=max(len(rows) + 10, int(worksheet.row_count)), cols=max(width, int(worksheet.col_count)))
        worksheet.clear()
        worksheet.update(values=rows, range_name="A1", value_input_option="RAW")
    except Exception as exc:
        logger.error("Google Sheets MOP update failed, restoring backup: %s", exc)
        try:
            worksheet.clear()
            if current:
                worksheet.update(values=current, range_name="A1", value_input_option="RAW")
        except Exception:
            logger.exception("MOP restore failed")
            _alert("Falha ao restaurar o MOP após erro de atualização. Use a aba 'MOP backup'.")
        raise MOPSheetsError(str(exc)) from exc

    if eps:
        after = after_eps
    else:
        before, after = _count_people(current), _count_people(rows)
    logger.info("MOP sheet updated (%s): %d -> %d people", eps or "todas", before, after)
    return before, after, ignored


def _mop_column_index(header_index: dict[str, int], *aliases: str) -> int | None:
    for alias in aliases:
        idx = header_index.get(alias.casefold())
        if idx is not None:
            return idx
    for alias in aliases:
        key = alias.casefold()
        for candidate, idx in header_index.items():
            if key in candidate:
                return idx
    return None


def _mop_counts(rows: list[list[Any]], idx: int | None) -> dict[str, int] | None:
    if idx is None:
        return None
    counts: dict[str, int] = {}
    display: dict[str, str] = {}
    for row in rows:
        text = str(row[idx] if idx < len(row) else "").strip()
        if not text:
            continue
        key = text.casefold()
        display.setdefault(key, text)
        counts[key] = counts.get(key, 0) + 1
    if not counts:
        return None
    return {display[key]: count for key, count in counts.items()}


def _read_mop_values(google_client: Any = None) -> list[list[str]] | None:
    """Valores da aba do MOP; None se não deu para ler."""
    google_client = google_client or _safe_google_client()
    if google_client is None:
        return None
    try:
        return _worksheet_by_gid(_open_mop_spreadsheet(google_client), MOP_WORKSHEET_GID).get_all_values()
    except Exception:
        logger.warning("MOP: could not read MOP worksheet", exc_info=True)
        return None


def _mop_summary_text(google_client: Any = None, eps: str | None = None, values: list[list[str]] | None = None) -> str:
    """Resumo compacto do MOP para injetar no prompt ("" se indisponível ou sem pessoas).

    Com eps, considera só as pessoas daquela EPS.
    """
    all_values = values if values is not None else _read_mop_values(google_client)
    if all_values is None:
        return ""
    if not all_values:
        return ""
    header_index: dict[str, int] = {}
    for idx, raw in enumerate(all_values[0]):
        text = str(raw).strip()
        if text:
            header_index.setdefault(text.casefold(), idx)
    data_rows = [row for row in all_values[1:] if any(str(cell).strip() for cell in row)]
    if eps:
        eps_col, email_col = _find_col(all_values[0], "EPS"), _find_col(all_values[0], "Email")
        data_rows = [r for r in data_rows if _eps_matches(_row_eps(r, eps_col, email_col), eps)]
    if not data_rows:
        return ""

    scope = f" da {eps}" if eps else ""
    lines: list[str] = [f"Total de pessoas{scope} no MOP: {len(data_rows)}"]

    def add_group(label: str, *aliases: str, limit: int = 15) -> None:
        counts = _mop_counts(data_rows, _mop_column_index(header_index, *aliases))
        if counts is None:
            return
        lines.extend(["", f"{label}:"])
        ordered = sorted(counts.items(), key=lambda item: (-item[1], item[0].casefold()))
        for value, count in ordered[:limit]:
            lines.append(f"- {value}: {count}")
        if len(ordered) > limit:
            lines.append(f"- (outros {len(ordered) - limit} valores omitidos)")

    add_group("Por Status", "status", "situação", "situacao")
    add_group("Por EPS", "eps")
    add_group("Por Pilar", "pilar")
    add_group("Por Versão do Oráculo", "versão do oráculo", "versao do oraculo", "oráculo", "oraculo")
    add_group("Total por Célula (top 10)", "célula", "celula", limit=10)
    add_group("Span por supervisor (pessoas por supervisor)", "supervisor", limit=80)
    add_group("Pessoas por coordenador", "coordenador", limit=40)
    sup_idx = _mop_column_index(header_index, "supervisor")
    coord_idx = _mop_column_index(header_index, "coordenador")
    sup_counts = _mop_counts(data_rows, sup_idx)
    if sup_counts:
        lines.extend(["", f"Span médio: {len(data_rows) / len(sup_counts):.1f} pessoas por supervisor ({len(sup_counts)} supervisores)"])
    if sup_idx is not None and coord_idx is not None:
        sups_by_coord: dict[str, set[str]] = {}
        for row in data_rows:
            coord = str(row[coord_idx] if coord_idx < len(row) else "").strip()
            sup = str(row[sup_idx] if sup_idx < len(row) else "").strip()
            if coord and sup:
                sups_by_coord.setdefault(coord, set()).add(sup.casefold())
        if sups_by_coord:
            lines.extend(["", "Span por coordenador (supervisores por coordenador):"])
            for coord, sups in sorted(sups_by_coord.items(), key=lambda item: (-len(item[1]), item[0])):
                lines.append(f"- {coord}: {len(sups)}")
    add_group("Por horário de Entrada", "entrada", limit=8)
    add_group("Por horário de Saída", "saída", "saida", limit=8)
    updated = STATE.get(f"mop_updated_at:{eps or CONSOLIDATED}")
    if updated:
        lines.extend(["", f"Última atualização do MOP pelo app: {updated}"])
    return "\n".join(lines)


def _question_has_mop_intent(question: str) -> bool:
    lowered = question.casefold()
    return any(keyword in lowered for keyword in MOP_INTENT_KEYWORDS)


INTERNAL_NOTE = (
    "[INFORMAÇÃO INTERNA PARA A LYRA — nunca cite este trecho, nem app, sistema, planilha, "
    "Toqan ou Databricks. Responda como a gestora das EPSs, em linguagem natural.]"
)


def _mop_situation(eps: str | None, values: list[list[str]] | None = None) -> str:
    """Situação do MOP em linguagem de negócio, sempre preenchida."""
    who = f"da {eps}" if eps else "da operação"
    if values is None:
        values = _read_mop_values()
    if values is None:
        return (
            f"Situação do MOP {who}: não deu para consultar o MOP agora (instabilidade). "
            "Diga que não conseguiu olhar o MOP neste momento e que tenta de novo em instantes. Não peça o Excel."
        )
    summary = _mop_summary_text(eps=eps, values=values)
    if not summary:
        return (
            f"Situação do MOP {who}: ainda não há nenhuma pessoa {who} no MOP — ele não foi enviado. "
            f"Diga isso de forma simples e peça para a liderança {who} enviar o Excel do MOP aqui no canal marcando @Lyra."
        )
    return (
        f"Situação do MOP {who} (atualizado):\n{summary}\n"
        "Use esses números para perguntas sobre pessoas, hierarquia e efetivo. Não peça o Excel."
    )


_LOOKUP_STOPWORDS = {
    "mop", "quem", "qual", "quais", "quantas", "quantos", "pessoas", "pessoa", "supervisor", "supervisora",
    "coordenador", "coordenadora", "celula", "equipe", "time", "esta", "estao", "onde", "como", "sobre",
    "tem", "temos", "para", "pela", "pelo", "dele", "dela", "deles", "delas", "esse", "essa", "isso",
    "lyra", "analista", "analistas", "operador", "operadores", "lista", "manda", "mandar", "favor",
    "turno", "entrada", "saida", "status", "ativo", "ativos", "ferias", "licenca", "maior", "menor",
    "span", "spam", "gestao", "gestor", "hoje", "agora", "todos", "todas", "aqui", "voces", "nosso", "nossa",
    "trabalha", "trabalham", "horario", "formato", "remoto", "presencial", "oraculo", "versao", "variante",
}
MOP_LOOKUP_LIMIT = 40


def _mop_lookup_rows(question: str, values: list[list[str]], eps: str | None) -> list[str]:
    """Linhas do MOP que batem com nomes/e-mails/supervisores citados na pergunta (só da EPS do canal)."""
    if not values or len(values) < 2:
        return []
    header = values[0]
    tokens = [t for t in re.findall(r"[a-z0-9.]+", _norm_text(question)) if len(t) >= 4 and t not in _LOOKUP_STOPWORDS]
    if not tokens:
        return []
    cols = {name: _find_col(header, name) for name in (
        "Nome", "Email", "Célula", "Supervisor", "Coordenador", "Entrada", "Saída", "Status",
        "Versão do Oráculo", "Formato de Trabalho", "EPS", "Pilar")}
    eps_col, email_col = cols["EPS"], cols["Email"]

    def cell(row: list[str], name: str) -> str:
        idx = cols[name]
        return str(row[idx]).strip() if idx is not None and idx < len(row) else ""

    scored: list[tuple[int, list[str]]] = []
    for row in values[1:]:
        if not any(str(c).strip() for c in row):
            continue
        if eps and not _eps_matches(_row_eps(row, eps_col, email_col), eps):
            continue
        blob = _norm_text(" ".join(cell(row, n) for n in ("Nome", "Email", "Supervisor", "Coordenador", "Célula")))
        score = sum(1 for t in set(tokens) if t in blob)
        if score:
            scored.append((score, row))
    if not scored:
        return []
    best = max(score for score, _ in scored)
    rows = [row for score, row in scored if score == best][:MOP_LOOKUP_LIMIT]
    lines = []
    for row in rows:
        user = cell(row, "Email").split("@")[0]
        lines.append(
            f"- {cell(row, 'Nome')} ({user}) · célula {cell(row, 'Célula') or '—'} · supervisor {cell(row, 'Supervisor') or '—'} "
            f"· coordenador {cell(row, 'Coordenador') or '—'} · {cell(row, 'Entrada') or '—'}–{cell(row, 'Saída') or '—'} "
            f"· {cell(row, 'Status') or '—'} · {cell(row, 'Versão do Oráculo') or '—'} · {cell(row, 'Formato de Trabalho') or '—'}"
        )
    total = len([1 for score, _ in scored if score == best])
    if total > len(rows):
        lines.append(f"- (+{total - len(rows)} pessoas)")
    return lines


_EXPORT_RE = re.compile(
    r"\b(manda|mande|mandar|envia|envie|enviar|extrai|extraia|extrair|exporta|exporte|exportar|baixa|baixar|"
    r"gera|gere|gerar|passa|passe|compartilha|compartilhe)\b.*\bmop\b|\bmop\b.*\b(excel|xlsx|planilha|arquivo)\b",
    re.IGNORECASE,
)


def _is_mop_export_request(question: str) -> bool:
    return bool(_EXPORT_RE.search(_norm_text(question)))


def _build_mop_xlsx(values: list[list[str]], eps: str | None) -> tuple[bytes, int]:
    """Excel do MOP (só a EPS do canal, ou tudo na visão interna). Devolve (bytes, pessoas)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    header = values[0]
    eps_col, email_col = _find_col(header, "EPS"), _find_col(header, "Email")
    rows = [
        r for r in values[1:]
        if any(str(c).strip() for c in r) and (not eps or _eps_matches(_row_eps(r, eps_col, email_col), eps))
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = f"MOP {eps}" if eps else "MOP Consolidado"
    ws.append(list(header))
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="EA1D2C")
    for r in rows:
        ws.append(list(r) + [""] * (len(header) - len(r)))
    for idx, name in enumerate(header, start=1):
        width = max([len(str(name))] + [len(str(r[idx - 1])) for r in rows[:300] if idx - 1 < len(r)])
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max(width + 2, 10), 45)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(rows)


def _send_mop_export(event: dict[str, Any], client: Any, eps: str | None) -> str | None:
    """Gera e envia o Excel do MOP na thread. Devolve texto de resposta se não deu para enviar."""
    values = _read_mop_values()
    if values is None:
        return "Não consegui puxar o MOP agora. Tenta de novo em alguns minutos?"
    if len(values) < 2:
        return "O MOP ainda está vazio por aqui."
    content, people = _build_mop_xlsx(values, eps)
    if people == 0:
        return f"Ainda não tenho ninguém da {eps} no MOP." if eps else "O MOP ainda está vazio por aqui."
    who = f" da {eps}" if eps else ""
    stamp = _now_brazil().strftime("%d-%m-%Y")
    uploader = client or slack_app.client
    uploader.files_upload_v2(
        channel=event.get("channel"),
        thread_ts=_reply_ts(event),
        content=content,
        filename=f"MOP{('_' + eps) if eps else '_Consolidado'}_{stamp}.xlsx",
        title=f"MOP{who} — {stamp}",
        initial_comment=f"Segue o MOP{who} atualizado 📋",
    )
    return None


def _inject_mop_context(question: str, eps: str | None = None) -> str:
    if not question:
        return question
    values = _read_mop_values() if _question_has_mop_intent(question) else None
    people = _mop_lookup_rows(question, values or [], eps) if values else []
    if not _question_has_mop_intent(question) and not people:
        return question
    extra = ""
    if people:
        extra = "\nPessoas do MOP citadas na pergunta (nome · usuário · célula · supervisor · coordenador · turno · status · Oráculo · formato):\n" + "\n".join(people)
    return f"{INTERNAL_NOTE}\n{_mop_situation(eps, values)}{extra}\n\nMensagem de quem perguntou: {question}"


def _mop_attachment(event: dict[str, Any]) -> dict[str, Any] | None:
    files = event.get("files") or []
    if not isinstance(files, list):
        return None
    return next((f for f in files if isinstance(f, dict) and _is_xlsx_attachment(f)), None)


def _mop_channel_allowed(event: dict[str, Any]) -> bool:
    if not MOP_ALLOWED_CHANNELS:
        return True
    return str(event.get("channel") or "") in MOP_ALLOWED_CHANNELS


def _process_mop_upload(event: dict[str, Any], client: Any, eps: str | None = None) -> str | None:
    """Processa um .xlsx; devolve None quando o evento não é upload de Excel.

    Com eps (canal de EPS): troca só as linhas daquela EPS.
    """
    attachment = _mop_attachment(event)
    if attachment is None:
        return None
    if not _mop_channel_allowed(event):
        raise MOPChannelNotAllowedError
    google_client = _get_google_client()
    if google_client is None:
        raise MOPGoogleCredentialsError
    workbook_bytes = _download_slack_file(attachment, client)
    rows, sheet_title, missing = _read_mop_workbook(workbook_bytes)
    before, after, ignored = _update_mop_sheet(rows, google_client, eps)
    _mark_mop_updated([eps] if eps else None)
    if eps:
        missing = [m for m in missing if _compact(m) != "eps"]

    lines = ["MOP atualizado na planilha ✅"]
    if ignored:
        lines.append(f"⚠️ {ignored} linha(s) de outra EPS no arquivo foram ignoradas.")
    if missing:
        lines.append(f"⚠️ Colunas não encontradas: {', '.join(missing)} — confira o modelo de 12 colunas.")
    if before and after < before * 0.5:
        lines.append(
            "⚠️ O número de pessoas caiu mais de 50%. Se não for esperado, reenvie o arquivo correto — "
            f"a versão anterior está na aba _{MOP_BACKUP_WORKSHEET_TITLE}_."
        )
    return "\n".join(lines)


def _mark_mop_updated(eps_list: list[str] | None) -> None:
    """Registra a data da atualização (None = MOP inteiro, todas as EPS)."""
    now = _now_brazil()
    stamp = now.strftime("%d/%m/%Y %H:%M")
    keys = [CONSOLIDATED] + list(EPS_CELULAS) if eps_list is None else eps_list
    for key in keys:
        STATE.set(f"mop_updated_at:{key}", stamp)
    STATE.set("mop_last_update", now.isoformat())


def _file_eps_set(rows: list[list[str]]) -> set[str]:
    """EPS presentes no arquivo (pela coluna EPS ou pelo e-mail)."""
    if not rows:
        return set()
    eps_col, email_col = _find_col(rows[0], "EPS"), _find_col(rows[0], "Email")
    found: set[str] = set()
    for row in rows[1:]:
        value = _row_eps(row, eps_col, email_col)
        for known in EPS_CELULAS:
            if value and _eps_matches(value, known):
                found.add(known)
    return found


def _message_text_blob(message: dict[str, Any]) -> str:
    """Texto + blocos + anexos de uma mensagem (mensagens de fluxo guardam o texto nos blocos)."""
    parts = [str(message.get("text") or "")]
    for field in ("blocks", "attachments"):
        if message.get(field):
            parts.append(json.dumps(message.get(field), ensure_ascii=False))
    return " ".join(parts)


def _is_mop_flow_message(event: dict[str, Any]) -> bool:
    """Excel postado no canal do fluxo, em mensagem ou thread que fala de MOP."""
    if not MOP_FLOW_CHANNEL or event.get("channel") != MOP_FLOW_CHANNEL:
        return False
    if _mop_attachment(event) is None:
        return False
    if "mop" in _message_text_blob(event).casefold():
        return True
    thread_ts = event.get("thread_ts")
    if not thread_ts or thread_ts == event.get("ts"):
        return False
    try:
        response = slack_app.client.conversations_replies(channel=MOP_FLOW_CHANNEL, ts=thread_ts, limit=1)
        messages = _resp_get(response, "messages", []) or []
        parent = next((m for m in messages if isinstance(m, dict) and m.get("ts") == thread_ts), None)
    except Exception:
        logger.warning("MOP flow: could not read parent message", exc_info=True)
        return False
    return parent is not None and "mop" in _message_text_blob(parent).casefold()


def _process_mop_flow(event: dict[str, Any], client: Any) -> str:
    """Atualiza o MOP a partir do Excel do fluxo. Arquivo com mais de uma EPS = MOP inteiro;
    arquivo com uma EPS só = troca só as linhas dela (não apaga a outra)."""
    google_client = _get_google_client()
    if google_client is None:
        raise MOPGoogleCredentialsError
    attachment = _mop_attachment(event)
    workbook_bytes = _download_slack_file(attachment, client)
    rows, sheet_title, missing = _read_mop_workbook(workbook_bytes)
    eps_found = _file_eps_set(rows)
    single_eps = next(iter(eps_found)) if len(eps_found) == 1 else None
    before, after, ignored = _update_mop_sheet(rows, google_client, single_eps)
    _mark_mop_updated([single_eps] if single_eps else None)

    lines = ["MOP atualizado na planilha ✅"]
    # Só fala mais se houver algo que precise de atenção
    if ignored:
        lines.append(f"⚠️ {ignored} linha(s) de outra EPS foram ignoradas.")
    if missing:
        lines.append(f"⚠️ Colunas não encontradas: {', '.join(missing)}.")
    if before and after < before * 0.5:
        lines.append("⚠️ O número de pessoas caiu mais de 50% — confira o arquivo. A versão anterior ficou guardada.")
    return "\n".join(lines)


def _handle_mop_flow_event(event: dict[str, Any], say: Callable[..., Any], client: Any) -> None:
    """Worker do fluxo: dedupe persistente (o mesmo arquivo pode chegar por dois eventos)."""
    ts = str(event.get("ts") or "")
    done_key = f"mop_flow:{ts}"
    with _IN_FLIGHT_LOCK:
        if ts in _IN_FLIGHT or STATE.get(done_key):
            return
        _IN_FLIGHT.add(ts)
    thread_ts = event.get("thread_ts") or ts
    try:
        if not _is_mop_flow_message(event):
            return
        STATE.set(done_key, "processando")
        message = _process_mop_flow(event, client)
        STATE.set(done_key, "ok")
        say(text=message, thread_ts=thread_ts)
    except MOPValidationError as exc:
        STATE.set(done_key, "recusado")
        say(text=f"Não atualizei o MOP: {exc}. O MOP atual foi mantido.", thread_ts=thread_ts)
    except (MOPDownloadError, MOPReadError) as exc:
        STATE.set(done_key, "erro")
        say(text="Não consegui abrir esse arquivo. Pode reenviar o Excel aqui na thread?", thread_ts=thread_ts)
        _alert(f"MOP do fluxo: falha ao ler o arquivo ({type(exc).__name__}: {str(exc)[:150]})")
    except Exception as exc:
        STATE.set(done_key, "erro")
        logger.exception("MOP flow update failed")
        say(text="Tive um problema pra atualizar o MOP agora. Já avisei o time responsável.", thread_ts=thread_ts)
        _alert(f"MOP do fluxo: falha ao atualizar ({type(exc).__name__}: {str(exc)[:150]})")
    finally:
        with _IN_FLIGHT_LOCK:
            _IN_FLIGHT.discard(ts)


# ---------------------------------------------------------------------------
# Report: leitura das filas
# ---------------------------------------------------------------------------
_NUM = r"\d{1,3}(?:[.,]\d+)?"
FILA_LINE_RE = re.compile(
    rf"\*(?P<fila>[^*\n]{{2,80}}?)\s*:?\s*\*\s*:?\s*(?:Mês|Mes|MTD)\s*(?P<mes>{_NUM})\s*%"
    rf"[^\n%]{{0,20}}?D\s*[-–]?\s*1\s*(?P<d1>{_NUM})\s*%",
    re.IGNORECASE,
)


def _parse_report_csat(text: str) -> dict[str, dict[str, float]]:
    """Extrai {fila: {"mes": x, "d1": y}} das linhas '• *Fila:* Mês X% · D-1 Y%'."""
    result: dict[str, dict[str, float]] = {}
    for match in FILA_LINE_RE.finditer(text or ""):
        fila = match.group("fila").strip().rstrip(":").strip()
        mes, d1 = _to_float(match.group("mes")), _to_float(match.group("d1"))
        if fila and mes is not None and d1 is not None:
            result.setdefault(fila, {"mes": mes, "d1": d1})
    return result


def _match_fila(name: str, filas: list[str]) -> str | None:
    """Casa um nome livre com uma fila conhecida (exato, depois o mais específico)."""
    target = _norm_text(name)
    if not target or not filas:
        return None
    normalized = {f: _norm_text(f) for f in filas}
    for fila, norm in normalized.items():
        if norm == target:
            return fila
    candidates = [f for f, n in normalized.items() if n and (n in target or target in n)]
    if not candidates:
        return None
    return max(candidates, key=lambda f: len(normalized[f]))


def _split_report(answer: str) -> tuple[str, str]:
    if THREAD_DELIMITER in answer:
        main, thread = answer.split(THREAD_DELIMITER, 1)
        return main.strip(), thread.strip()
    return answer.strip(), ""


def _is_error_answer(answer: str) -> bool:
    first = next((line.strip() for line in str(answer or "").splitlines() if line.strip()), "")
    return not first or first.upper().lstrip("*_ ").startswith("ERRO")


# ---------------------------------------------------------------------------
# Datas de compromisso
# ---------------------------------------------------------------------------
_NO_DATE = {"", "a definir", "sem prazo", "não informado", "nao informado", "—", "-", "n/a"}
_WEEKDAYS = {
    "segunda": 0, "terça": 1, "terca": 1, "quarta": 2, "quinta": 3,
    "sexta": 4, "sábado": 5, "sabado": 5, "domingo": 6,
}


def _midnight(day: date) -> datetime:
    return datetime.combine(day, datetime_time(0, 0), tzinfo=BRAZIL_TIMEZONE)


def _parse_commitment_date(value: Any, today: datetime | date | None = None) -> datetime | None:
    """Entende dd/mm[/aaaa], aaaa-mm-dd, hoje, amanhã, dia da semana, semana que vem, fim do mês…"""
    text = str(value or "").strip().casefold()
    if text in _NO_DATE:
        return None
    base = _as_date(today or _today_brazil())

    for pattern, fmt in (
        (r"\b(\d{1,2}/\d{1,2}/\d{4})\b", "%d/%m/%Y"),
        (r"\b(\d{1,2}-\d{1,2}-\d{4})\b", "%d-%m-%Y"),
        (r"\b(\d{4}-\d{1,2}-\d{1,2})\b", "%Y-%m-%d"),
    ):
        match = re.search(pattern, text)
        if match:
            try:
                return _midnight(datetime.strptime(match.group(1), fmt).date())
            except ValueError:
                return None
    match = re.search(r"\b(\d{1,2})/(\d{1,2})(?:/(\d{2}))?\b", text)
    if match:
        day, month = int(match.group(1)), int(match.group(2))
        year = 2000 + int(match.group(3)) if match.group(3) else base.year
        try:
            parsed = date(year, month, day)
            if not match.group(3) and parsed < base - timedelta(days=180):
                parsed = date(year + 1, month, day)
        except ValueError:
            return None
        return _midnight(parsed)

    if "depois de amanh" in text:
        return _midnight(base + timedelta(days=2))
    if "amanh" in text or "tomorrow" in text:
        return _midnight(base + timedelta(days=1))
    if "hoje" in text or "today" in text:
        return _midnight(base)
    if re.search(r"fi(m|nal) do m[eê]s", text):
        first_next = date(base.year + (base.month == 12), base.month % 12 + 1, 1)
        return _midnight(first_next - timedelta(days=1))
    if re.search(r"semana que vem|pr[oó]xima semana", text):
        return _midnight(base + timedelta(days=7))
    if re.search(r"fi(m|nal) da semana", text):
        return _midnight(base + timedelta(days=(4 - base.weekday()) % 7))
    match = re.search(r"\b(\d{1,2})\s*dias?\b", text)
    if match:
        return _midnight(base + timedelta(days=int(match.group(1))))
    for name, weekday in _WEEKDAYS.items():
        if re.search(rf"\b{name}", text):
            return _midnight(base + timedelta(days=(weekday - base.weekday()) % 7))
    return None


def _deadline_fields(raw: Any, today: datetime | date) -> tuple[str, str]:
    """(Data Fim, Duração em dias) a partir do prazo informado."""
    parsed = _parse_commitment_date(raw, today)
    if parsed is None:
        return "a definir", "—"
    return _format_date(parsed), str(max(0, (parsed.date() - _as_date(today)).days))


# ---------------------------------------------------------------------------
# Compromissos
# ---------------------------------------------------------------------------
def _compromissos_ws(google_client: Any) -> Any:
    return _ensure_worksheet(_open_spreadsheet(google_client), COMPROMISSOS_WORKSHEET_TITLE, COMPROMISSOS_HEADERS)


def _commitment_status(row: dict[str, str], today: datetime | date | None = None) -> str:
    """concluído | encerrado | sem resposta | prazo a definir | vence hoje | prazo vencido | expirado | em andamento"""
    base = _as_date(today or _today_brazil())
    if _row_value(row, "Sucesso").casefold() == "encerrado":
        return "encerrado"
    post = _row_value(row, "Resultado Pós")
    if post and post.casefold() not in {"a definir"}:
        return "concluído"
    if _row_value(row, "Ação Tomada").casefold() == SEM_RESPOSTA:
        return "sem resposta"
    deadline = _parse_commitment_date(_row_value(row, "Data Fim"), base)
    if deadline is None:
        return "prazo a definir"
    if deadline.date() == base:
        return "vence hoje"
    if deadline.date() < base - timedelta(days=EXPIRED_AFTER_DAYS):
        return "expirado"
    if deadline.date() < base:
        return "prazo vencido"
    return "em andamento"


_STATUS_PRIORITY = {
    "prazo vencido": 0, "vence hoje": 1, "sem resposta": 2, "prazo a definir": 3,
    "em andamento": 4, "concluído": 5,
}


def _commitment_fila(row: dict[str, str], filas: list[str] | None = None) -> str:
    fila = _row_value(row, "Fila")
    if fila:
        return fila
    if filas:  # linhas antigas, sem coluna Fila: tenta achar no texto do problema
        return _match_fila(_row_value(row, "Problema Identificado"), filas) or FILA_GERAL
    return FILA_GERAL


def _get_open_compromissos(
    google_client: Any, today: datetime | date | None = None, tenant_key: str = CONSOLIDATED
) -> str:
    """Contexto de compromissos para o prompt do report (só o que está em aberto, só desta EPS)."""
    base = _as_date(today or _today_brazil())
    try:
        rows = _worksheet_rows_with_header(_compromissos_ws(google_client))
    except Exception:
        logger.warning("Could not read commitments from Google Sheets", exc_info=True)
        return ""

    selected: list[tuple[int, date, str, dict[str, str]]] = []
    for _, row in rows:
        if not _same_tenant(row, tenant_key):
            continue
        status = _commitment_status(row, base)
        if status in {"encerrado", "expirado"}:
            continue
        deadline = _parse_commitment_date(_row_value(row, "Data Fim"), base)
        if status == "concluído":
            # mostra resultado só de quem fechou nos últimos 2 dias
            if deadline is None or deadline.date() < base - timedelta(days=2):
                continue
        if status == "sem resposta":
            start = _parse_commitment_date(_row_value(row, "Data Inicio"), base)
            if start is None or start.date() < base - timedelta(days=SEM_RESPOSTA_WINDOW_DAYS):
                continue
        sort_day = deadline.date() if deadline else base + timedelta(days=365)
        selected.append((_STATUS_PRIORITY.get(status, 9), sort_day, status, row))

    if not selected:
        return ""
    selected.sort(key=lambda item: (item[0], item[1]))
    lines = ["COMPROMISSOS EM ABERTO (aba Compromissos, lida agora pelo app):"]
    for _, _, status, row in selected[:OPEN_COMMITMENTS_LIMIT]:
        fila = _commitment_fila(row)
        action = _row_value(row, "Ação Tomada") or SEM_RESPOSTA
        owner = _row_value(row, "Responsável") or "Liderança"
        deadline_text = _row_value(row, "Data Fim") or "a definir"
        pre = _row_value(row, "Resultado Pré") or "a definir"
        if status == "em andamento":
            lines.append(f"- Em andamento: {fila} · {action} · {owner} · prazo {deadline_text}")
        elif status in {"prazo vencido", "vence hoje"}:
            label = "vence hoje" if status == "vence hoje" else f"venceu em {deadline_text}"
            lines.append(f"- Prazo {label}: {fila} · {action} · {owner} · CSAT pré {pre} — pedir o status")
        elif status == "concluído":
            post = _row_value(row, "Resultado Pós")
            success = _row_value(row, "Sucesso") or "a definir"
            lines.append(f"- Concluído: {fila} · {action} · CSAT {pre} → {post} — funcionou: {success}")
        elif status == "sem resposta":
            lines.append(f"- Sem resposta: report de {_row_value(row, 'Data Inicio')} — cobrar novamente")
        else:
            lines.append(f"- Sem prazo: {fila} · {action} · {owner} — pedir prazo")
    if len(selected) > OPEN_COMMITMENTS_LIMIT:
        lines.append(f"- (+{len(selected) - OPEN_COMMITMENTS_LIMIT} compromissos com prazo mais distante)")
    lines.extend([
        "",
        "Em andamento: não peça plano novo, só pergunte o status. Vencidos: peça o status. "
        "Sem resposta: cobre de novo. Ofensores novos: peça ação · responsável · prazo na thread até 14h.",
    ])
    return "\n".join(lines)


_REPORT_KPIS_RE = re.compile(r"(CSAT D-1 [\d.,]+%(?: \(meta [\d.,]+%\))?(?: · TMA [\d.]+ s(?: \(plan [\d.]+ s\))?)?)")
_REPORT_PROBLEM_RE = re.compile(r"Problema do D-1(?: \(CSAT\))?\s*:?\**\s*:?\s*([^\n]+)", re.IGNORECASE)


def _report_problem(report_text: str) -> str:
    """Resumo do que foi cobrado, tirado do bloco do report: indicadores do D-1 e o problema do D-1."""
    text = str(report_text or "").replace("*", "")
    parts = []
    kpis = _REPORT_KPIS_RE.search(text)
    if kpis:
        parts.append(kpis.group(1))
    problem = _REPORT_PROBLEM_RE.search(text)
    if problem:
        parts.append(problem.group(1).strip())
    return " · ".join(parts)


def _commitment_row(
    fila: str, acao: str, responsavel: str, prazo: str, pre: str, report_ts: str, today: datetime | date,
    tenant_key: str = CONSOLIDATED, problema_report: str = "",
) -> list[str]:
    data_fim, duracao = _deadline_fields(prazo, today)
    if problema_report:
        problema = f"{fila} · {problema_report}"
    else:
        problema = f"{fila} · CSAT D-1 {pre}" if pre else f"{fila} · ofensor do report"
    return [
        problema, acao, _format_date(today), data_fim, duracao,
        pre or "a definir", "", "a definir", fila, responsavel or "Liderança", report_ts, tenant_key,
    ]


def _sem_resposta_row(
    report_ts: str, today: datetime | date, tenant_key: str = CONSOLIDATED,
    fila: str = "", pre: str = "", problema_report: str = "",
) -> list[str]:
    problema = "Report sem resposta da liderança"
    if fila:
        problema = f"{fila} · sem resposta ao report" + (f" · {problema_report}" if problema_report else "")
    return [
        problema, SEM_RESPOSTA, _format_date(today), "—", "—",
        pre or "a definir", "", "a definir", fila or FILA_GERAL, "Liderança", report_ts, tenant_key,
    ]


def _row_tenant(row: dict[str, str]) -> str:
    return _row_value(row, "EPS") or CONSOLIDATED


def _same_tenant(row: dict[str, str], tenant_key: str) -> bool:
    return _compact(_row_tenant(row)) == _compact(tenant_key)


def _check_deadline_compromissos(
    google_client: Any, csat: dict[str, dict[str, float]], today: datetime | date | None = None,
    tenant_key: str = CONSOLIDATED,
) -> int:
    """Preenche Resultado Pós/Sucesso dos compromissos vencidos usando o D-1 da fila.

    Só fecha quando o prazo já passou (o D-1 do report cobre o dia do prazo).
    """
    if not csat:
        return 0
    base = _as_date(today or _today_brazil())
    filas = list(csat)
    worksheet = _compromissos_ws(google_client)
    updated = 0
    for row_number, row in _worksheet_rows_with_header(worksheet):
        if not _same_tenant(row, tenant_key):
            continue
        if _commitment_status(row, base) not in {"prazo vencido", "expirado"}:
            continue
        fila = _match_fila(_commitment_fila(row, filas), filas)
        pre = _to_float(_row_value(row, "Resultado Pré"))
        if fila is None or pre is None:
            continue
        post = csat[fila]["d1"]
        worksheet.update_cell(row_number, COL_POS, _fmt_pct(post))
        worksheet.update_cell(row_number, COL_SUCESSO, "Sim" if post > pre else "Não")
        updated += 1
    return updated


def _strip_json_code_blocks(answer: str) -> str:
    text = str(answer or "").strip()
    match = re.fullmatch(r"```(?:json)?\s*(.*?)\s*```", text, flags=re.IGNORECASE | re.DOTALL)
    return match.group(1).strip() if match else text


def _parse_extraction(answer: str) -> dict[str, Any] | None:
    text = _strip_json_code_blocks(answer)
    parsed: Any = None
    try:
        parsed = json.loads(text)
    except (json.JSONDecodeError, TypeError):
        match = re.search(r"(\{.*\}|\[.*\])", text, flags=re.DOTALL)
        if match:
            try:
                parsed = json.loads(match.group(1))
            except json.JSONDecodeError:
                parsed = None
    if isinstance(parsed, list):
        return {"novos": [i for i in parsed if isinstance(i, dict)], "atualizacoes": [], "feedback": ""}
    if isinstance(parsed, dict):
        return {
            "novos": [i for i in (parsed.get("novos") or []) if isinstance(i, dict)],
            "atualizacoes": [i for i in (parsed.get("atualizacoes") or []) if isinstance(i, dict)],
            "feedback": str(parsed.get("feedback") or "").strip(),
        }
    return None


def _extract_commitments_via_toqan(
    messages: list[str],
    filas: list[str],
    existing: list[tuple[int, dict[str, str]]],
    today: datetime | date,
    context: str = "",
) -> dict[str, Any]:
    """Pede ao Toqan os compromissos novos e as atualizações dos já registrados."""
    base = _as_date(today)
    weekday = ["segunda", "terça", "quarta", "quinta", "sexta", "sábado", "domingo"][base.weekday()]
    existing_lines = "\n".join(
        f"linha {n} · {_commitment_fila(r)} · {_row_value(r, 'Ação Tomada')} · "
        f"{_row_value(r, 'Responsável') or 'Liderança'} · prazo {_row_value(r, 'Data Fim') or 'a definir'}"
        for n, r in existing
        if _row_value(r, "Ação Tomada").casefold() != SEM_RESPOSTA
    ) or "(nenhum)"
    prompt = (
        "TAREFA INTERNA DO APP — NÃO é mensagem para o Slack. Responda APENAS com JSON válido, "
        "sem markdown, sem texto antes ou depois.\n\n"
        "Você é um extrator de compromissos de plano de ação da liderança das EPSs.\n"
        f"Hoje é {_format_date(base)} ({weekday}).\n"
        f"Filas do report: {', '.join(filas) if filas else '(não identificadas)'}\n"
        f"Compromissos já registrados nesta thread:\n{existing_lines}\n\n"
        + (f"Report desta fila (o que foi cobrado):\n{context[:1500]}\n\n" if context else "")
        +
        "Mensagens novas da liderança:\n" + "\n---\n".join(messages) + "\n\n"
        "Formato de saída:\n"
        '{"novos":[{"fila":"...","acao":"...","responsavel":"...","prazo":"dd/mm/aaaa ou a definir"}],'
        '"atualizacoes":[{"linha":12,"acao":"...","responsavel":"...","prazo":"dd/mm/aaaa"}],'
        '"feedback":"..."}\n\n'
        "Regras:\n"
        "- novos: planos de ação novos. fila = exatamente um nome da lista de filas; se não der pra saber, use \"Geral\".\n"
        "- atualizacoes: quando a mensagem completa ou corrige um compromisso já registrado (ex.: informa o prazo "
        "que faltava). Use o número da linha e inclua só os campos que mudaram.\n"
        "- prazo: converta para dd/mm/aaaa a partir de hoje (\"sexta\" = próxima sexta; \"semana que vem\" = hoje + 7 dias; "
        "\"fim do mês\" = último dia do mês). Sem prazo → \"a definir\".\n"
        "- responsavel: quem executa; se não mencionado, \"Liderança\".\n"
        "- Status de compromisso (\"concluído\", \"em andamento\") não é compromisso novo.\n"
        '- Não invente. Sem plano nem atualização → {"novos":[],"atualizacoes":[],"feedback":"..."}\n'
        "- feedback: sua avaliação de analista de CX sobre o que a liderança respondeu, para postar na thread "
        "(1 a 3 frases, tom de colega exigente e respeitoso, formatação Slack com *um asterisco*). Compare com o report: "
        "a ação ataca o problema cobrado (motivo de maior impacto, dispersão/Q4)? É concreta (o quê, para quem, quando)? "
        "Tem responsável, prazo e como medir o efeito? Se o plano é bom, reconheça em uma frase e diga o que vai acompanhar. "
        "Se é genérico (\"treinamento\", \"acompanhar\", \"alinhar\") ou não ataca a causa, aponte o que falta e faça UMA pergunta objetiva. "
        "Se a mensagem não é plano nem status (ex.: só \"ok\", \"vamos ver\"), peça o plano no formato ação · responsável · prazo. "
        "Nunca cite nome de operador. Vazio só se a mensagem não tiver relação com o report."
    )
    answer = ask_toqan(prompt)
    parsed = _parse_extraction(answer)
    if parsed is None:
        logger.warning("Commitment extraction was not JSON; keeping raw message as commitment")
        raw = " / ".join(m.strip() for m in messages if m.strip())[:200]
        return {"novos": [{"fila": FILA_GERAL, "acao": raw, "responsavel": "Liderança", "prazo": "a definir"}],
                "atualizacoes": [], "feedback": ""}
    return parsed


def _human_thread_message_records(messages: list[dict[str, Any]], report_ts: str) -> list[tuple[str, str]]:
    records: list[tuple[str, str]] = []
    for message in messages:
        if not isinstance(message, dict) or message.get("ts") == report_ts:
            continue
        if message.get("bot_id") or message.get("bot_profile") or message.get("subtype") == "bot_message":
            continue
        text = _strip_mentions(str(message.get("text") or ""))
        timestamp = str(message.get("ts") or "").strip()
        if text and timestamp:
            records.append((timestamp, text))
    return records


def _fetch_thread(channel_id: str, report_ts: str) -> list[dict[str, Any]]:
    messages: list[dict[str, Any]] = []
    cursor = None
    for _ in range(10):
        kwargs: dict[str, Any] = {"channel": channel_id, "ts": report_ts, "limit": 200}
        if cursor:
            kwargs["cursor"] = cursor
        response = slack_app.client.conversations_replies(**kwargs)
        page = _resp_get(response, "messages", []) or []
        messages.extend(m for m in page if isinstance(m, dict))
        cursor = (_resp_get(response, "response_metadata", {}) or {}).get("next_cursor")
        if not cursor:
            break
    return messages


_THREAD_LOCKS: dict[str, threading.Lock] = {}
_THREAD_LOCKS_GUARD = threading.Lock()


def _thread_lock(report_ts: str) -> threading.Lock:
    with _THREAD_LOCKS_GUARD:
        return _THREAD_LOCKS.setdefault(report_ts, threading.Lock())


def _read_report_thread(
    channel_id: str, report_ts: str, phase: str, today: datetime | date | None = None,
    tenant_key: str = CONSOLIDATED, fila: str = "", register_no_reply: bool = True,
) -> dict[str, int]:
    """Lê a thread (uma leitura por vez por thread: agendada e ao vivo não se atropelam)."""
    with _thread_lock(report_ts):
        return _read_report_thread_locked(channel_id, report_ts, phase, today, tenant_key, fila, register_no_reply)


def _read_report_thread_locked(
    channel_id: str, report_ts: str, phase: str, today: datetime | date | None = None,
    tenant_key: str = CONSOLIDATED, fila: str = "", register_no_reply: bool = True,
) -> dict[str, int]:
    """Lê a thread do report e registra/atualiza compromissos. Levanta exceção em falha.

    phase: "first" (2h depois; registra "sem resposta" se ninguém respondeu),
           "late" (18h) ou "next" (manhã seguinte, antes do report).
    """
    base = _as_date(today or _today_brazil())
    messages = _fetch_thread(channel_id, report_ts)
    report_text = "\n".join(
        str(m.get("text") or "")
        for m in messages
        if m.get("ts") == report_ts or m.get("bot_id") or m.get("subtype") == "bot_message"
    )
    if fila:  # bloco de uma fila só: a fila é conhecida, o CSAT vem da linha Resultado
        values = _block_csat(fila, report_text)
        csat = {fila: values} if values else {}
        filas = [fila]
    else:
        csat = _parse_report_csat(report_text)
        filas = list(csat)

    records = _human_thread_message_records(messages, report_ts)
    captured_key = f"captured:{report_ts}"
    captured = {t for t in STATE.get(captured_key).split(",") if t}
    new_records = [(t, text) for t, text in records if t not in captured]

    google_client = _get_google_client()
    if google_client is None:
        raise MOPGoogleCredentialsError("GOOGLE_CREDENTIALS_JSON não configurado")
    worksheet = _compromissos_ws(google_client)
    thread_rows = [(n, r) for n, r in _worksheet_rows_with_header(worksheet) if _row_value(r, "Thread") == report_ts]

    if not new_records:
        if phase == "first" and register_no_reply and not records and not thread_rows:
            pre_sr = _fmt_pct(csat[fila]["d1"]) if fila and fila in csat else ""
            worksheet.append_row(_sem_resposta_row(report_ts, base, tenant_key, fila=fila, pre=pre_sr,
                                                   problema_report=_report_problem(report_text)),
                                 value_input_option="RAW")
        return {"novos": 0, "atualizados": 0}

    extraction = _extract_commitments_via_toqan([text for _, text in new_records], filas, thread_rows, base,
                                                context=report_text)

    # novos
    new_rows: list[list[str]] = []
    missing_deadline: list[str] = []
    for item in extraction["novos"]:
        action = str(item.get("acao") or "").strip()
        if not action:
            continue
        item_fila = fila or _match_fila(str(item.get("fila") or ""), filas) or FILA_GERAL
        pre = _fmt_pct(csat[item_fila]["d1"]) if item_fila in csat else ""
        prazo = str(item.get("prazo") or "a definir").strip() or "a definir"
        row = _commitment_row(
            item_fila, action, str(item.get("responsavel") or "").strip(), prazo, pre, report_ts, base, tenant_key,
            problema_report=_report_problem(report_text) if fila else "",
        )
        new_rows.append(row)
        if row[COL_DATA_FIM - 1] == "a definir":
            missing_deadline.append(action)
    if new_rows:
        worksheet.append_rows(new_rows, value_input_option="RAW")

    # atualizações (só linhas desta thread)
    valid_rows = {n: r for n, r in thread_rows}
    updated = 0
    for item in extraction["atualizacoes"]:
        n = _to_int(item.get("linha"))
        if n not in valid_rows:
            continue
        changed = False
        if str(item.get("acao") or "").strip():
            worksheet.update_cell(n, COL_ACAO, str(item["acao"]).strip())
            changed = True
        if str(item.get("responsavel") or "").strip():
            worksheet.update_cell(n, COL_RESPONSAVEL, str(item["responsavel"]).strip())
            changed = True
        if str(item.get("prazo") or "").strip():
            start = _parse_commitment_date(_row_value(valid_rows[n], "Data Inicio"), base) or _midnight(base)
            data_fim, duracao = _deadline_fields(item["prazo"], start)
            if data_fim != "a definir":
                worksheet.update_cell(n, COL_DATA_FIM, data_fim)
                worksheet.update_cell(n, COL_DURACAO, duracao)
                changed = True
        updated += int(changed)

    # liderança respondeu: encerra o "sem resposta" desta thread
    if new_rows or updated:
        for n, r in thread_rows:
            if _row_value(r, "Ação Tomada").casefold() == SEM_RESPOSTA:
                worksheet.update_cell(n, COL_ACAO, f"sem resposta (respondido em {_format_date(base)})")
                worksheet.update_cell(n, COL_SUCESSO, "encerrado")

    STATE.set(captured_key, ",".join(sorted(captured | {t for t, _ in new_records})))

    feedback = str(extraction.get("feedback") or "").strip()
    answered = {t for t in STATE.get(f"answered:{report_ts}").split(",") if t}
    if feedback and all(t in answered for t, _ in new_records):
        feedback = ""  # a Lyra já respondeu essas mensagens na conversa
    if new_rows or updated or feedback:
        chunks = []
        if feedback:
            chunks.append(feedback)
        if new_rows or updated:
            parts = []
            if new_rows:
                parts.append(f"{len(new_rows)} novo(s) compromisso(s)")
            if updated:
                parts.append(f"{updated} atualizado(s)")
            chunks.append(f"Registrado: {' e '.join(parts)} ✅")
        text = "\n\n".join(chunks)
        if missing_deadline:
            text += "\nQual o prazo para:\n" + "\n".join(f"• {a[:120]}" for a in missing_deadline)
        slack_app.client.chat_postMessage(channel=channel_id, text=text, thread_ts=report_ts)
    return {"novos": len(new_rows), "atualizados": updated}


# ---------------------------------------------------------------------------
# Owners das filas, escalation e Excel de operadores
# ---------------------------------------------------------------------------
_MENTION_ID_RE = re.compile(r"<@(U[A-Z0-9]+)>")
_BULLET_FILA_RE = re.compile(r"^\s*[•\-]\s*\*(?P<fila>[^*\n]+?)\s*:?\s*\*")


_SLACK_USERS_CACHE: dict[str, Any] = {"at": None, "by_name": {}}
_UNRESOLVED_OWNERS_ALERTED: set[str] = set()


def _slack_users_by_name() -> dict[str, list[str]]:
    """{nome normalizado: [IDs]} a partir do Slack (real_name e display_name). Cache de 6h."""
    at = _SLACK_USERS_CACHE["at"]
    if at and _now_brazil() - at < timedelta(hours=6):
        return _SLACK_USERS_CACHE["by_name"]
    by_name: dict[str, list[str]] = {}
    try:
        cursor = None
        for _ in range(50):
            kwargs: dict[str, Any] = {"limit": 200}
            if cursor:
                kwargs["cursor"] = cursor
            resp = slack_app.client.users_list(**kwargs)
            for u in _resp_get(resp, "members", []) or []:
                if u.get("deleted") or u.get("is_bot"):
                    continue
                prof = u.get("profile") or {}
                for name in {u.get("real_name"), prof.get("real_name"), prof.get("display_name")}:
                    if name:
                        by_name.setdefault(_norm_text(name), [])
                        if u["id"] not in by_name[_norm_text(name)]:
                            by_name[_norm_text(name)].append(u["id"])
            cursor = (_resp_get(resp, "response_metadata", {}) or {}).get("next_cursor")
            if not cursor:
                break
    except Exception:
        logger.warning("Could not list Slack users (falta o escopo users:read?)", exc_info=True)
        return _SLACK_USERS_CACHE["by_name"]
    _SLACK_USERS_CACHE.update(at=_now_brazil(), by_name=by_name)
    return by_name


def _owner_mention(owner: str) -> str:
    """'U123' → <@U123>; '@Nome Sobrenome' → <@ID> achado no Slack (ou o nome em texto, com alerta)."""
    owner = owner.strip()
    if re.fullmatch(r"[UW][A-Z0-9]{6,}", owner):
        return f"<@{owner}>"
    name = owner.lstrip("@").strip()
    by_name = _slack_users_by_name()
    ids = by_name.get(_norm_text(name), [])
    if not ids:
        # nome no Slack costuma ser mais curto (ex.: "Carlos Porto"): aceita quando os nomes de um
        # estão todos no outro, com o mesmo primeiro nome e pelo menos 2 nomes em comum
        stop = {"da", "de", "do", "das", "dos", "e"}
        wanted = [w for w in _norm_text(name).split() if w not in stop]
        found: set[str] = set()
        for slack_name, uids in by_name.items():
            tokens = [w for w in slack_name.split() if w not in stop]
            if len(tokens) < 2 or not wanted or tokens[0] != wanted[0]:
                continue
            common = set(tokens) & set(wanted)
            if len(common) >= 2 and (set(tokens) <= set(wanted) or set(wanted) <= set(tokens)):
                found.update(uids)
        ids = sorted(found)
    if len(ids) == 1:
        return f"<@{ids[0]}>"
    if name not in _UNRESOLVED_OWNERS_ALERTED:
        _UNRESOLVED_OWNERS_ALERTED.add(name)
        motivo = "mais de uma pessoa com esse nome" if ids else "não achei no Slack"
        _alert(f"Owner '{name}': {motivo}. Coloque o ID do Slack em FILA_OWNERS. Por enquanto vai o nome em texto.")
    return f"@{name}"


def _owners_for(fila_label: str) -> list[str]:
    fila = _match_fila(fila_label, list(FILA_OWNERS))
    return FILA_OWNERS.get(fila, []) if fila else []


def _tag_owners(main: str) -> str:
    """Marca os owners nas linhas da seção Cobranças (cada linha começa com • *Fila:*)."""
    lines = main.splitlines()
    in_cobrancas = False
    for i, line in enumerate(lines):
        low = _norm_text(line)
        if "cobrancas" in low and line.strip().startswith("*"):
            in_cobrancas = True
            continue
        if in_cobrancas and line.strip().startswith("*") and not line.strip().startswith(("*•", "•")):
            in_cobrancas = False  # começou outra seção
        if not in_cobrancas:
            continue
        match = _BULLET_FILA_RE.match(line)
        if not match:
            continue
        mentions = [m for m in (_owner_mention(o) for o in _owners_for(match.group("fila"))) if m and m not in line]
        if mentions:
            lines[i] = line.rstrip() + " " + " ".join(mentions)
    return "\n".join(lines)


def _escalation_check(tenant: "Tenant", report_ts: str, today: date) -> str:
    """5h depois do report: cobra na thread quem foi marcado e não respondeu."""
    messages = _fetch_thread(tenant.channel, report_ts)
    parent = next((m for m in messages if m.get("ts") == report_ts), {})
    tagged = list(dict.fromkeys(_MENTION_ID_RE.findall(str(parent.get("text") or ""))))
    humans = [m for m in messages if m.get("ts") != report_ts and not m.get("bot_id") and m.get("subtype") != "bot_message"]
    responders = {str(m.get("user") or "") for m in humans}
    if tagged:
        pending = [uid for uid in tagged if uid not in responders]
        if not pending:
            return "todos responderam"
        who = " ".join(f"<@{uid}>" for uid in pending)
        text = f"{who} algum retorno sobre o plano? Preciso de *ação · responsável · prazo* aqui na thread pra gente acompanhar. 🙏"
    else:
        if humans:
            return "houve resposta"
        text = "Lideranças, algum retorno sobre o report? Preciso do plano aqui na thread no formato *ação · responsável · prazo*. 🙏"
    slack_app.client.chat_postMessage(channel=tenant.channel, text=text, thread_ts=report_ts)
    return "cobrado"


_OPERATORS_EXCEL_RE = re.compile(
    r"\b(excel|planilha|xlsx|arquivo|lista|base|relacao)\b.*\b(operador|operadores|analista|analistas|quadrante|quadrantes|q1|q2|q3|q4|pessoas)\b"
    r"|\b(operador|operadores|analista|analistas|quadrante|quadrantes|q4)\b.*\b(excel|planilha|xlsx|arquivo)\b"
)
OPERATORS_CSV_HEADER = ["email", "fila", "avaliacoes", "promotores", "csat", "meta", "atingimento", "quadrante"]


# Resultado individual pedido em texto ("resultado por operador de ontem", "CSAT dos analistas"): vai em Excel,
# porque nome/e-mail de operador não aparece no Slack.
_OPERATORS_RESULT_RE = re.compile(
    r"\bpor (operador|analista|atendente)\b|\bresultado (individual|dos operadores|dos analistas|dos atendentes)\b"
    r"|\b(quem|quais|lista|listar|resultado|traz|traga|mostra|manda)\b.*\b(ofensor|ofensores)\b"
    r"|\b(ofensor|ofensores)\b (de|do|da|em) (ontem|hoje|csat|d 1|mes|semana)\b"
    r"|\bpiores (operadores|analistas|atendentes)\b"
    r"|\b(csat|resultado|quadrante|ranking)\b.*\b(dos|de cada|por) (operador|operadores|analista|analistas|atendente|atendentes)\b"
)


def _is_operators_excel_request(question: str) -> bool:
    text = _norm_text(question)
    if re.search(r"\bmop\b", text):  # pedido de MOP tem prioridade
        return False
    return bool(_OPERATORS_EXCEL_RE.search(text) or _OPERATORS_RESULT_RE.search(text))


def _operators_period(question: str) -> tuple[str, str]:
    """(instrução de período para o agente, rótulo para o arquivo)."""
    text = _norm_text(question)
    day = _today_brazil() - timedelta(days=1)
    if re.search(r"\b(ontem|d 1|d1)\b", text):
        return f"D-1 ({day.strftime('%d/%m/%Y')}). Só operadores com 3+ avaliações no dia.", f"de ontem ({day.strftime('%d/%m')})"
    if re.search(r"\b(mes|mensal)\b", text):
        return "mês corrente até o D-1. Só operadores com 5+ avaliações.", "do mês"
    return "semana corrente (D-7 a D-1). Só operadores com 5+ avaliações.", "da semana"


def _eps_in_text(question: str) -> str | None:
    words = set(_norm_text(question).split())
    return next((e for e in EPS_CELULAS if _norm_text(e) in words), None)


def _operators_prompt(eps: str | None, period: str = "semana corrente (D-7 a D-1). Só operadores com 5+ avaliações.") -> str:
    return (
        f"{_scope_instruction(eps)}\n\n"
        "TAREFA INTERNA DO APP — base de operadores para Excel. Não é mensagem para o Slack.\n"
        "Se você usa subagentes, peça ao subagente de indicadores com "
        f"'EPS: {eps or 'iFood'}', 'Pedido: base de operadores (CSV)' e 'Período: {period}'\n"
        f"Período: {period} Só filas da lista oficial da EPS.\n"
        "Quadrante pelo atingimento = CSAT do operador ÷ meta da fila: Q1 ≥ 100% · Q2 95–99% · Q3 90–94% · Q4 < 90%.\n"
        "Responda APENAS um CSV separado por ';', sem markdown, com este cabeçalho exato:\n"
        "email;fila;avaliacoes;promotores;csat;meta;atingimento;quadrante\n"
        "Números com ponto decimal (csat, meta e atingimento em %, ex.: 72.5). Uma linha por operador.\n"
        "Última linha: TOTAL;<quantidade de operadores>\n"
        "Se a consulta falhar: responda só 'ERRO: <motivo>'."
    )


def _parse_operators_csv(answer: str) -> tuple[list[list[str]], int | None]:
    import csv as _csv
    text = _strip_json_code_blocks(answer)
    rows: list[list[str]] = []
    total: int | None = None
    header_seen = False
    for row in _csv.reader(io.StringIO(text), delimiter=";"):
        cells = [c.strip() for c in row]
        if not cells or not any(cells):
            continue
        if not header_seen:
            if [c.lower() for c in cells[:len(OPERATORS_CSV_HEADER)]] == OPERATORS_CSV_HEADER:
                header_seen = True
            continue
        if cells[0].upper() == "TOTAL":
            total = _to_int(cells[1] if len(cells) > 1 else "")
            continue
        if len(cells) >= len(OPERATORS_CSV_HEADER) and "@" in cells[0]:
            rows.append(cells[:len(OPERATORS_CSV_HEADER)])
    return rows, total


# --- Gráficos sob pedido (PNG montado pelo app; a Lyra só traz os dados) -------------
_CHART_RE = re.compile(r"\b(grafico|graficos|chart|plota|plotar|plote|visualizacao|visual)\b")
CHART_COLORS = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100"]  # ordem fixa (paleta validada para daltonismo)


def _is_chart_request(question: str) -> bool:
    text = _norm_text(question)
    return bool(_CHART_RE.search(text)) and not re.search(r"\b(o que|como funciona|consegue gerar)\b.*\?$", text)


def _chart_prompt(eps: str | None, question: str, context: str) -> str:
    return (
        f"{_scope_instruction(eps)}\n\n{context}"
        "TAREFA INTERNA DO APP — a pessoa pediu um gráfico. Busque os dados com o subagente certo (TMA sempre pelo SQL "
        "oficial da Lyra Indicadores) e devolva APENAS JSON válido:\n"
        '{"titulo": "TMA por célula — D-1 x mês (s)", "tipo": "barra", "unidade": "s", '
        '"categorias": ["CX Suporte", "Madrugada"], '
        '"series": [{"nome": "D-1", "valores": [526, 689]}, {"nome": "Mês", "valores": [724, 672]}], '
        '"referencia": {"nome": "meta", "valor": null}, "nota": "1 frase factual com o principal destaque"}\n'
        "Regras: tipo = \"barra\" (comparar categorias) ou \"linha\" (evolução no tempo, categorias = datas dd/mm em ordem). "
        "No máximo 4 séries e 20 categorias. unidade = \"%\", \"s\" ou \"\". Números com ponto decimal. "
        "referencia só quando houver UMA meta/plan comum a todas as categorias; senão null. Respeite o escopo da EPS. "
        "Se não der para consultar: só ERRO: <motivo>.\n\n"
        f"Pedido: {question}"
    )


def _build_chart_png(spec: dict[str, Any]) -> bytes | None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.ticker

    cats = [str(c) for c in (spec.get("categorias") or [])][:20]
    series = [s_ for s_ in (spec.get("series") or []) if isinstance(s_, dict)][:4]
    if not cats or not series:
        return None
    unit = str(spec.get("unidade") or "")
    fmt = (lambda v: f"{v:.1f}%".replace(".", ",")) if unit == "%" else \
          (lambda v: f"{_fmt_int(v)} s") if unit == "s" else (lambda v: _bfmt(v, 1))
    vals = [[(_num(v) if _num(v) is not None else float("nan")) for v in (s_.get("valores") or [])][:len(cats)] for s_ in series]
    vals = [v + [float("nan")] * (len(cats) - len(v)) for v in vals]
    ink, muted, grid = "#2b2b29", "#6b6a64", "#e6e5e0"
    plt.rcParams.update({"font.size": 10, "axes.edgecolor": grid, "axes.labelcolor": muted,
                         "xtick.color": muted, "ytick.color": muted, "text.color": ink})
    kind = str(spec.get("tipo") or "barra").lower()
    horizontal = kind != "linha"
    height = max(3.2, 0.42 * len(cats) * (1 + 0.35 * (len(series) - 1)) + 1.2) if horizontal else 4.2
    fig, ax = plt.subplots(figsize=(9, height), dpi=150)
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    for side in ("top", "right"):
        ax.spines[side].set_visible(False)
    if horizontal:
        n = len(series)
        bar_h = 0.8 / n
        order = list(range(len(cats)))[::-1]  # primeira categoria no topo
        for i, (s_, v) in enumerate(zip(series, vals)):
            ys = [o + (n - 1 - i) * bar_h - 0.4 + bar_h / 2 for o in order]
            ax.barh(ys, v, height=bar_h * 0.9, color=CHART_COLORS[i], label=str(s_.get("nome") or ""), zorder=2)
            for y, x in zip(ys, v):  # rótulo de dados em todas as barras
                if x == x:
                    ax.text(x, y, " " + fmt(x), va="center", ha="left", fontsize=7.5 if n > 1 else 8.5, color=ink)
        ax.set_yticks(order)
        ax.set_yticklabels(cats)
        ax.xaxis.grid(True, color=grid, linewidth=0.8, zorder=0)
        ax.spines["left"].set_visible(False)
        ax.tick_params(axis="y", length=0)
        ref = spec.get("referencia") or {}
        if isinstance(ref, dict) and _num(ref.get("valor")) is not None:
            ax.axvline(_num(ref["valor"]), color=muted, linestyle=(0, (4, 3)), linewidth=1, zorder=3)
            ax.text(_num(ref["valor"]), len(cats) - 0.45, f" {ref.get('nome') or 'meta'} {fmt(_num(ref['valor']))}",
                    fontsize=8, color=muted, va="bottom")
        right = max([x for v in vals for x in v if x == x] or [1])
        ax.set_xlim(0, right * 1.15)
    else:
        for i, (s_, v) in enumerate(zip(series, vals)):
            ax.plot(cats, v, color=CHART_COLORS[i], linewidth=2, marker="o", markersize=5,
                    label=str(s_.get("nome") or ""), zorder=3)
            for j, y in enumerate(v):  # rótulo de dados em todos os pontos: acima nos picos, abaixo nos vales
                if y != y:
                    continue
                neigh = [v[k] for k in (j - 1, j + 1) if 0 <= k < len(v) and v[k] == v[k]]
                below = bool(neigh) and y < sum(neigh) / len(neigh)
                if len(series) > 1:
                    below = i % 2 == 1
                ax.annotate(fmt(y), (j, y), textcoords="offset points", xytext=(0, -13 if below else 8), ha="center",
                            fontsize=7 if len(cats) > 12 else 8, color=ink,
                            bbox=dict(boxstyle="round,pad=0.15", facecolor="white", edgecolor="none", alpha=0.85))
        ax.yaxis.grid(True, color=grid, linewidth=0.8, zorder=0)
        ref = spec.get("referencia") or {}
        if isinstance(ref, dict) and _num(ref.get("valor")) is not None:
            ax.axhline(_num(ref["valor"]), color=muted, linestyle=(0, (4, 3)), linewidth=1)
            ax.annotate(f"{ref.get('nome') or 'meta'} {fmt(_num(ref['valor']))}", (1, _num(ref["valor"])),
                        xycoords=("axes fraction", "data"), xytext=(4, 0), textcoords="offset points",
                        fontsize=8, color=muted, va="center", ha="left")
        if len(cats) > 10:
            plt.setp(ax.get_xticklabels(), rotation=45, ha="right")
        ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: _bfmt(v, 1)))
    if len(series) > 1:
        ax.legend(frameon=False, loc="lower right" if horizontal else "best", fontsize=9)
    ax.set_title(str(spec.get("titulo") or ""), loc="left", fontsize=12, color=ink, pad=12)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor="white")
    plt.close(fig)
    return buf.getvalue()


def _send_chart(event: dict[str, Any], client: Any, eps: str | None, question: str,
                send_progress: Callable[[str], Any] | None = None) -> str | None:
    answer = ask_toqan(_chart_prompt(eps, question, _thread_context(event)), progress_callback=send_progress)
    if _is_error_answer(answer):
        return "Não consegui puxar esses dados agora. Tenta de novo em alguns minutos?"
    spec = _block_data_from_json(answer) or {}
    png = _build_chart_png(spec)
    if not png:
        return "Não consegui montar o gráfico com esse pedido. Me diz o indicador, o período e o recorte (ex.: TMA por célula, D-1 x mês)."
    nota = str(spec.get("nota") or "").strip()
    title = str(spec.get("titulo") or "Gráfico")
    (client or slack_app.client).files_upload_v2(
        channel=event.get("channel"), thread_ts=_reply_ts(event),
        content=png, filename=re.sub(r"[^\w\-]+", "_", title)[:60] + ".png", title=title,
        initial_comment=(f"📊 {nota}" if nota else "📊 Segue o gráfico."),
    )
    return None


_DATA_EXCEL_RE = re.compile(r"\b(excel|xls|xlsx|planilha|csv|tabela em arquivo|arquivo)\b")


def _is_data_excel_request(question: str) -> bool:
    """Pedido de dados em planilha (ex.: "manda o CSAT por fila em excel"). MOP e operadores têm fluxo próprio."""
    text = _norm_text(question)
    asks = re.search(r"\b(manda|mande|envia|envie|gera|gere|gerar|exporta|exporte|exportar|extrai|extraia|"
                     r"quero|queria|preciso|passa|passe|baixar|tira|tire|monta|monte|montar|coloca|coloque|colocar|pode|consegue|em|numa|num)\b", text)
    return bool(_DATA_EXCEL_RE.search(text) and asks) and not re.search(r"\bmop\b", text)


def _data_excel_prompt(eps: str | None, question: str, context: str) -> str:
    return (
        f"{_scope_instruction(eps)}\n\n{context}"
        "TAREFA INTERNA DO APP — a pessoa pediu os dados em planilha. Busque os dados com o subagente certo "
        "(Lyra Indicadores para números; Lyra Qualidade para causa raiz/comentários) e responda APENAS:\n"
        "ARQUIVO: <nome curto do arquivo, sem extensão>\n"
        "<CSV separado por ';', com cabeçalho em linguagem de negócio (sem nome de coluna técnica), uma linha por registro, "
        "números com ponto decimal e sem símbolo de %>\n"
        "Nada além disso (sem markdown, sem comentário). Respeite o escopo da EPS. Nunca dado pessoal de cliente. "
        "E-mail de operador só se a pessoa pedir lista de operadores, e só da EPS do canal. "
        "Se não der para consultar, responda só ERRO: <motivo>.\n\n"
        f"Pedido: {question}"
    )


def _build_data_xlsx(header: list[str], rows: list[list[str]], title: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = (title or "Dados")[:31]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="EA1D2C")
    for r in rows:
        ws.append([(_to_float(v) if _to_float(v) is not None and re.fullmatch(r"-?[\d.,]+", v.strip()) else v) for v in r])
    for i, name in enumerate(header, start=1):
        width = max([len(str(name))] + [len(str(r[i - 1])) for r in rows if len(r) >= i]) + 2
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(width, 10), 50)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_data_csv(answer: str) -> tuple[str, list[str], list[list[str]]] | None:
    text = _strip_json_code_blocks(answer or "").strip().strip("`")
    name = "Dados"
    lines = [l for l in text.splitlines() if l.strip()]
    if lines and lines[0].upper().startswith("ARQUIVO:"):
        name = lines.pop(0).split(":", 1)[1].strip() or name
    lines = [l for l in lines if ";" in l]
    if len(lines) < 2:
        return None
    header = [h.strip() for h in lines[0].split(";")]
    rows = [[c.strip() for c in l.split(";")] for l in lines[1:]]
    rows = [r + [""] * (len(header) - len(r)) for r in rows]
    return name, header, rows


def _send_data_excel(event: dict[str, Any], client: Any, eps: str | None, question: str) -> str | None:
    answer = ask_toqan(_data_excel_prompt(eps, question, _thread_context(event)))
    if _is_error_answer(answer):
        return "Não consegui puxar esses dados agora. Tenta de novo em alguns minutos?"
    parsed = _parse_data_csv(answer)
    if not parsed:
        return "Não consegui montar a planilha com esse pedido. Me diz quais colunas você quer (ex.: fila, CSAT, meta, TMA)."
    name, header, rows = parsed
    if eps:  # trava: e-mail de operador só da EPS do canal
        rows = [r for r in rows if not any("@" in c and _eps_pattern(eps) not in c.lower() for c in r)]
    safe = re.sub(r"[^\w\- ]+", "", name).strip().replace(" ", "_") or "Dados"
    stamp = _now_brazil().strftime("%d-%m-%Y")
    (client or slack_app.client).files_upload_v2(
        channel=event.get("channel"),
        thread_ts=_reply_ts(event),
        content=_build_data_xlsx(header, rows, name),
        filename=f"{safe}{('_' + eps) if eps else ''}_{stamp}.xlsx",
        title=name,
        initial_comment=f"Segue a planilha 📋 ({len(rows)} linha{'s' if len(rows) != 1 else ''})",
    )
    return None


def _build_operators_xlsx(rows: list[list[str]], eps: str | None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    def num(value: str) -> float | str:
        f = _to_float(value)
        return f if f is not None else value

    wb = Workbook()
    ws = wb.active
    ws.title = f"Operadores {eps}" if eps else "Operadores"
    header = ["E-mail", "Fila", "Avaliações", "Promotores", "CSAT (%)", "Meta (%)", "Atingimento (%)", "Quadrante"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="EA1D2C")
    fills = {"Q1": "C6EFCE", "Q2": "E2F0D9", "Q3": "FFF2CC", "Q4": "F8CBAD"}
    order = {"Q4": 0, "Q3": 1, "Q2": 2, "Q1": 3}
    rows = sorted(rows, key=lambda r: (order.get(r[7].upper(), 9), _to_float(r[6]) or 0))
    for r in rows:
        ws.append([r[0], r[1], num(r[2]), num(r[3]), num(r[4]), num(r[5]), num(r[6]), r[7].upper()])
        q = r[7].upper()
        if q in fills:
            ws.cell(row=ws.max_row, column=8).fill = PatternFill("solid", fgColor=fills[q])
    for col, width in zip("ABCDEFGH", (38, 26, 11, 11, 10, 10, 15, 11)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _operators_rows(eps: str | None, question: str) -> tuple[list[list[str]], int | None, str, str | None] | str:
    """Base de operadores do período pedido: (linhas, total, rótulo do período, EPS usada) ou texto de erro."""
    period, period_label = _operators_period(question)
    if eps is None:
        eps = _eps_in_text(question)  # visão completa pedindo uma EPS: recorta nela
    answer = ask_toqan(_operators_prompt(eps, period))
    if _is_error_answer(answer):
        return "Não consegui puxar os dados dos operadores agora. Tenta de novo em alguns minutos?"
    rows, total = _parse_operators_csv(answer)
    if not rows:
        return "Não consegui montar a base dos operadores agora. Tenta de novo em alguns minutos?"
    if eps:
        rows = [r for r in rows if _eps_pattern(eps) in r[0].lower()]  # trava: só e-mails da EPS
    fila = _fila_in_question(question, eps)
    if fila:
        rows = [r for r in rows if _norm_text(r[1]) == _norm_text(fila)]
    return rows, total, period_label, eps


def _upload_operators_xlsx(event: dict[str, Any], client: Any, rows: list[list[str]], eps: str | None,
                           period_label: str, comment: str) -> None:
    stamp = _now_brazil().strftime("%d-%m-%Y")
    who = f" da {eps}" if eps else ""
    (client or slack_app.client).files_upload_v2(
        channel=event.get("channel"),
        thread_ts=_reply_ts(event),
        content=_build_operators_xlsx(rows, eps),
        filename=f"Operadores{('_' + eps) if eps else ''}_{stamp}.xlsx",
        title=f"Operadores{who} — CSAT e quadrante {period_label}",
        initial_comment=comment,
    )


def _send_operators_excel(event: dict[str, Any], client: Any, eps: str | None, question: str = "") -> str | None:
    """Pede a base ao agente, monta o Excel e posta na thread. Devolve texto se não deu."""
    got = _operators_rows(eps, question)
    if isinstance(got, str):
        return got
    rows, total, period_label, eps = got
    if not rows:
        return "Não achei operadores com avaliações suficientes nesse recorte."
    who = f" da {eps}" if eps else ""
    note = f" (vieram {len(rows)} de {total} — a lista pode estar incompleta)" if total and total > len(rows) else ""
    _upload_operators_xlsx(event, client, rows, eps, period_label,
                           f"Segue a planilha dos operadores{who} com CSAT e quadrante {period_label} 📋{note}")
    return None


OPERATORS_INLINE_MAX = 15
_OFFENDERS_RE = re.compile(r"\b(ofensor|ofensores|piores|q4|detrator|detratores)\b")


def _send_operators_inline(event: dict[str, Any], client: Any, eps: str | None, question: str,
                           say: Callable[..., Any], in_thread: bool) -> str | None:
    """Visão completa (e-mail iFood / canal interno): resultado por operador direto no chat. "Ofensores" = Q4,
    do pior para o melhor. Mais de 15 linhas: mostra as 15 piores e manda a planilha completa junto."""
    got = _operators_rows(eps, question)
    if isinstance(got, str):
        return got
    rows, _, period_label, eps = got
    fila = _fila_in_question(question, eps)
    offenders = bool(_OFFENDERS_RE.search(_norm_text(question)))
    data = sorted(rows, key=lambda r: (_to_float(r[6]) if _to_float(r[6]) is not None else 999,
                                       -(_to_float(r[2]) or 0)))
    if offenders:
        data = [r for r in data if str(r[7]).strip().upper() == "Q4"]
    where = " · ".join(x for x in (fila, eps if not fila else None) if x) or "todas as filas"
    when = re.sub(r"^(de|do|da) ", "", period_label)
    title = f"*{'Ofensores de CSAT' if offenders else 'CSAT por operador'} · {where} · {when}*"
    if not data:
        return f"{title}\nNinguém no Q4 nesse recorte (abaixo de 90% da meta, com avaliações suficientes). 👏"
    lines = []
    for i, r in enumerate(data[:OPERATORS_INLINE_MAX], 1):
        user = r[0].split("@")[0]
        csat, ating = _to_float(r[4]), _to_float(r[6])
        fila_txt = f" · {r[1]}" if not fila else ""
        lines.append(f"{i}. `{user}`{fila_txt} — CSAT {_bfmt(csat) if csat is not None else r[4]}% "
                     f"({r[3]}/{r[2]} promotoras) · {_bfmt(ating, 0) if ating is not None else r[6]}% da meta · {r[7]}")
    rule = "Q4 = abaixo de 90% da meta da fila" if offenders else "do menor para o maior atingimento"
    text = f"{title}\n_{rule}_\n\n" + "\n".join(lines)
    if len(data) > OPERATORS_INLINE_MAX:
        text += f"\n\n…e mais {len(data) - OPERATORS_INLINE_MAX}. A lista completa vai na planilha aqui embaixo 📋"
    _say(say, text, event, in_thread)
    if len(data) > OPERATORS_INLINE_MAX:
        _upload_operators_xlsx(event, client, data, eps, period_label, "Lista completa 📋")
    return None


# ---------------------------------------------------------------------------
# Report diário (um por EPS + consolidado opcional)
# ---------------------------------------------------------------------------
class ReportGenerationError(RuntimeError):
    """O Toqan não devolveu um report publicável."""


_FILA_HEADER_RE = re.compile(
    r"^\s*(?:[🔴🟢⚪🟡]\s*)?\*{0,2}\s*(?:Governan[çc]a\s*·\s*)?\*{0,2}\s*`?(?P<fila>[^`*\n·—]+?)`?\s*\*{0,2}\s*[·—–-]\s*\d{1,2}/\d{1,2}\b.*$",
    re.MULTILINE,
)


def _split_blocks(answer: str, use_headers: bool = False) -> list[tuple[str, str]]:
    """[(fila, texto do bloco)] a partir dos marcadores ===FILA: x===. Vazio se não houver marcador."""
    marks = list(FILA_BLOCK_RE.finditer(answer or ""))
    if not marks and use_headers:  # report diário sem marcadores: usa os títulos de fila que o modelo escreveu
        marks = list(_FILA_HEADER_RE.finditer(answer or ""))
    blocks = []
    for i, m in enumerate(marks):
        end = marks[i + 1].start() if i + 1 < len(marks) else len(answer)
        text = answer[m.end():end].strip()
        if text:
            blocks.append((m.group("fila").strip(), text))
    return blocks


def _tag_block(fila: str, text: str) -> tuple[str, bool]:
    """Marca o owner na linha *Cobrança:* (ou 👉) do bloco. Devolve (texto, tem_cobrança)."""
    lines = text.splitlines()
    has_cobranca = False
    for i, line in enumerate(lines):
        stripped = line.strip()
        if stripped.startswith("👉") or _norm_text(stripped.lstrip("•").strip().strip("*")).startswith("cobranca"):
            has_cobranca = True
            mentions = [m for m in (_owner_mention(o) for o in _owners_for(fila)) if m and m not in line]
            if mentions:
                lines[i] = line.rstrip() + " " + " ".join(mentions)
    return "\n".join(lines), has_cobranca


_BLOCK_D1_RE = re.compile(rf"\*?\s*({_NUM})\s*%\s*\*?\s*em\s+D\s*[-–]?\s*1", re.IGNORECASE)
_BLOCK_MES_RE = re.compile(rf"\bm[eê]s\s*:?\s*\*?\s*({_NUM})\s*%", re.IGNORECASE)

# --- layout do bloco: montado pelo app (não depende do modelo seguir o formato) ---
_SNUM = r"[-−–]?\s*\d{1,4}(?:[.,]\d+)?"
_X_D1_AFTER_RE = re.compile(rf"D\s*[-–]?\s*1\s*:?\s*\*?\s*({_NUM})\s*%", re.IGNORECASE)
_X_N_RE = re.compile(r"\bn\s*=\s*(\d[\d.]*)", re.IGNORECASE)
_X_VAR_RE = re.compile(rf"([▲▼])\s*({_NUM})")
_X_META_RE = re.compile(rf"\bmeta\s*:?\s*\*?\s*({_NUM})\s*%", re.IGNORECASE)
_X_OFENSOR_RE = re.compile(r"maior\s+(?:ofensor|impacto)\s*:?\s*\*?\s*:?\s*(.+)", re.IGNORECASE)
_X_PP_RE = re.compile(rf"({_SNUM})\s*p\.?\s*p", re.IGNORECASE)
_X_Q4_RE = re.compile(r"(?:\bQ4\s*:?\s*(\d+))|(?:(\d+)\s*(?:operadores\s*)?no\s*Q4)", re.IGNORECASE)


def _num(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("−", "-").replace("–", "-").replace(" ", "").replace("%", "")
    if text.count(",") == 1 and "." in text:  # 1.240,5
        text = text.replace(".", "")
    try:
        return float(text.replace(",", "."))
    except ValueError:
        return None


def _bfmt(value: float, decimals: int = 1) -> str:
    return f"{value:.{decimals}f}".replace(".", ",")


def _fmt_meta(value: float) -> str:
    return _bfmt(value, 0) if float(value).is_integer() else _bfmt(value)


def _fmt_int(value: float) -> str:
    return f"{int(round(value)):,}".replace(",", ".")


def _block_data_from_json(text: str) -> dict[str, Any] | None:
    start, end = text.find("{"), text.rfind("}")
    if start < 0 or end <= start:
        return None
    try:
        data = json.loads(text[start:end + 1])
    except json.JSONDecodeError:
        return None
    return data if isinstance(data, dict) else None


def _block_data_from_text(text: str) -> dict[str, Any]:
    """Extrai os números de um bloco em texto livre (quando o modelo não devolve o JSON)."""
    clean = text.replace("**", "*")
    data: dict[str, Any] = {}
    if re.search(r"sem\s+volume", clean, re.IGNORECASE) and not _X_N_RE.search(clean):
        return {"sem_volume": True}
    d1 = _BLOCK_D1_RE.search(clean) or _X_D1_AFTER_RE.search(clean)
    data["d1"] = _num(d1.group(1)) if d1 else None
    for key, rx in (("n", _X_N_RE), ("meta", _X_META_RE), ("mes", _BLOCK_MES_RE)):
        m = rx.search(clean)
        data[key] = _num(m.group(1).replace(".", "") if key == "n" else m.group(1)) if m else None
    var = _X_VAR_RE.search(clean)
    if var:
        data["var"] = (_num(var.group(2)) or 0) * (-1 if var.group(1) == "▼" else 1)
    for line in clean.splitlines():
        m = _X_OFENSOR_RE.search(line)
        if m:
            rest = m.group(1).strip().strip("*")
            pp = _X_PP_RE.search(rest)
            if pp:
                value = _num(pp.group(1)) or 0
                data["ofensor_pp"] = -abs(value)
            name = re.split(r"\s*(?:·|\(|—| - |,\s*tira|\btira\b|\|)", rest, maxsplit=1)[0]
            data["ofensor"] = re.sub(r"^D\s*-\s*1\s*:?\s*", "", name.strip(" *:."), flags=re.IGNORECASE)
            mes_part = re.search(r"\|\s*m[eê]s\s*:?\s*(.+)", rest, re.IGNORECASE)
            if mes_part:
                data["ofensor_mes"] = re.split(r"\s*(?:·|\(|—|\btira\b)", mes_part.group(1), maxsplit=1)[0].strip(" *:.")
                pp_mes = _X_PP_RE.search(mes_part.group(1))
                if pp_mes:
                    data["ofensor_mes_pp"] = -abs(_num(pp_mes.group(1)) or 0)
            break
    q4 = _X_Q4_RE.search(clean)
    if q4:
        data["q4"] = _num(q4.group(1) or q4.group(2))
    for q in ("q1", "q2", "q3"):
        m = re.search(rf"\b{q}\s*:?\s*(\d+)", clean, re.IGNORECASE)
        if m:
            data[q] = _num(m.group(1))
    reinc = re.search(r"(\d+)\s*reincident", clean, re.IGNORECASE)
    if reinc:
        data["reincidentes"] = _num(reinc.group(1))
    cob = re.search(r"cobran[çc]a\s*:?\s*\*?\s*:?\s*(.+?)(?:\s*(?:Plano|Status|O prazo da ação) |$)", clean, re.IGNORECASE)
    if cob:
        data["cobranca"] = cob.group(1).strip(" *")
    pin = re.search(r"📌\s*(.+)", clean)
    if pin:
        data["compromisso"] = pin.group(1).split("—")[0].strip(" *")
    for line in clean.splitlines():
        if line.strip().startswith("👉") or "cobran" in _norm_text(line)[:12]:
            questions = re.findall(r"[^.?!—:]*\?", line)
            if questions:
                data["pergunta"] = questions[0].strip(" *")
            break
    return data


def _block_status(data: dict[str, Any]) -> str:
    if data.get("sem_volume") or data.get("d1") is None:
        return "sem_volume"
    if data.get("meta") is None:
        return "sem_meta"
    return "abaixo" if data["d1"] < data["meta"] else "ok"


QUAD_KEYS = ("q1", "q2", "q3", "q4")
DISPERSION_ALERT_PP = 5.0  # variação de Q3+Q4 (p.p. dos operadores) que conta como mudança de dispersão


def _quad(data: dict[str, Any], key: str) -> dict[str, float] | None:
    """Quantidade de operadores por quadrante ({q1..q4}) no mês ou no D-1."""
    raw = data.get(key)
    if key == "quad_mes" and not isinstance(raw, dict) and any(data.get(q) is not None for q in QUAD_KEYS):
        raw = {q: data.get(q) for q in QUAD_KEYS}  # formato antigo: q1..q4 soltos
    if not isinstance(raw, dict):
        return None
    values = {q: _num(raw.get(q)) or 0.0 for q in QUAD_KEYS}
    return values if sum(values.values()) > 0 else None


def _quad_shares(quad: dict[str, float]) -> dict[str, float]:
    total = sum(quad.values())
    return {q: 100.0 * v / total for q, v in quad.items()}


def _quad_line(label: str, quad: dict[str, float]) -> str:
    shares = _quad_shares(quad)
    parts = " · ".join(f"{q.upper()} {round(shares[q]):d}%" for q in QUAD_KEYS)
    return f"`{label} │ {parts}`"


DISPERSION_ALERT = 0.05  # variação do índice (CSAT Q4 ÷ CSAT Q1) que conta como mudança


def _disp_index(data: dict[str, Any], key: str) -> float | None:
    """Dispersão (régua do Cockpit) = CSAT dos operadores do Q4 ÷ CSAT dos operadores do Q1. Perto de 1 = homogêneo."""
    raw = data.get(key)
    if not isinstance(raw, dict):
        return None
    q1, q4 = _num(raw.get("q1")), _num(raw.get("q4"))
    if not q1 or q4 is None:
        return None
    return q4 / q1


def _dispersion(data: dict[str, Any]) -> tuple[str, float, float] | None:
    """('aumentou' | 'diminuiu' | 'estável', índice mês, índice D-1). Índice caindo = dispersão aumentando."""
    mes, d1 = _disp_index(data, "disp_mes"), _disp_index(data, "disp_d1")
    if mes is None or d1 is None:
        return None
    delta = d1 - mes
    trend = "aumentou" if delta <= -DISPERSION_ALERT else "diminuiu" if delta >= DISPERSION_ALERT else "estável"
    return trend, mes, d1


def _same_reason(a: Any, b: Any) -> bool:
    na, nb = _norm_text(a), _norm_text(b)
    return bool(na) and bool(nb) and (na == nb or na in nb or nb in na)


def _fmt_tempo(seconds: float) -> str:
    """Tempos sempre em segundos (preferência do time), ex.: 802 s."""
    return f"{_fmt_int(seconds)} s"


def _tma_data(fila: str, data: dict[str, Any]) -> dict[str, Any] | None:
    """TMA = tempo por interação × interações por ticket, no D-1 e no mês, contra o plan da célula."""
    if not TMA_IN_REPORT or any(_norm_text(fila) == _norm_text(f) for f in NO_TMA_FILAS):
        return None
    tpi_d1, int_d1 = _num(data.get("tma_tpi_d1")), _num(data.get("tma_int_d1"))
    if tpi_d1 is None or int_d1 is None:
        return None
    tpi_mes, int_mes = _num(data.get("tma_tpi_mes")), _num(data.get("tma_int_mes"))
    plan_key = next((k for k in TMA_PLAN if _norm_text(k) == _norm_text(fila)), None)  # nome exato: "Review" ≠ "CX Review - AeC"
    plan = TMA_PLAN.get(plan_key) if plan_key else None
    out: dict[str, Any] = {
        "d1": tpi_d1 * int_d1, "tpi_d1": tpi_d1, "int_d1": int_d1,
        "mes": (tpi_mes * int_mes) if tpi_mes is not None and int_mes is not None else None,
        "tratados": _num(data.get("tma_tratados_d1")),
        "plan": None,
    }
    if plan:
        tempo_unico, inter_plan = plan
        out.update(plan=tempo_unico, int_plan=inter_plan, tpi_plan=tempo_unico / inter_plan if inter_plan else None)
    return out


def _tma_status(tma: dict[str, Any] | None) -> str:
    if not tma:
        return "sem_dado"
    if tma.get("plan") is None:
        return "sem_plan"
    return "acima" if tma["d1"] > tma["plan"] else "ok"


def _render_block(fila: str, data: dict[str, Any], day: date) -> tuple[str, str] | None:
    """(título no canal, corpo na thread). None = fila sem CSAT no D-1 (não entra no report).
    Corpo em duas seções bem separadas: D-1 (o que é cobrado) e Mês (contexto), depois leitura e cobrança."""
    status = _block_status(data)
    if status == "sem_volume":
        return None
    dd = day.strftime("%d/%m")
    d1, meta, mes = data["d1"], data.get("meta"), data.get("mes")
    tma = _tma_data(fila, data)
    tstatus = _tma_status(tma)
    off = status == "abaixo" or tstatus == "acima"
    has_target = meta is not None or tstatus in {"ok", "acima"}
    emoji = "🔴" if off else ("🟢" if has_target else "⚪")
    title = f"{emoji} *Governança ·* `{fila}` *— D-1 {dd}* · CSAT D-1 {_bfmt(d1)}%" + (
        f" (meta {_fmt_meta(meta)}%)" if meta is not None else "")
    if tma:
        title += f" · TMA {_fmt_tempo(tma['d1'])}" + (f" (plan {_fmt_tempo(tma['plan'])})" if tma.get("plan") else "")

    def csat_vs_meta(value: float) -> str:
        if meta is None:
            return "⚪ sem meta"
        if value < meta:
            return f"🔴 {_bfmt(meta - value)} p.p. abaixo da meta {_fmt_meta(meta)}%"
        return "🟢 na meta" if value == meta else f"🟢 {_bfmt(value - meta)} p.p. acima da meta {_fmt_meta(meta)}%"

    def tma_vs_plan(value: float) -> str:
        if not tma or tma.get("plan") is None:
            return "⚪ plan a definir"
        diff = value - tma["plan"]
        if diff > 0:
            return f"🔴 {_fmt_tempo(diff)} acima do plan {_fmt_tempo(tma['plan'])}"
        return f"🟢 {_fmt_tempo(-diff)} abaixo do plan {_fmt_tempo(tma['plan'])}" if diff < -0.5 else "🟢 no plan"

    def problema(name_key: str, pp_key: str) -> str:
        text = str(data.get(name_key) or "").strip()
        if text and data.get(pp_key) is not None:
            text += f" · tira {_bfmt(abs(_num(data[pp_key]) or 0))} p.p."
        return text

    def quad_text(quad: dict[str, float] | None, disp_key: str) -> str:
        if not quad:
            return ""
        shares = _quad_shares(quad)
        text = " · ".join(f"{q.upper()} {round(shares[q]):d}%" for q in QUAD_KEYS)
        idx = _disp_index(data, disp_key)
        return text + (f" · dispersão {_bfmt(idx, 2)}" if idx is not None else "")

    lines: list[str] = []
    # ---- D-1: o que está sendo cobrado
    lines.append(f"📍 *D-1 · {dd}* — é o que estamos cobrando")
    extra = []
    if data.get("n") is not None:
        extra.append(f"n={_fmt_int(data['n'])}")
    if data.get("var") is not None:
        extra.append(f"{'▼' if data['var'] < 0 else '▲'} {_bfmt(abs(data['var']))} p.p. vs dia anterior")
    lines.append(f"• CSAT do D-1 *{_bfmt(d1)}%* → {csat_vs_meta(d1)}" + (f" ({' · '.join(extra)})" if extra else ""))
    if tma:
        lines.append(f"• TMA do D-1 *{_fmt_tempo(tma['d1'])}* → {tma_vs_plan(tma['d1'])}"
                     + (f" ({_fmt_int(tma['tratados'])} tratados)" if tma.get("tratados") is not None else ""))
    p_d1 = problema("ofensor", "ofensor_pp")
    if p_d1:
        lines.append(f"• Problema do D-1: {p_d1}")
    quad_mes, quad_d1 = _quad(data, "quad_mes"), _quad(data, "quad_d1")
    q_d1 = quad_text(quad_d1, "disp_d1")
    if q_d1:
        lines.append(f"• Quadrantes do D-1: {q_d1}")

    # ---- Mês: contexto
    month_lines = []
    if mes is not None:
        month_lines.append(f"• CSAT do mês *{_bfmt(mes)}%* → {csat_vs_meta(mes)}")
    if tma and tma.get("mes") is not None:
        month_lines.append(f"• TMA do mês *{_fmt_tempo(tma['mes'])}* → {tma_vs_plan(tma['mes'])}")
    p_mes = problema("ofensor_mes", "ofensor_mes_pp")
    if p_mes:
        month_lines.append(f"• Problema do mês: {p_mes}")
    q_mes = quad_text(quad_mes, "disp_mes")
    if q_mes:
        month_lines.append(f"• Quadrantes do mês: {q_mes}")
    if month_lines:
        lines.append("")
        lines.append(f"🗓️ *Mês · 01 a {dd}* — contexto")
        lines.extend(month_lines)

    # ---- Leitura (o que os números dizem)
    leitura = []
    if p_d1 and p_mes:
        leitura.append("Mesmo problema no D-1 e no mês: é estrutural, não pontual."
                       if _same_reason(data.get("ofensor"), data.get("ofensor_mes"))
                       else "O problema do D-1 é diferente do mês: vale entender o que mudou ontem.")
    disp = _dispersion(data)
    reinc = data.get("reincidentes")
    reinc_txt = ""
    if reinc is not None:
        n_re = int(_num(reinc) or 0)
        reinc_txt = f" {n_re} reincidente{'s' if n_re != 1 else ''} no Q4."
    if disp:
        trend, idx_mes, idx_d1 = disp
        change = f"{_bfmt(idx_mes, 2)} no mês → {_bfmt(idx_d1, 2)} no D-1"
        leitura.append({"aumentou": f"📈 Dispersão aumentou ({change}): piora no comportamento operacional.",
                        "diminuiu": f"📉 Dispersão diminuiu ({change}): operação mais homogênea.",
                        "estável": f"Dispersão estável ({change})."}[trend] + reinc_txt)
    elif reinc_txt:
        leitura.append(reinc_txt.strip())
    if tstatus == "acima" and tma.get("tpi_plan") and tma.get("int_plan"):
        r_int = tma["int_d1"] / tma["int_plan"]
        r_tpi = tma["tpi_d1"] / tma["tpi_plan"]
        if r_int >= r_tpi:
            leitura.append(f"TMA puxado por *mais interações por ticket* ({_bfmt(tma['int_d1'], 2)} vs "
                           f"{_bfmt(tma['int_plan'], 2)} do plan): cliente voltando = não resolveu no primeiro contato.")
        else:
            leitura.append(f"TMA puxado por *tempo por interação* ({_fmt_tempo(tma['tpi_d1'])} vs "
                           f"{_fmt_tempo(tma['tpi_plan'])} do plan): atendimento mais longo = condução, procedimento ou ferramenta.")
    if leitura:
        lines.append("")
        lines.append("💡 *Leitura*")
        lines.extend(f"↳ {t}" for t in leitura)

    # ---- Compromisso e cobrança / destaque
    # Ação já rodando para o mesmo problema: cobra status e efeito, não ação nova (prazo vencido: resultado ou novo plano).
    compromisso = str(data.get("compromisso") or "").strip()
    cobre = bool(compromisso) and data.get("compromisso_cobre") is not False
    vencido = False
    pm = re.search(r"prazo\s*(\d{1,2})/(\d{1,2})", compromisso)
    if pm:
        try:
            prazo = date(day.year, int(pm.group(2)), int(pm.group(1)))
            if prazo.month == 12 and day.month == 1:
                prazo = prazo.replace(year=day.year - 1)
            vencido = prazo <= day  # venceu até o D-1
        except ValueError:
            pass
    if compromisso or off or has_target:
        lines.append("")
    if compromisso:
        label = "📌 *Ação com prazo vencido:*" if vencido else "📌 *Ação em andamento:*"
        lines.append(f"{label} {compromisso}" + ("" if off and cobre else " — qual o status?"))
    if off:
        cobranca = str(data.get("cobranca") or data.get("pergunta") or "").strip()
        cobranca = re.split(r"\s*(Plano|Status) aqui na thread", cobranca, maxsplit=1)[0].strip()
        if not cobranca:
            gaps = []
            if status == "abaixo":
                gaps.append(f"CSAT {_bfmt(meta - d1)} p.p. abaixo da meta")
            if tstatus == "acima":
                gaps.append(f"TMA {_fmt_tempo(tma['d1'] - tma['plan'])} acima do plan")
            cobranca = (f"{' e '.join(gaps)} no D-1 — a ação em andamento está segurando o problema?" if cobre
                        else f"{' e '.join(gaps)} no D-1 — o que está sendo feito pra corrigir?")
        if not cobranca.endswith("?"):
            cobranca = cobranca.rstrip(".") + "?"
        if cobre and not vencido:
            suffix = "Status aqui na thread até 14h: *o que já foi feito · efeito até agora · próximo passo*."
        elif cobre:
            suffix = ("O prazo da ação venceu: foi feita e qual o resultado? Se não resolveu, novo plano aqui na thread "
                      "até 14h: *ação · responsável · prazo*.")
        else:
            suffix = "Plano aqui na thread até 14h: *ação · responsável · prazo*."
        lines.append(f"👉 *Cobrança:* {cobranca} {suffix}")
    elif has_target:
        elogio = str(data.get("elogio") or "").strip() or ("CSAT e TMA dentro do esperado, bom trabalho, time."
                                                             if tma else "CSAT dentro do esperado, bom trabalho, time.")
        if not tma:
            elogio = re.sub(r"\s*(e|,)?\s*TMA[^.,;]*", "", elogio).strip() or "CSAT dentro do esperado, bom trabalho, time."
        lines.append(f"✅ *Destaque:* {elogio}")
    return title, "\n".join(lines).strip()


_LAST_BLOCK_DATA: dict[str, dict[str, Any]] = {}


def _prepare_blocks(blocks: list[tuple[str, str]], day: date) -> list[tuple[str, str, str]]:
    """[(fila, título, corpo)] no layout fixo. Filas sem CSAT no D-1 ficam fora.
    Ordem: fora do alvo (maior gap de CSAT, depois maior excesso de TMA), no alvo, sem meta."""
    prepared = []
    for fila, text in blocks:
        data = _block_data_from_json(text) or _block_data_from_text(text)
        status = _block_status(data)
        if status == "sem_volume":
            continue
        rendered = _render_block(fila, data, day)
        if rendered is None:
            continue
        title, body = rendered
        _LAST_BLOCK_DATA[fila] = data
        tma = _tma_data(fila, data)
        tstatus = _tma_status(tma)
        csat_gap = (data["meta"] - data["d1"]) if status == "abaixo" else 0.0
        tma_excess = (tma["d1"] / tma["plan"] - 1) if tstatus == "acima" else 0.0
        off = status == "abaixo" or tstatus == "acima"
        order = 0 if off else (1 if (data.get("meta") is not None or tstatus == "ok") else 2)
        prepared.append((order, -csat_gap, -tma_excess, fila, title, body))
    prepared.sort(key=lambda item: item[:3])
    return [(fila, title, body) for *_, fila, title, body in prepared]


def _block_csat(fila: str, text: str) -> dict[str, float] | None:
    """D-1 e mês do bloco: "*93,5%* em D-1 … mês 94,6%" (aceita o formato antigo "Mês X% · D-1 Y%")."""
    d1 = _BLOCK_D1_RE.search(text or "") or _X_D1_AFTER_RE.search(text or "")
    if d1:
        mes = _BLOCK_MES_RE.search(text or "")
        mes_val = _to_float(mes.group(1)) if mes else None
        d1_val = _to_float(d1.group(1))
        return {"d1": d1_val, "mes": mes_val if mes_val is not None else d1_val}
    parsed = _parse_report_csat(text)
    return next(iter(parsed.values()), None) if parsed else None


def _report_blocks(value: str) -> list[dict[str, Any]]:
    """Estado do report: lista de blocos [{fila, ts, cobranca}] (aceita o formato antigo, só o ts)."""
    if not value:
        return []
    if value.startswith("["):
        try:
            data = json.loads(value)
            return [b for b in data if isinstance(b, dict) and b.get("ts")]
        except json.JSONDecodeError:
            return []
    return [{"fila": "", "ts": value, "cobranca": True}]


def _report_prompt(tenant: Tenant, compromissos_text: str) -> str:
    scope = _scope_instruction(tenant.eps)
    if tenant.eps:
        scope += (
            f" Este é o report da {tenant.eps}: use os nomes das células da {tenant.eps} como filas; "
            "células sem volume em D-1 aparecem como 'sem volume em D-1'. "
            "Não use a palavra iFood para se referir a totais — o total do report é o total da EPS."
        )
    parts = [scope]
    if compromissos_text:
        parts.append(compromissos_text)
    filas_tenant = _celulas_of(tenant.eps) if tenant.eps else list(TMA_PLAN)
    plan_lines = [f"- {f}: {_fmt_tempo(TMA_PLAN[f][0])}" for f in filas_tenant if f in TMA_PLAN]
    if plan_lines:
        parts.append("PLAN DE TMA (TMA acima do plan = fora do alvo):\n" + "\n".join(plan_lines))
    parts.append(DAILY_REPORT_PROMPT)
    return "\n\n".join(parts)


ESTEIRA_ENABLED = os.getenv("ESTEIRA_ENABLED", "0").strip() not in {"0", "false", "não", "nao"}
ESTEIRA_PILARES_EPS = {"operador", "tempo"}
_ESTEIRA_ORDER = ["Operador", "Tempo", "Política iFood", "Parceiro / logística", "Produto / sistema"]


def _esteira_prompt(tenant: Tenant, fila: str, gap: float, day: date) -> str:
    """Esteira do mês consolidado (dia 1 até D-1)."""
    return (
        f"{_scope_instruction(tenant.eps)}\n\n"
        "TAREFA INTERNA DO APP — NÃO é mensagem para o Slack. Chame o subagente *Lyra Qualidade* com exatamente:\n"
        f"EPS: {tenant.eps or 'iFood'}\nPedido: esteira de impacto\nFila: {fila}\n"
        f"Período: mês ({day.replace(day=1).strftime('%d/%m')} a {day.strftime('%d/%m')})\n"
        f"Gap: {_bfmt(gap, 2)} p.p.\n\n"
        "Devolva SOMENTE o JSON que ela responder (ou a linha ERRO: que vier), sem comentar."
    )


# Metas de CSAT por fila (as mesmas do prompt da Lyra Indicadores). Usadas na esteira sob demanda.
FILA_METAS: dict[str, float] = {
    "CX Suporte": 70, "Madrugada": 70, "CX Review - AeC": 70, "CX Super Cliente - AeC": 94,
    "CX Review - CSU": 70, "CX Super Cliente - CSU": 94, "PF Cancelamento": 80, "PF Super Clientes": 88,
    "PF Revisão Manual": 85, "CX Denúncia": 85,
}
ESTEIRA_MAX_POLLS = 72  # 12 min: a esteira não pode segurar a conversa
ESTEIRA_MIN_GAP_PP = 0.3  # abaixo disso o gap é pequeno demais para dividir em p.p.: mostra só %
_CAUSE_RE = re.compile(
    r"\b(por que|porque|pq|motivo|motivos|causa|causas|explica|o que (esta|ta) puxando|de onde vem)\b.*"
    r"\b(csat|caiu|cai|caindo|piorou|queda|gap|detrator|detratores|ruim|abaixo)\b"
    r"|\b(csat|gap|detratores)\b.*\b(por que|porque|pq|causa|causas|motivo)\b"
)


_KNOWLEDGE_RE = re.compile(
    r"\b(clube|assinatura|estorno|estornar|reembolso|reembolsar|cobranca (duplicada|nao reconhecida)|ype|recolhimento|"
    r"cnl|m111|cliente nao localizado|voucher|cupom|retratacao|bko|backoffice|oraculo|rosie|will|wills|nexus|clubia|"
    r"agentforce|matriz de decisao|platina|vip|gold|silver|bronze|off-target|super cliente|pos-entrega|pos entrega|"
    r"item faltante|item errado|eta|allocating|one hub|politica|processo|fluxo de atendimento|pode cancelar|pode estornar)\b"
)


_DATA_WORDS_RE = re.compile(r"\b(csat|tma|tmi|nps|meta|plan|report|volume|quadrante|dispersao|operadores|ranking|"
                            r"grafico|planilha|excel|response rate|reabertura|sla|mop|span)\b")


def _is_knowledge_question(question: str) -> bool:
    """Dúvida de processo/política/agente de IA: a resposta vem da Lyra Conhecimento, não de dado nem de chute."""
    text = _norm_text(question)
    if _DATA_WORDS_RE.search(text):  # "CSAT da Super Cliente", "TMA do VIP": é dado, não processo
        return False
    return bool(_KNOWLEDGE_RE.search(text)) and not _is_cause_request(question)


KNOWLEDGE_INSTRUCTION = (
    "INFORMAÇÃO INTERNA PARA A LYRA: esta é uma dúvida de processo, política ou agente de IA. Chame o subagente "
    "*Lyra Conhecimento* com `Pedido: <a pergunta, completa>` (uma chamada por pergunta, se houver mais de uma) e "
    "responda com o que ela devolver. NÃO use outra busca nem conhecimento geral (CDC, \"geralmente\", outras empresas). "
    "Se ela devolver NÃO SEI, diga que isso não está na sua base, que a referência é o One Hub (Matriz CX) e que vai "
    "confirmar com o time.\n\n"
)


_PLAN_RE = re.compile(r"\b(sugir\w*|sugere|sugestao|sugestoes|monta|monte|montar|cria|crie|criar|estrutur\w*|propo\w*|"
                      r"me ajuda|ajuda a|ideia|ideias|o que (eu )?(faco|fazer|podemos fazer))\b.*\b(acao|acoes|plano|planos)\b"
                      r"|\b(acao|acoes|plano)\b.*\b(sugir\w*|sugere|sugestao|ideia)")


def _is_plan_request(question: str) -> bool:
    return bool(_PLAN_RE.search(_norm_text(question)))


PLAN_INSTRUCTION = (
    "INFORMAÇÃO INTERNA PARA A LYRA: pedido de sugestão de plano de ação. Antes de responder, chame a *Lyra Conhecimento* "
    "com `Pedido: quais fluxos e políticas valem para <o motivo/problema da fila> e o que é decisão da operação x do iFood` "
    "e siga as regras de plano do seu prompt: nenhum passo de compensação/voucher/cupom/estorno/cancelamento fora da matriz "
    "de decisão; causa só a que o dado mostra (use a Leitura do report desta thread); número no recorte certo (fila x motivo); "
    "interações por ticket não é reabertura nem FCR; régua nova só como 'sugestão a validar'; não diga que tem/segue planilha "
    "ou arquivo; nada de sanção ou RH. No máximo 3 ações (o quê · quem · prazo · como medir), curtas.\n\n"
)

_FILE_WORDS_RE = re.compile(r"(planilha|arquivo|excel|\.xlsx|\.csv|\.pdf|\.pptx|anexo)", re.IGNORECASE)
_FILE_PROMISE_RE = re.compile(r"\b(tem uma|tenho uma|segue|seguem|anexei|anexo|te mandei|te enviei|enviei|mandei|"
                              r"deixei|preparei|montei|gerei|disponivel|baixar|se quiser)\b")


def _strip_file_promises(answer: str) -> str:
    """O modelo às vezes diz que 'tem uma planilha' que ninguém gerou: quem manda arquivo é o app. Tira essas linhas."""
    if not answer:
        return answer
    kept = [line for line in answer.split("\n")
            if not (_FILE_WORDS_RE.search(line) and _FILE_PROMISE_RE.search(_norm_text(line)))]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(kept)).strip() if len(kept) != len(answer.split("\n")) else answer


def _is_cause_request(question: str) -> bool:
    return bool(_CAUSE_RE.search(_norm_text(question)))


def _fila_in_question(question: str, eps: str | None) -> str | None:
    candidates = _celulas_of(eps) if eps else [f for cel in EPS_CELULAS.values() for f in cel]
    text = _norm_text(question)
    hits = [f for f in candidates if _norm_text(f) in text]
    if hits:
        return max(hits, key=len)
    return _match_fila(question, candidates) if candidates else None


def _esteira_on_demand_prompt(eps: str | None, fila: str, day: date, d1_only: bool) -> str:
    meta = FILA_METAS.get(fila)
    periodo = f"D-1 {day.strftime('%d/%m')}" if d1_only else f"mês ({day.replace(day=1).strftime('%d/%m')} a {day.strftime('%d/%m')})"
    return (
        f"{_scope_instruction(eps)}\n\n"
        "TAREFA INTERNA DO APP — NÃO é mensagem para o Slack. Chame o subagente *Lyra Qualidade* com exatamente:\n"
        f"EPS: {eps or 'iFood'}\nPedido: esteira de impacto\nFila: {fila}\nPeríodo: {periodo}\n"
        + (f"Meta: {_fmt_meta(meta)}%\nGap: calcule (meta − CSAT da fila no período; 0 se estiver na meta) e devolva no campo \"gap\"\n"
           if meta is not None else "Gap: 0\n")
        + "\nDevolva SOMENTE o JSON que ela responder (ou a linha ERRO: que vier), sem comentar."
    )


def _post_esteira_on_demand(event: dict[str, Any], eps: str | None, question: str) -> bool:
    """Pergunta de causa ('por que o CSAT caiu?'): posta a esteira de impacto da fila na thread. True = postou."""
    block = _report_block_for(str(event.get("channel") or ""), str(event.get("thread_ts") or ""))
    fila = block[1].get("fila") if block else _fila_in_question(question, eps)
    if not fila:
        return False
    text_q = _norm_text(question)
    d1_only = bool(re.search(r"\b(ontem|d-1|d1|dia anterior)\b", text_q)) and "mes" not in text_q
    day = _today_brazil().date() - timedelta(days=1)
    try:
        answer = ask_toqan(_esteira_on_demand_prompt(eps, fila, day, d1_only), max_attempts=ESTEIRA_MAX_POLLS)
        if _is_error_answer(answer):
            return False
        data = _block_data_from_json(answer) or {}
        gap = _num(data.get("gap")) or 0.0
        text = _render_esteira(data, gap, periodo="do D-1" if d1_only else "do mês", fila=fila)
        if not text:
            return False
        _post_sections(str(event.get("channel") or ""), text, thread_ts=_reply_ts(event))
        return True
    except Exception:
        logger.exception("On-demand esteira failed (%s)", fila)
        return False


def _variation_periods(question: str) -> tuple[tuple[date, date], tuple[date, date], str, str]:
    """(período atual, referência, rótulo atual, rótulo referência). Padrão: mês até D-1 x mês anterior até o mesmo dia."""
    d1 = _today_brazil().date() - timedelta(days=1)
    text = _norm_text(question)
    fmt = lambda a, b: f"{a.strftime('%d/%m')}–{b.strftime('%d/%m')}" if a != b else a.strftime("%d/%m")
    if re.search(r"\b(ontem|d-1|d1|dia anterior)\b", text):
        ref = (d1.replace(day=1), d1 - timedelta(days=1)) if d1.day > 1 else (d1 - timedelta(days=7), d1 - timedelta(days=1))
        return (d1, d1), ref, f"D-1 ({fmt(d1, d1)})", f"mês até D-2 ({fmt(*ref)})"
    if re.search(r"\bsemana\b", text):
        cur = (d1 - timedelta(days=6), d1)
        ref = (d1 - timedelta(days=13), d1 - timedelta(days=7))
        return cur, ref, f"últimos 7 dias ({fmt(*cur)})", f"7 dias anteriores ({fmt(*ref)})"
    cur = (d1.replace(day=1), d1)
    prev_end = d1.replace(day=1) - timedelta(days=1)
    ref_end = prev_end.replace(day=min(d1.day, prev_end.day))
    ref = (prev_end.replace(day=1), ref_end)
    return cur, ref, f"este mês ({fmt(*cur)})", f"mês anterior ({fmt(*ref)})"


def _variation_prompt(eps: str | None, fila: str, cur: tuple[date, date], ref: tuple[date, date]) -> str:
    iso = lambda d: d.isoformat()
    return (
        f"{_scope_instruction(eps)}\n\n"
        "TAREFA INTERNA DO APP — NÃO é mensagem para o Slack. Chame a *Lyra Indicadores* com:\n"
        f"EPS: {eps or 'iFood'}\nPedido: variação do CSAT\nFila: {fila}\n"
        f"Período atual: {iso(cur[0])} a {iso(cur[1])}\nReferência: {iso(ref[0])} a {iso(ref[1])}\n\n"
        "Use EXATAMENTE esses dois períodos e as fórmulas da seção 7.2 dela (efeito mix e efeito performance por motivo "
        "de contato; response rate). Devolva APENAS JSON válido:\n"
        '{"csat_ref": 94.8, "csat_atual": 93.7, "n_ref": 4120, "n_atual": 3980, "mix_pp": -0.8, "performance_pp": -0.3, '
        '"top_mix": [{"motivo": "pedido atrasado", "share_ref": 18.0, "share_atual": 26.0, "csat_ref": 85.8, "efeito_pp": -0.7}], '
        '"top_performance": [{"motivo": "cobrança não reconhecida", "csat_ref": 92.0, "csat_atual": 88.0, "share_atual": 9.0, "efeito_pp": -0.36}], '
        '"rr_ref": 27.1, "rr_atual": 26.9, '
        '"leitura": "1 frase factual sobre o que os números mostram, sem afirmar nada que eles não mostrem"}\n'
        "Percentuais em %, p.p. com sinal. Até 3 motivos em cada lista. Se a consulta falhar: só ERRO: <motivo>."
    )


MIX_DOMINANT_SHARE = float(os.environ.get("MIX_DOMINANT_SHARE", "0.8"))  # mix responde por ≥80% da queda
MIX_PERF_TOLERANCE_PP = 0.2  # performance estável (queda de até 0,2 p.p.)


def _mix_dominant(data: dict[str, Any]) -> bool:
    """Queda explicada pelo efeito mix: mudou o tipo de contato, não o atendimento. Nesse caso não se cobra plano."""
    c_ref, c_cur = _num(data.get("csat_ref")), _num(data.get("csat_atual"))
    mix, perf = _num(data.get("mix_pp")), _num(data.get("performance_pp"))
    if None in (c_ref, c_cur, mix, perf):
        return False
    total = c_cur - c_ref
    if total >= 0 or mix >= 0:
        return False
    return abs(mix) >= MIX_DOMINANT_SHARE * abs(total) and perf >= -MIX_PERF_TOLERANCE_PP


def _render_variation(fila: str, data: dict[str, Any], label_cur: str, label_ref: str) -> str | None:
    c_ref, c_cur = _num(data.get("csat_ref")), _num(data.get("csat_atual"))
    if c_ref is None or c_cur is None:
        return None
    total = c_cur - c_ref
    mix, perf = _num(data.get("mix_pp")), _num(data.get("performance_pp"))
    sign = lambda v: ("+" if v > 0 else "−" if v < 0 else "") + _bfmt(abs(v), 2 if abs(v) < 0.1 else 1)
    arrow = "📉" if total < 0 else "📈" if total > 0 else "➖"
    lines = [f"{arrow} *Variação do CSAT ·* `{fila}` — {label_ref} → {label_cur}"]
    n_txt = ""
    if data.get("n_ref") is not None and data.get("n_atual") is not None:
        n_txt = f" · n {_fmt_int(_num(data['n_ref']) or 0)} → {_fmt_int(_num(data['n_atual']) or 0)}"
    lines.append(f"CSAT *{_bfmt(c_ref)}% → {_bfmt(c_cur)}%* ({sign(total)} p.p.){n_txt}")
    if mix is not None:
        lines.append("")
        lines.append(f"🔀 *Efeito mix: {sign(mix)} p.p.* — mudou o tipo de contato que chegou")
        for m in (data.get("top_mix") or [])[:3]:
            if not isinstance(m, dict):
                continue
            part = f"• {m.get('motivo')}: {_bfmt(_num(m.get('share_ref')) or 0)}% → {_bfmt(_num(m.get('share_atual')) or 0)}% das avaliações"
            if m.get("csat_ref") is not None:
                part += f" (CSAT {_bfmt(_num(m['csat_ref']) or 0)}%)"
            if m.get("efeito_pp") is not None:
                part += f" · {sign(_num(m['efeito_pp']) or 0)} p.p."
            lines.append(part)
    if perf is not None:
        lines.append("")
        lines.append(f"🎯 *Efeito performance: {sign(perf)} p.p.* — mudou o atendimento do mesmo motivo")
        for m in (data.get("top_performance") or [])[:3]:
            if not isinstance(m, dict):
                continue
            part = f"• {m.get('motivo')}: CSAT {_bfmt(_num(m.get('csat_ref')) or 0)}% → {_bfmt(_num(m.get('csat_atual')) or 0)}%"
            if m.get("efeito_pp") is not None:
                part += f" · {sign(_num(m['efeito_pp']) or 0)} p.p."
            lines.append(part)
    if mix is not None and perf is not None and abs((mix + perf) - total) > 0.3:
        lines.append(f"⚠️ Mix + performance ({sign(mix + perf)} p.p.) não fecham com a variação ({sign(total)} p.p.): vou revisar a conta.")
    rr_ref, rr_cur = _num(data.get("rr_ref")), _num(data.get("rr_atual"))
    if rr_ref is not None and rr_cur is not None:
        d = rr_cur - rr_ref
        note = "mudou a amostra: parte da variação pode ser de quem respondeu" if abs(d) >= 3 else "estável, não distorce a leitura"
        lines.append("")
        lines.append(f"📨 *Response rate:* {_bfmt(rr_ref)}% → {_bfmt(rr_cur)}% ({sign(d)} p.p.) — {note}")
    leitura = str(data.get("leitura") or "").strip()
    if leitura:
        lines.append("")
        lines.append(f"💡 {leitura}")
    if _mix_dominant(data):
        lines.append(f"👉 *Leitura:* a queda vem do efeito mix ({sign(mix)} p.p.): chegou mais contato de motivo com CSAT "
                     "naturalmente menor, e o atendimento de cada motivo se manteve. Não é caso de plano de ação da EPS; "
                     "é contexto para acompanhar o volume desses motivos.")
    elif perf is not None:
        quem = (f"o efeito performance ({sign(perf)} p.p.) e os pilares Operador e Tempo da esteira"
                if perf < 0 else "os pilares Operador e Tempo da esteira (a performance não piorou)")
        lines.append(f"👉 *Da EPS cobrar:* {quem}. Mix e amostra são contexto, não culpa da operação.")
    return "\n".join(lines)


def _render_esteira(data: dict[str, Any], gap: float, periodo: str = "do mês", fila: str = "") -> str | None:
    """Esteira de impacto (cascata): gap da meta dividido entre os pilares pela proporção de detratores."""
    pilares = [p for p in (data.get("pilares") or []) if isinstance(p, dict) and (_num(p.get("detratores")) or 0) > 0
               and _norm_text(p.get("pilar")) != _norm_text("Sem causa clara")]
    total = sum(_num(p.get("detratores")) or 0 for p in pilares)
    if not pilares or total <= 0:
        return None
    rank = {_norm_text(n): i for i, n in enumerate(_ESTEIRA_ORDER)}
    pilares.sort(key=lambda p: (-(_num(p.get("detratores")) or 0), rank.get(_norm_text(p.get("pilar")), 99)))
    width = max(len(str(p.get("pilar"))) for p in pilares)
    has_gap = gap >= ESTEIRA_MIN_GAP_PP
    rows, eps_pp, eps_pct = [], 0.0, 0.0
    top_share = max((_num(p.get("detratores")) or 0) for p in pilares) / total
    for p in pilares:
        share = (_num(p.get("detratores")) or 0) / total
        impact = gap * share
        filled = max(1, round(share / top_share * 10))  # maior pilar = barra cheia
        bar = "█" * filled + "░" * (10 - filled)
        rows.append(f"{str(p.get('pilar')).ljust(width)}  {bar}  " + (f"{_bfmt(impact, 2)} p.p.  " if has_gap else "")
                    + f"{round(share * 100):d}%")
        if any(_norm_text(p.get("pilar")).startswith(k) for k in ESTEIRA_PILARES_EPS):
            eps_pp += impact
            eps_pct += share
    n_txt = f"{_fmt_int(total)} detratores classificados"
    alvo = f" · `{fila}`" if fila else ""
    if has_gap:
        head = f"de onde vem o gap de {_bfmt(gap, 1)} p.p."
    elif gap > 0:
        head = f"gap pequeno ({_bfmt(gap, 1)} p.p.); peso de cada causa nos detratores"
    else:
        head = "na meta; peso de cada causa nos detratores"
    lines = [f"🧱 *Esteira de impacto {periodo}*{alvo} — {head} ({n_txt})",
             "```" + "\n".join(rows) + "```"]
    top = pilares[0]
    if top.get("padrao"):
        lines.append(f"↳ *{top.get('pilar')}* é o maior pilar: {str(top['padrao']).strip().rstrip('.')}.")
    leitura = str(data.get("leitura") or "").strip()
    if leitura:
        lines.append(f"↳ {leitura}")
    tempos = data.get("tempos") or {}
    d_fr, p_fr = _num(tempos.get("primeira_resposta_detratores_s")), _num(tempos.get("primeira_resposta_promotores_s"))
    if d_fr is not None and p_fr:
        lines.append(f"⏱️ Primeira resposta: detratores {_fmt_tempo(d_fr)} x promotores {_fmt_tempo(p_fr)}"
                     + (f" ({_bfmt(d_fr / p_fr, 1)}x)" if p_fr else ""))
    if eps_pct > 0:
        parte = f"{_bfmt(eps_pp, 2)} p.p. do gap ({round(eps_pct * 100):d}%)" if has_gap else f"{round(eps_pct * 100):d}% dos detratores"
        lines.append(f"👉 *Da EPS (Operador + Tempo): {parte}.* "
                     "É isso que o plano precisa atacar. Política, parceiro e produto viram insumo para o iFood.")
    return "\n".join(lines)


def _post_esteiras(tenant: Tenant, posted: list[dict[str, Any]], day: date) -> None:
    """Depois do report: para cada fila com CSAT do D-1 abaixo da meta, posta na thread a esteira do mês consolidado
    (gap = meta − CSAT do mês; mês na meta = só o peso de cada causa)."""
    problems: list[str] = []
    for block in posted:
        data = _LAST_BLOCK_DATA.get(block.get("fila", ""), {})
        d1, meta, mes = _num(data.get("d1")), _num(data.get("meta")), _num(data.get("mes"))
        if not block.get("cobranca") or d1 is None or meta is None or d1 >= meta:
            continue
        gap = max(0.0, meta - mes) if mes is not None else meta - d1
        try:
            answer = ask_toqan(_esteira_prompt(tenant, block["fila"], gap, day))
            if _is_error_answer(answer):
                problems.append(f"{block['fila']}: {answer.strip()[:150]}")
                continue
            parsed = _block_data_from_json(answer)
            text = _render_esteira(parsed or {}, gap)
            if text:
                _post_sections(tenant.channel, text, thread_ts=block["ts"])
            else:
                problems.append(f"{block['fila']}: a Lyra Qualidade não devolveu os pilares ({answer.strip()[:120]!r})")
        except Exception as exc:
            logger.exception("[%s] Esteira de impacto failed (%s)", tenant.key, block.get("fila"))
            problems.append(f"{block.get('fila')}: {type(exc).__name__}")
    if problems:
        _alert(f"[{tenant.key}] Esteira de impacto não saiu para: " + " | ".join(problems[:4]))


def _run_daily_report(now: datetime | None = None, tenant: Tenant | None = None) -> str:
    """Gera e publica o report de um tenant; devolve o ts da mensagem principal."""
    now = now or _now_brazil()
    tenant = tenant or TENANTS[0]
    google_client = _safe_google_client()
    compromissos_text = _get_open_compromissos(google_client, now, tenant.key) if google_client else ""

    answer = ask_toqan(_report_prompt(tenant, compromissos_text))
    if _is_error_answer(answer):
        raise ReportGenerationError(answer.strip()[:300] or "resposta vazia")

    blocks = _prepare_blocks(_split_blocks(answer, use_headers=True), now.date() - timedelta(days=1))
    if blocks:
        posted: list[dict[str, Any]] = []
        csat: dict[str, dict[str, float]] = {}
        for fila, title, body in blocks:
            tagged, has_cobranca = _tag_block(fila, body)
            response = slack_app.client.chat_postMessage(channel=tenant.channel, text=title)
            ts = _resp_get(response, "ts") or (_resp_get(response, "message", {}) or {}).get("ts")
            if not ts:
                continue
            if tagged:
                try:  # detalhe e cobrança dentro da thread do título
                    _post_sections(tenant.channel, tagged, thread_ts=ts)
                except Exception:
                    logger.exception("Could not post block detail in thread (%s)", fila)
            posted.append({"fila": fila, "ts": str(ts), "cobranca": has_cobranca})
            values = _block_csat(fila, body)
            if values:
                csat[fila] = values
        if not posted:
            raise ReportGenerationError("Slack não devolveu o ts de nenhum bloco")
        STATE.set(f"report:{tenant.key}:{now.date().isoformat()}", json.dumps(posted, ensure_ascii=False))
        if ESTEIRA_ENABLED:
            threading.Thread(target=_post_esteiras, args=(tenant, list(posted), now.date() - timedelta(days=1)),
                             daemon=True).start()
        report_ts = posted[0]["ts"]
        if not csat:
            _alert(f"[{tenant.key}] Não consegui ler o CSAT D-1 dos blocos — pré/pós dos compromissos não será calculado hoje.")
    else:
        # formato antigo: mensagem principal + ===THREAD===
        main, thread = _split_report(answer)
        if not main:
            raise ReportGenerationError("mensagem principal vazia")
        csat = _parse_report_csat(answer)
        main = _tag_owners(main)
        response = slack_app.client.chat_postMessage(channel=tenant.channel, text=main)
        report_ts = _resp_get(response, "ts") or (_resp_get(response, "message", {}) or {}).get("ts")
        if not report_ts:
            raise ReportGenerationError("Slack não devolveu o ts do report")
        STATE.set(f"report:{tenant.key}:{now.date().isoformat()}", str(report_ts))
        if thread:
            try:
                slack_app.client.chat_postMessage(channel=tenant.channel, text=thread, thread_ts=report_ts)
            except Exception:
                logger.exception("Could not post report thread")

    if google_client is not None and csat:
        try:
            closed = _check_deadline_compromissos(google_client, csat, now, tenant.key)
            logger.info("[%s] Daily report: %d commitments closed with pós", tenant.key, closed)
        except Exception:
            logger.warning("Daily report: could not update commitment deadlines", exc_info=True)
    return str(report_ts)


# ---------------------------------------------------------------------------
# Calendário do MOP
# ---------------------------------------------------------------------------
def _is_business_day(day: date) -> bool:
    if day.weekday() >= 5:
        return False
    return not (_BR_HOLIDAYS is not None and day in _BR_HOLIDAYS)


def _next_business_day(day: date) -> date:
    while not _is_business_day(day):
        day += timedelta(days=1)
    return day


def _mop_fire_dates_for_month(year: int, month: int) -> list[date]:
    """1º dia útil e dia 15 (ou próximo dia útil). Fim de semana e feriado nacional pulados."""
    return [_next_business_day(date(year, month, 1)), _next_business_day(date(year, month, 15))]


def _is_mop_day(day: date) -> bool:
    return day in _mop_fire_dates_for_month(day.year, day.month)


# ---------------------------------------------------------------------------
# Saúde dos agentes (Lyra, Nexus, Vera, ClubIA...) e do token do Databricks
# ---------------------------------------------------------------------------
def _parse_any_date(value: str) -> date | None:
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value.strip(), fmt).date()
        except ValueError:
            continue
    return None


def _looks_like_token_error(text: str) -> bool:
    low = str(text or "").casefold()
    return any(hint in low for hint in TOKEN_ERROR_HINTS)


def _monitored_agents() -> list[tuple[str, str | None]]:
    """(nome, api_key). A Lyra usa a TOQAN_API_KEY (None)."""
    return [("Lyra", None)] + [(name, key) for name, key in MONITORED_AGENTS.items()]


def _ask_agent(api_key: str | None, prompt: str) -> str:
    conversation_id, request_id = create_conversation(prompt, api_key)
    return get_answer(conversation_id, request_id, api_key=api_key, max_attempts=AGENT_CHECK_MAX_POLLS)


def _check_agent(name: str, api_key: str | None) -> tuple[str, str]:
    """('ok' | 'token' | 'erro' | 'fora', detalhe). Nunca levanta exceção."""
    no_data = name.casefold() in NO_DATA_AGENTS
    try:
        answer = _ask_agent(api_key, PING_PROMPT if no_data else HEALTHCHECK_PROMPT).strip()
    except ToqanAPIError as exc:
        if "timed out" in str(exc):
            return "fora", "não respondeu em 5 minutos"
        return "fora", f"API do Toqan: {str(exc)[:120]}"
    except Exception as exc:
        return "fora", f"{type(exc).__name__}: {str(exc)[:120]}"
    first = next((line.strip() for line in answer.splitlines() if line.strip()), "")
    if first.upper().lstrip("*_ ").startswith("OK"):
        return "ok", first[:60]
    if not first:
        return "fora", "resposta vazia"
    if no_data:
        return "ok", "respondeu (sem dados)"
    if _looks_like_token_error(answer):
        return "token", first[:200]
    if first.upper().lstrip("*_ ").startswith("ERRO"):
        return "erro", first[:200]
    # respondeu, mas fora do padrão: está vivo; não alarma
    return "ok", "respondeu (fora do padrão)"


_STATUS_LABEL = {
    "ok": "✅ funcionando",
    "token": "🔴 token do Databricks expirado/sem permissão",
    "erro": "🟡 erro ao consultar dados",
    "fora": "🔴 sem resposta",
}


def _token_expiry() -> date | None:
    renewed = _parse_any_date(STATE.get("token_renewed") or DATABRICKS_TOKEN_RENEWED or "")
    return renewed + timedelta(days=DATABRICKS_TOKEN_VALIDITY_DAYS) if renewed else None


def _mark_token_renewed(day: date) -> date:
    STATE.set("token_renewed", day.isoformat())
    return day + timedelta(days=DATABRICKS_TOKEN_VALIDITY_DAYS)


def _run_agent_checks(now: datetime, remind: bool = False) -> dict[str, tuple[str, str]]:
    """Testa todos os agentes e grava o status. Avisa quando quebra, quando volta e,
    com remind=True (testes agendados), de novo enquanto continuar parado."""
    results: dict[str, tuple[str, str]] = {}
    broke: list[str] = []
    still: list[str] = []
    recovered: list[str] = []
    token_fixed = False
    for name, key in _monitored_agents():
        status, detail = _check_agent(name, key)
        results[name] = (status, detail)
        previous = STATE.get(f"agent_status:{name}")
        STATE.set(f"agent_status:{name}", status)
        if status != "ok" and previous != status:
            STATE.set(f"agent_down_since:{name}", now.strftime("%d/%m %Hh%M"))
            broke.append(f"• *{name}:* {_STATUS_LABEL[status]} — {detail}")
        elif status != "ok" and remind:
            since = STATE.get(f"agent_down_since:{name}") or "?"
            still.append(f"• *{name}:* {_STATUS_LABEL[status]} desde {since} — {detail}")
        elif status == "ok" and previous and previous != "ok":
            recovered.append(name)
            token_fixed = token_fixed or previous == "token"
    if broke:
        text = "Agente(s) com problema:\n" + "\n".join(broke)
        if any(results[n][0] == "token" for n, _ in _monitored_agents()):
            others = [a for a in SHARED_TOKEN_AGENTS if a not in results]
            if others:
                text += f"\n⚠️ O token é o mesmo de {', '.join(others)} — eles também devem estar sem acesso aos dados."
            text += "\nPra token: Databricks → Settings → Developer → Access tokens, gere um novo e atualize no Toqan. Depois me manda \"troquei o token\"."
        _alert(text)
    if still:
        _alert("Continua parado (ninguém resolveu ainda):\n" + "\n".join(still)
               + "\nSe for token: confira se trocou em TODOS os agentes no Toqan — é comum esquecer um.")
    if recovered:
        _alert(f"Voltou a funcionar ✅ {', '.join(recovered)}")
    if token_fixed:
        expiry = _mark_token_renewed(now.date())  # voltou depois de erro de token = token trocado
        logger.info("Token renewal inferred; next expiry %s", expiry)
    return results


def _agents_summary(results: dict[str, tuple[str, str]]) -> str:
    lines = ["*Status dos agentes:*"]
    for name, (status, detail) in results.items():
        note = f" — {detail}" if status != "ok" else (
            " (respondeu, mas não confirmou a consulta — confira o prompt dele)" if detail.startswith("respondeu (fora")
            else " (não usa token — só confiro se responde)" if detail == "respondeu (sem dados)" else "")
        lines.append(f"• *{name}:* {_STATUS_LABEL[status]}{note}")
    pending = [n for n, (st, _) in results.items() if st == "token"]
    if pending:
        lines.append(f"⚠️ Ainda sem acesso: *{', '.join(pending)}*. Faltou trocar o token nele(s) no Toqan?")
    if len(results) == 1:
        lines.append("_Só estou testando a Lyra. Pra eu testar cada agente/subagente, preencha MONITORED_AGENTS no Railway._")
    expiry = _token_expiry()
    if expiry:
        lines.append(f"Token do Databricks: vence em {_format_date(expiry)} (lembro você na véspera).")
    return "\n".join(lines)


def _token_reminder_tick(now: datetime) -> None:
    expiry = _token_expiry()
    if not expiry or now.time() < TOKEN_REMINDER_TIME:
        return
    today = now.date()
    key = f"token_reminder:{today.isoformat()}"
    if STATE.get(key):
        return
    days_left = (expiry - today).days
    if days_left == 1:
        _alert(f"Lembrete: o token do Databricks (Lyra, Nexus, Vera, ClubIA…) vence amanhã ({_format_date(expiry)}). "
               "Bom momento pra gerar o novo e atualizar no Toqan — depois me manda \"troquei o token\".")
    elif days_left == 0:
        _alert("O token do Databricks vence HOJE. Gere o novo e atualize no Toqan — depois me manda \"troquei o token\".")
    elif days_left < 0:
        _alert(f"O token do Databricks venceu em {_format_date(expiry)}. Se já trocou, me manda \"troquei o token\".")
    STATE.set(key, str(days_left))


_ACCESS_ERROR_RE = re.compile(
    r"(databricks|consulta|acesso|query|sql|dados).{0,80}(token|401|403|unauthori|expir|permiss|credencia|invalid access)"
    r"|(token|401|403|unauthori|expir|permiss|credencia|invalid access).{0,80}(databricks|consulta|acesso|query|sql|dados)",
    re.IGNORECASE | re.DOTALL,
)
PASSIVE_CHECK_COOLDOWN = timedelta(hours=1)


def _passive_access_check(answer: str) -> None:
    """Uma conversa voltou com erro de acesso aos dados: testa os agentes na hora (no máx. 1x por hora)."""
    if not answer or not _ACCESS_ERROR_RE.search(answer[:1500]):
        return
    now = _now_brazil()
    last = STATE.get("passive_check_at")
    if last:
        try:
            if now - datetime.fromisoformat(last) < PASSIVE_CHECK_COOLDOWN:
                return
        except ValueError:
            pass
    STATE.set("passive_check_at", now.isoformat())

    def run() -> None:
        try:
            results = _run_agent_checks(now)
            if all(st == "ok" for st, _ in results.values()):
                _alert("Uma conversa voltou com erro de acesso aos dados, mas no teste agora todos responderam. Fico de olho.")
        except Exception:
            logger.exception("Passive access check failed")

    threading.Thread(target=run, daemon=True).start()


def _agent_health_tick(now: datetime) -> None:
    _token_reminder_tick(now)
    today = now.date()
    due = [t for t in AGENT_CHECK_TIMES if now.time() >= t]
    if not due:
        return
    slot = max(due)
    key = f"agent_check:{today.isoformat()}:{slot.strftime('%H%M')}"
    if STATE.get(key):
        return
    STATE.set(key, "rodando")
    _run_agent_checks(now, remind=True)
    STATE.set(key, "ok")


_DATA_TEST_RE = re.compile(
    r"\b(testa|teste|testar|verifica|verificar|checa|checar|status|como estao|como esta)\b.*"
    r"\b(databricks|token|dados|base|agentes?|nexus|vera|clubia|lyra)\b"
)
_TOKEN_RENEWED_RE = re.compile(r"\b(troquei|renovei|atualizei|gerei|trocamos|renovamos)\b.*\btokens?\b")


def _is_data_test_request(question: str) -> bool:
    return bool(_DATA_TEST_RE.search(_norm_text(question)))


def _is_token_renewed_message(question: str) -> bool:
    return bool(_TOKEN_RENEWED_RE.search(_norm_text(question)))


_RUN_REPORT_RE = re.compile(r"\b(roda|rode|rodar|dispara|dispare|disparar|gera|gere|gerar|manda|mande|solta|solte)\b.*\breport\b")


def _report_targets(question: str) -> list[Tenant] | None:
    """Tenants pedidos em "roda o report [da aec|da csu]"; None se não for esse comando."""
    text = _norm_text(question)
    if not _RUN_REPORT_RE.search(text):
        return None
    named = [t for t in TENANTS if t.eps and re.search(rf"\b{_compact(t.eps)}\b", text.replace(" ", " "))]
    return named or list(TENANTS)


def _run_report_now(targets: list[Tenant]) -> list[str]:
    """Roda o report na hora (teste manual). Devolve uma linha de resultado por tenant."""
    lines = []
    now = _now_brazil()
    for tenant in targets:
        try:
            _run_daily_report(now, tenant)
            lines.append(f"• *{tenant.key}:* report postado no canal ✅")
        except Exception as exc:
            detail = f"{type(exc).__name__}: {str(exc)[:200]}"
            if _looks_like_token_error(str(exc)):
                detail += " (parece o token do Databricks)"
            lines.append(f"• *{tenant.key}:* não saiu — {detail}")
    return lines


# ---------------------------------------------------------------------------
# Agendador único
# ---------------------------------------------------------------------------
_report_next_attempt_at: dict[str, datetime] = {}
_scheduler_thread: threading.Thread | None = None
_scheduler_last_tick: datetime | None = None


def _skey(kind: str, tenant: Tenant, day: date) -> str:
    return f"{kind}:{tenant.key}:{day.isoformat()}"


def _posted_at(ts: str) -> datetime | None:
    try:
        return datetime.fromtimestamp(float(ts), tz=BRAZIL_TIMEZONE)
    except Exception:
        return None


def _run_thread_read(tenant: Tenant, report_value: str, phase: str, state_key: str, today: date) -> None:
    """Lê a thread de cada bloco (fila) do report."""
    errors = []
    for block in _report_blocks(report_value):
        try:
            result = _read_report_thread(
                tenant.channel, block["ts"], phase, today, tenant.key,
                fila=block.get("fila", ""), register_no_reply=bool(block.get("cobranca", True)),
            )
            logger.info("[%s] Thread read %s %s: %s", tenant.key, phase, block.get("fila"), result)
        except Exception as exc:
            logger.exception("[%s] Thread read %s failed", tenant.key, phase)
            errors.append(f"{block.get('fila') or 'report'}: {type(exc).__name__}: {str(exc)[:120]}")
    STATE.set(state_key, "erro" if errors else "ok")
    if errors:
        _alert(f"[{tenant.key}] Falha ao ler thread(s) do report ({phase}): " + " | ".join(errors[:3]))


def _tenant_tick(tenant: Tenant, now: datetime) -> None:
    today = now.date()

    # 1) Releitura da thread de ontem (respostas tardias) antes do report
    if now.time() >= PREVIOUS_THREAD_READ_TIME:
        yesterday = today - timedelta(days=1)
        y_ts = STATE.get(_skey("report", tenant, yesterday))
        if y_ts and STATE.get(_skey("read_first", tenant, yesterday)) and not STATE.get(_skey("read_next", tenant, yesterday)):
            _run_thread_read(tenant, y_ts, "next", _skey("read_next", tenant, yesterday), today)

    # 2) Report diário (com recuperação após restart e novas tentativas)
    if not STATE.get(_skey("report", tenant, today)):
        fail_key = _skey("report_fail", tenant, today)
        if DAILY_REPORT_TIME <= now.time() < REPORT_CATCHUP_UNTIL:
            failures = _to_int(STATE.get(fail_key))
            next_at = _report_next_attempt_at.get(tenant.key)
            if failures < REPORT_MAX_ATTEMPTS and (next_at is None or now >= next_at):
                try:
                    _run_daily_report(now, tenant)
                    _report_next_attempt_at.pop(tenant.key, None)
                except Exception as exc:
                    failures += 1
                    STATE.set(fail_key, str(failures))
                    _report_next_attempt_at[tenant.key] = now + timedelta(minutes=REPORT_RETRY_MINUTES)
                    logger.exception("[%s] Daily report attempt %d failed", tenant.key, failures)
                    detail = f"{type(exc).__name__}: {str(exc)[:250]}"
                    if _looks_like_token_error(str(exc)):
                        detail += " — parece o token do Databricks: gere um novo e atualize no Toqan"
                    if failures >= REPORT_MAX_ATTEMPTS:
                        _alert(f"[{tenant.key}] Report de {_format_date(today)} não saiu após {failures} tentativas. Último erro: {detail}")
                        try:
                            slack_app.client.chat_postMessage(channel=tenant.channel, text=REPORT_FAILED_CHANNEL_MESSAGE)
                        except Exception:
                            logger.exception("Could not post report failure notice")
                    else:
                        _alert(f"[{tenant.key}] Tentativa {failures} do report falhou ({detail}). Nova tentativa em {REPORT_RETRY_MINUTES} min.")
        elif now.time() >= REPORT_CATCHUP_UNTIL and not STATE.get(_skey("report_missed", tenant, today)) \
                and _to_int(STATE.get(fail_key)) < REPORT_MAX_ATTEMPTS:
            STATE.set(_skey("report_missed", tenant, today), "1")
            _alert(f"[{tenant.key}] O report de {_format_date(today)} não rodou (app fora do ar na janela 09h–13h?).")

    # 2b) Escalation: 5h depois do report, cobra quem foi marcado e não respondeu
    blocks = _report_blocks(STATE.get(_skey("report", tenant, today)))
    if blocks:
        esc_key = _skey("escalation", tenant, today)
        posted_at = _posted_at(blocks[0]["ts"]) or now
        if now >= posted_at + ESCALATION_DELAY and not STATE.get(esc_key):
            results = []
            for block in blocks:
                if not block.get("cobranca", True):
                    continue
                try:
                    results.append(_escalation_check(tenant, block["ts"], today))
                except Exception as exc:
                    results.append("erro")
                    _alert(f"[{tenant.key}] Falha na cobrança de 5h ({block.get('fila')}): {type(exc).__name__}")
            STATE.set(esc_key, ",".join(results) or "sem cobrança")

    # 3) Leituras da thread de hoje
    report_ts = STATE.get(_skey("report", tenant, today))
    if report_ts:
        first_block = _report_blocks(report_ts)
        posted = (_posted_at(first_block[0]["ts"]) if first_block else None) or now
        first_due = max(datetime.combine(today, THREAD_READ_FIRST_TIME, tzinfo=BRAZIL_TIMEZONE),
                        posted + THREAD_READ_FIRST_MIN_DELAY)
        first_key = _skey("read_first", tenant, today)
        if now >= first_due and not STATE.get(first_key):
            _run_thread_read(tenant, report_ts, "first", first_key, today)
        late_due = max(datetime.combine(today, THREAD_READ_LATE_TIME, tzinfo=BRAZIL_TIMEZONE),
                       posted + THREAD_READ_LATE_MIN_DELAY)
        late_key = _skey("read_late", tenant, today)
        if now >= late_due and STATE.get(first_key) and not STATE.get(late_key):
            _run_thread_read(tenant, report_ts, "late", late_key, today)



def _scheduler_tick(now: datetime | None = None) -> None:
    global _scheduler_last_tick
    now = now or _now_brazil()
    _scheduler_last_tick = now
    for tenant in TENANTS:
        try:
            _tenant_tick(tenant, now)
        except Exception:
            logger.exception("[%s] Scheduler tick failed", tenant.key)
    try:
        _mop_missing_check(now)
    except Exception:
        logger.exception("MOP missing check failed")
    try:
        _agent_health_tick(now)
    except Exception:
        logger.exception("Agent health check failed")


def _mop_missing_check(now: datetime) -> None:
    """Dia de MOP (1º dia útil / dia 15) e nada atualizado até 18h -> DM para o responsável. Nada no canal."""
    today = now.date()
    if not _is_mop_day(today) or now.time() < MOP_MISSING_ALERT_TIME:
        return
    key = f"mop_missing_check:{today.isoformat()}"
    if STATE.get(key):
        return
    last = STATE.get("mop_last_update")
    updated = False
    if last:
        try:
            updated = datetime.fromisoformat(last).date() >= today - timedelta(days=MOP_UPDATE_GRACE_DAYS)
        except ValueError:
            updated = False
    STATE.set(key, "ok" if updated else "alerta")
    if not updated:
        _alert(
            f"Hoje ({_format_date(today)}) é dia de atualização do MOP e até as 18h ele não foi atualizado "
            "pelo fluxo Reportar Oráculo."
        )


def _check_mop_access() -> None:
    """No boot: confirma que a service account abre a planilha do MOP."""
    if _safe_google_client() is None:
        return
    if _read_mop_values() is None:
        _alert(
            "Não consigo abrir a planilha do MOP. Compartilhe-a como Editor com "
            "lyra-mop@lyra-510815.iam.gserviceaccount.com e confira MOP_SPREADSHEET_ID / MOP_WORKSHEET_GID."
        )


def _scheduler_loop() -> None:
    STATE.get("boot")  # força a carga do estado
    try:
        _check_mop_access()
    except Exception:
        logger.exception("MOP access check failed")
    if not STATE.persistent:
        logger.warning("State store em memória: um restart pode repetir rotinas do dia.")
    logger.info("Scheduler started for %s", TENANTS)
    while True:
        try:
            _scheduler_tick()
        except Exception:
            logger.exception("Scheduler tick failed")
        time.sleep(SCHEDULER_TICK_SECONDS)


def _start_scheduler() -> None:
    global _scheduler_thread
    if _scheduler_thread is not None and _scheduler_thread.is_alive():
        return
    _scheduler_thread = threading.Thread(target=_scheduler_loop, name="lyra-scheduler", daemon=True)
    _scheduler_thread.start()


# ---------------------------------------------------------------------------
# Slack events
# ---------------------------------------------------------------------------
SLACK_SECTION_MAX = 2900


def _section_blocks(text: str) -> list[dict[str, Any]] | None:
    """Mensagem montada pelo app em seções (separadas por linha em branco) -> Block Kit com divisória entre elas.
    None = uma seção só (vai como texto normal)."""
    parts = [p.strip() for p in re.split(r"\n\s*\n", text or "") if p.strip()]
    if len(parts) < 2:
        return None
    blocks: list[dict[str, Any]] = []
    for i, part in enumerate(parts):
        if i:
            blocks.append({"type": "divider"})
        chunk = ""
        for line in part.split("\n"):
            if chunk and len(chunk) + len(line) + 1 > SLACK_SECTION_MAX:
                blocks.append({"type": "section", "text": {"type": "mrkdwn", "text": chunk}})
                chunk = ""
            chunk = f"{chunk}\n{line}" if chunk else line[:SLACK_SECTION_MAX]
        if chunk:
            blocks.append({"type": "section", "text": {"type": "mrkdwn", "text": chunk}})
    return blocks if len(blocks) <= 50 else None


def _post_sections(channel: str, text: str, thread_ts: str | None = None) -> Any:
    """Posta com divisórias entre as seções (D-1 | Mês | Leitura | Cobrança…). Se o Slack recusar os blocos, vai como texto."""
    kwargs: dict[str, Any] = {"channel": channel, "text": text}
    if thread_ts:
        kwargs["thread_ts"] = thread_ts
    blocks = _section_blocks(text)
    if blocks:
        try:
            return slack_app.client.chat_postMessage(**kwargs, blocks=blocks)
        except Exception:
            logger.warning("Blocks rejected; posting as plain text", exc_info=True)
    return slack_app.client.chat_postMessage(**kwargs)


def _say_sections(say: Callable[..., Any], answer: str, event: dict[str, Any], in_thread: bool = True) -> None:
    kwargs: dict[str, Any] = {"text": answer}
    ts = _reply_ts(event) if (in_thread or event.get("channel_type") == "im") else None
    if ts:
        kwargs["thread_ts"] = ts
    blocks = _section_blocks(answer)
    if blocks:
        try:
            say(**kwargs, blocks=blocks)
            return
        except Exception:
            logger.warning("Blocks rejected; replying as plain text", exc_info=True)
    say(**kwargs)


def _reply_ts(event: dict[str, Any]) -> str | None:
    """Onde a Lyra responde: sempre na thread da pergunta, em canal e na DM (duas perguntas seguidas não se misturam)."""
    return event.get("thread_ts") or event.get("ts")


def _say(say: Callable[..., Any], answer: str, event: dict[str, Any], in_thread: bool) -> None:
    ts = _reply_ts(event) if (in_thread or event.get("channel_type") == "im") else None
    if ts:
        say(text=answer, thread_ts=ts)
    else:
        say(text=answer)


ALERT_CHANNEL_IDS = {t for t in _parse_id_list(ALERT_CHANNEL) if t[:1] in {"C", "G"}}
MONITORING_HELP = ("Aqui eu cuido do monitoramento dos agentes. Pode pedir *status dos agentes* "
                   "ou avisar *troquei o token* que eu testo todos na hora.")


def _monitoring_channel_reply(question: str, say: Callable[..., Any], event: dict[str, Any], in_thread: bool) -> None:
    """Canal de alerta do time: só comandos de monitoramento (sem dados de EPS)."""
    if _is_token_renewed_message(question):
        expiry = _mark_token_renewed(_now_brazil().date())
        _say(say, f"Anotado ✅ Token novo vale até {_format_date(expiry)}; lembro um dia antes. "
                  "Vou testar os agentes agora…", event, True)
        _say(say, _agents_summary(_run_agent_checks(_now_brazil())), event, True)
    elif _is_data_test_request(question):
        _say(say, "Testando os agentes, leva uns minutos… ⏳", event, True)
        _say(say, _agents_summary(_run_agent_checks(_now_brazil())), event, True)
    else:
        _say(say, MONITORING_HELP, event, in_thread)


THREAD_CONTEXT_MAX_CHARS = 4000


def _thread_context(event: dict[str, Any]) -> str:
    """Conversa dentro de uma thread: manda junto o que já foi dito ali (ex.: o bloco do report da fila),
    para a Lyra não perguntar o que já está na tela. Marca a mensagem como respondida (sem crítica duplicada depois)."""
    thread_ts = str(event.get("thread_ts") or "")
    channel = str(event.get("channel") or "")
    if not thread_ts or thread_ts == str(event.get("ts") or "") or event.get("channel_type") == "im":
        return ""
    try:
        messages = _fetch_thread(channel, thread_ts)
    except Exception:
        logger.warning("Could not read thread context", exc_info=True)
        return ""
    parts = []
    for m in messages:
        if m.get("ts") == event.get("ts"):
            continue
        who = "Lyra" if (m.get("bot_id") or m.get("subtype") == "bot_message") else "Liderança"
        text = _strip_mentions(str(m.get("text") or "")).strip()
        if text:
            parts.append(f"[{who}] {text}")
    if not parts:
        return ""
    body = "\n".join(parts)
    if len(body) > THREAD_CONTEXT_MAX_CHARS:
        body = body[:1500] + "\n[…]\n" + body[-(THREAD_CONTEXT_MAX_CHARS - 1500):]
    block = _report_block_for(channel, thread_ts)
    head = "INFORMAÇÃO INTERNA PARA A LYRA — CONTEXTO DA THREAD (já está na tela de quem fala com você; use, não peça de novo)"
    if block:
        head += f". Esta thread é o report da fila {block[1].get('fila')}: avalie o que a pessoa disse com esses números"
        key = f"answered:{thread_ts}"
        answered = {t for t in STATE.get(key).split(",") if t}
        STATE.set(key, ",".join(sorted(answered | {str(event.get('ts'))})))
        # registra o compromisso na hora (a crítica já vai na resposta da conversa, não repete)
        tenant, blk, _ = block

        def register() -> None:
            try:
                time.sleep(LIVE_REPLY_DELAY_SECONDS)
                _read_report_thread(tenant.channel, thread_ts, "live", _today_brazil(), tenant.key,
                                    fila=blk.get("fila", ""), register_no_reply=False)
            except Exception:
                logger.exception("[%s] Commitment registration after mention failed", tenant.key)

        threading.Thread(target=register, daemon=True).start()
    return f"{head}:\n{body}\n\n"


DM_CONTEXT_MESSAGES = 12
DM_CONTEXT_HOURS = 3


def _dm_context(event: dict[str, Any]) -> str:
    """DM: manda junto as últimas mensagens da conversa (ou da thread da DM), para 'ontem' ou 'da AeC' completar o
    pedido anterior em vez de virar pergunta nova."""
    if event.get("channel_type") != "im":
        return ""
    channel = str(event.get("channel") or "")
    thread_ts = str(event.get("thread_ts") or "")
    try:
        if thread_ts and thread_ts != str(event.get("ts") or ""):
            messages = _fetch_thread(channel, thread_ts)
        else:
            oldest = str(time.time() - DM_CONTEXT_HOURS * 3600)
            resp = slack_app.client.conversations_history(channel=channel, limit=DM_CONTEXT_MESSAGES + 1, oldest=oldest)
            messages = list(reversed(resp.get("messages") or []))
    except Exception:
        logger.warning("Could not read DM context", exc_info=True)
        return ""
    parts = []
    for m in messages[-(DM_CONTEXT_MESSAGES + 1):]:
        if m.get("ts") == event.get("ts"):
            continue
        who = "Lyra" if (m.get("bot_id") or m.get("subtype") == "bot_message") else "Pessoa"
        text = _strip_mentions(str(m.get("text") or "")).strip()
        if text:
            parts.append(f"[{who}] {text}")
    if not parts:
        return ""
    body = "\n".join(parts)
    if len(body) > THREAD_CONTEXT_MAX_CHARS:
        body = body[-THREAD_CONTEXT_MAX_CHARS:]
    return ("INFORMAÇÃO INTERNA PARA A LYRA — CONVERSA ATÉ AQUI (a mensagem nova pode só completar o pedido anterior, "
            f"ex.: 'ontem' ou 'da AeC'; junte as duas e responda):\n{body}\n\n")


def _last_dm_human_text(event: dict[str, Any]) -> str:
    """Última mensagem da pessoa antes desta, na DM (para juntar um complemento curto ao pedido anterior)."""
    if event.get("channel_type") != "im":
        return ""
    try:
        resp = slack_app.client.conversations_history(channel=str(event.get("channel") or ""), limit=6,
                                                       oldest=str(time.time() - DM_CONTEXT_HOURS * 3600))
    except Exception:
        return ""
    for m in resp.get("messages") or []:  # mais recente primeiro
        if m.get("ts") == event.get("ts") or m.get("bot_id") or m.get("subtype") == "bot_message":
            continue
        return _strip_mentions(str(m.get("text") or "")).strip()
    return ""


def _date_context(question: str) -> str:
    """Data de hoje para a conversa (o agente não sabe que dia é) e correção de data no futuro digitada errado
    (ex.: "29/11" pedido em 30/10 = 29/10)."""
    today = _as_date(_today_brazil())
    yesterday = today - timedelta(days=1)
    weekday = ["segunda", "terça", "quarta", "quinta", "sexta", "sábado", "domingo"][today.weekday()]
    note = (f"INFORMAÇÃO INTERNA PARA A LYRA — hoje é {today.strftime('%d/%m/%Y')} ({weekday}); ontem (D-1) = "
            f"{yesterday.strftime('%d/%m/%Y')}.")
    fixes = []
    for m in re.finditer(r"\b(\d{1,2})/(\d{1,2})(?:/(\d{2,4}))?\b", question or ""):
        d, mo = int(m.group(1)), int(m.group(2))
        y = int(m.group(3)) if m.group(3) else today.year
        y = y + 2000 if y < 100 else y
        try:
            asked = date(y, mo, d)
        except ValueError:
            continue
        if asked < today:
            continue
        guess = None
        if (asked.year, asked.month) != (today.year, today.month):  # mês digitado errado: mesmo dia, mês atual/anterior
            for back in range(0, 2):
                mm, yy = today.month - back, today.year
                if mm < 1:
                    mm, yy = mm + 12, yy - 1
                try:
                    cand = date(yy, mm, d)
                except ValueError:
                    continue
                if cand < today:
                    guess = cand
                    break
        if guess:
            fixes.append(f"{m.group(0)} ainda não aconteceu; quase certamente é {guess.strftime('%d/%m/%Y')}: use essa "
                         f"data e avise em meia frase (\"considerei {guess.strftime('%d/%m')}\")")
        else:
            fixes.append(f"{m.group(0)} ainda não aconteceu: diga que ainda não tem dado desse dia e traga o D-1")
    if fixes:
        note += " " + "; ".join(fixes) + "."
    return note + "\n\n"


DEFAULTS_INSTRUCTION = (
    "INFORMAÇÃO INTERNA PARA A LYRA — RESPONDA, NÃO DEVOLVA PERGUNTA: se faltar o indicador, use CSAT; se faltar o "
    "período, use o D-1 (ontem) com o mês como contexto; se faltar a fila, traga todas as filas do escopo. Diga em meia "
    "frase o recorte que usou. Só pergunte quando for impossível responder sem a informação.\n\n")


def _process_event(event: dict[str, Any], say: Callable[..., Any], in_thread: bool, client: Any = None) -> None:
    """Processa upload de MOP ou repassa a mensagem para o Toqan."""
    event_key = _event_key(event)
    with _IN_FLIGHT_LOCK:
        if event_key in _IN_FLIGHT:
            return
        _IN_FLIGHT.add(event_key)

    question = _strip_mentions(str(event.get("text") or ""))
    try:
        allowed, eps = _event_scope(event)
        if not allowed and str(event.get("channel") or "") in ALERT_CHANNEL_IDS:
            _monitoring_channel_reply(question, say, event, in_thread)
            return
        if not allowed:
            is_im = event.get("channel_type") == "im"
            _say(say, DM_NOT_ALLOWED_MESSAGE if is_im else
                 f"{CHANNEL_NOT_CONFIGURED_MESSAGE} (canal `{event.get('channel')}`)", event, in_thread)
            logger.warning("Channel not configured: %s (scopes: %s)", event.get("channel"), sorted(CHANNEL_SCOPE))
            return
        mop_answer = _process_mop_upload(event, client, eps)
        if mop_answer is not None:
            _say(say, mop_answer, event, in_thread=True)
            return
        if not question:
            _say(say, "Oi! Sou a Lyra. Como posso ajudar?", event, in_thread)
            return

        if eps is None and _is_token_renewed_message(question):
            expiry = _mark_token_renewed(_now_brazil().date())
            reminder = expiry - timedelta(days=1)
            _say(say, f"Anotado ✅ Token novo vale até {_format_date(expiry)}; te lembro dia {_format_date(reminder)}. "
                      "Vou testar os agentes agora…", event, in_thread)
            _say(say, _agents_summary(_run_agent_checks(_now_brazil())), event, in_thread)
            return
        targets = _report_targets(question) if eps is None else None
        if targets:
            _say(say, f"Rodando o report agora ({', '.join(t.key for t in targets)})… pode levar alguns minutos.", event, in_thread)
            _say(say, "\n".join(_run_report_now(targets)), event, in_thread)
            return
        if eps is None and _is_data_test_request(question):
            _say(say, _agents_summary(_run_agent_checks(_now_brazil())), event, in_thread)
            return

        op_question = question
        if not _is_operators_excel_request(question) and len(question.split()) <= 4:
            prev = _last_dm_human_text(event)  # "ontem" respondendo a um pedido de operadores
            if prev and _is_operators_excel_request(prev):
                op_question = f"{prev} {question}"
        if _is_operators_excel_request(op_question):
            wants_file = bool(re.search(r"\b(excel|planilha|xlsx|arquivo)\b", _norm_text(op_question)))
            if eps is None and not wants_file:  # visão completa (e-mail iFood / canal interno): direto no chat
                _say(say, "Vou puxar o resultado por operador, já te mando. ⏳", event, in_thread=True)
                problem = _send_operators_inline(event, client, eps, op_question, say, in_thread)
                if problem:
                    _say(say, problem, event, in_thread)
                return
            _say(say, "Resultado individual vai em planilha, para não expor operadores aqui. Já te mando. ⏳",
                 event, in_thread=True)
            problem = _send_operators_excel(event, client, eps, op_question)
            if problem:
                _say(say, problem, event, in_thread)
            return

        if _is_mop_export_request(question):
            problem = _send_mop_export(event, client, eps)
            if problem:
                _say(say, problem, event, in_thread)
            return

        if _is_chart_request(question):
            _say(say, "Vou montar o gráfico, já te mando aqui. ⏳", event, in_thread=True)
            problem = _send_chart(event, client, eps, question, send_progress=lambda m: _say(say, m, event, in_thread=True))
            if problem:
                _say(say, problem, event, in_thread)
            return

        if _is_data_excel_request(question):
            _say(say, "Vou montar a planilha, já te mando aqui. ⏳", event, in_thread=True)
            problem = _send_data_excel(event, client, eps, question)
            if problem:
                _say(say, problem, event, in_thread)
            return

        def send_progress(message: str) -> None:
            _say(say, message, event, in_thread=True)

        extra = ""
        if _is_cause_request(question):
            fila_q = _fila_in_question(question, eps)
            block_q = _report_block_for(str(event.get("channel") or ""), str(event.get("thread_ts") or ""))
            if fila_q or block_q:
                # esteira (Qualidade) em paralelo com a análise de variação (Indicadores): a pessoa não fica esperando
                _say(say, "Vou levantar as causas: primeiro a análise da variação (mix de motivos e response rate); "
                          "se a queda não for só mix, em seguida a esteira de impacto por pilar. Leva alguns minutos ⏳",
                          event, in_thread=True)
                fila_v = block_q[1].get("fila") if block_q else fila_q
                cur, ref, l_cur, l_ref = _variation_periods(question)
                var_answer = ask_toqan(_variation_prompt(eps, fila_v, cur, ref), progress_callback=send_progress)
                var_data = {} if _is_error_answer(var_answer) else (_block_data_from_json(var_answer) or {})
                rendered = _render_variation(fila_v, var_data, l_cur, l_ref) if var_data else None
                if rendered:
                    _say_sections(say, rendered, event)
                    if not _mix_dominant(var_data):
                        threading.Thread(target=_post_esteira_on_demand, args=(event, eps, question), daemon=True).start()
                    return
                threading.Thread(target=_post_esteira_on_demand, args=(event, eps, question), daemon=True).start()
                extra = ("INFORMAÇÃO INTERNA PARA A LYRA: a esteira de impacto (pilares Operador, Tempo, Política, Parceiro, "
                         "Produto) é postada separadamente pelo app nesta thread. NÃO chame a Lyra Qualidade e NÃO liste pilares: "
                         "chame só a Lyra Indicadores (Pedido: variação do CSAT) e responda a variação (efeito mix x performance "
                         "e response rate). Feche com 1 frase sobre o que é da EPS cobrar.\n\n")
        if not extra and _is_plan_request(question):
            extra = PLAN_INSTRUCTION
        elif not extra and _is_knowledge_question(question):
            extra = KNOWLEDGE_INSTRUCTION
        context = _thread_context(event) or _dm_context(event)
        prompt = (f"{_scope_instruction(eps)}\n\n{_date_context(question)}{context}{extra}"
                  f"{DEFAULTS_INSTRUCTION if not extra else ''}"
                  f"{_inject_mop_context(question, eps)}")
        answer = ask_toqan(prompt, progress_callback=send_progress)
        _passive_access_check(answer)
        answer = _strip_file_promises(answer)
        conv_blocks = _split_blocks(answer or "")
        if conv_blocks:
            day = _today_brazil() - timedelta(days=1)
            for fila, title, body in _prepare_blocks(conv_blocks, _as_date(day)):
                _say_sections(say, title + ("\n\n" + _tag_block(fila, body)[0] if body else ""), event)
            return
        if THREAD_DELIMITER in (answer or ""):
            # report pedido em conversa: principal + detalhe, sem o marcador aparecer
            main, detail = _split_report(answer)
            _say(say, _tag_owners(main) or EMPTY_ANSWER_MESSAGE, event, in_thread)
            if detail:
                _say(say, detail, event, in_thread=True)
            return
        _say(say, answer or EMPTY_ANSWER_MESSAGE, event, in_thread)
    except MOPChannelNotAllowedError:
        _say(say, MOP_CHANNEL_NOT_ALLOWED_MESSAGE, event, in_thread=True)
    except MOPGoogleCredentialsError as exc:
        logger.warning("MOP upload: Google credentials error: %s", exc)
        _say(say, MOP_CREDENTIALS_ERROR_MESSAGE, event, in_thread=True)
    except MOPDownloadError as exc:
        _say(say, f"Não consegui baixar o arquivo: {exc}" if str(exc) else MOP_DOWNLOAD_ERROR_MESSAGE, event, in_thread=True)
    except MOPValidationError as exc:
        _say(say, f"Não atualizei o MOP: {exc}. A planilha atual foi mantida.", event, in_thread=True)
    except MOPReadError as exc:
        logger.warning("MOP upload: Excel read error: %s", exc)
        _say(say, MOP_READ_ERROR_MESSAGE, event, in_thread=True)
    except MOPSheetsError as exc:
        logger.warning("MOP upload could not update Google Sheets: %s", exc)
        _say(say, MOP_SHEETS_ERROR_MESSAGE, event, in_thread=True)
    except ToqanAPIError as exc:
        logger.exception("Lyra could not process Slack event")
        _say(say, TIMEOUT_ERROR_MESSAGE if "polling timed out" in str(exc) else ERROR_MESSAGE, event, in_thread)
    except Exception:
        logger.exception("Lyra could not process Slack event")
        _say(say, ERROR_MESSAGE, event, in_thread)
    finally:
        with _IN_FLIGHT_LOCK:
            _IN_FLIGHT.discard(event_key)


def _dispatch(event: dict[str, Any], say: Callable[..., Any], client: Any, ack: Callable[[], Any]) -> None:
    ack()
    channel = event.get("channel")
    timestamp = event.get("ts")
    if channel and timestamp:
        try:
            client.reactions_add(channel=channel, timestamp=timestamp, name=PROCESSING_REACTION)
        except Exception:
            logger.warning("Could not add processing reaction", exc_info=True)

    in_thread = event.get("channel_type") != "im"
    event_key = _event_key(event)

    def worker() -> None:
        try:
            _process_event(event, say, in_thread, client)
        finally:
            if channel and timestamp:
                try:
                    client.reactions_remove(channel=channel, timestamp=timestamp, name=PROCESSING_REACTION)
                except Exception:
                    logger.debug("Could not remove processing reaction", exc_info=True)
                try:
                    client.reactions_add(channel=channel, timestamp=timestamp, name=DONE_REACTION)
                except Exception:
                    logger.debug("Could not add done reaction", exc_info=True)

    try:
        threading.Thread(target=worker, name=f"lyra-{event_key}", daemon=True).start()
    except Exception:
        logger.exception("Could not start Lyra worker")


@slack_app.event("app_mention")
def handle_app_mention(event, say, client, ack):
    if OWNERS_FLOW_CHANNEL and event.get("channel") == OWNERS_FLOW_CHANNEL:
        _route_chamados_event(event)  # marcar a Lyra no fluxo de chamados = complemento do chamado
        ack()
        return
    if MOP_FLOW_CHANNEL and event.get("channel") == MOP_FLOW_CHANNEL:
        # No canal do fluxo a Lyra só processa o Excel do MOP; não conversa ali.
        ack()
        if _mop_attachment(event) is not None:
            threading.Thread(target=_handle_mop_flow_event, args=(event, say, client), daemon=True).start()
        return
    _dispatch(event, say, client, ack)


LIVE_REPLY_DELAY_SECONDS = int(os.getenv("LIVE_REPLY_DELAY_SECONDS", "60"))
_LIVE_PENDING: set[str] = set()
_LIVE_LOCK = threading.Lock()
_BOT_USER_ID: str | None = None


def _bot_user_id() -> str:
    global _BOT_USER_ID
    if _BOT_USER_ID is None:
        try:
            _BOT_USER_ID = str(_resp_get(slack_app.client.auth_test(), "user_id") or "")
        except Exception:
            _BOT_USER_ID = ""
    return _BOT_USER_ID


def _report_block_for(channel: str, thread_ts: str) -> tuple[Tenant, dict[str, Any], date] | None:
    """Bloco do report (de hoje ou de ontem) cuja thread recebeu a mensagem."""
    today = _today_brazil().date()
    for tenant in TENANTS:
        if tenant.channel != channel:
            continue
        for day in (today, today - timedelta(days=1)):
            for block in _report_blocks(STATE.get(_skey("report", tenant, day))):
                if block.get("ts") == thread_ts:
                    return tenant, block, day
    return None


def _maybe_live_thread_reply(event: dict[str, Any]) -> None:
    """Resposta na thread de um bloco do report: a Lyra lê, registra e comenta o plano (sem esperar 11h/18h)."""
    thread_ts = str(event.get("thread_ts") or "")
    if not thread_ts or thread_ts == str(event.get("ts") or "") or event.get("subtype") not in (None, "file_share"):
        return
    bot = _bot_user_id()
    if bot and f"<@{bot}>" in str(event.get("text") or ""):
        return  # marcou a Lyra: a conversa é tratada pela menção
    found = _report_block_for(str(event.get("channel") or ""), thread_ts)
    if not found:
        return
    tenant, block, day = found
    with _LIVE_LOCK:
        if thread_ts in _LIVE_PENDING:
            return  # já tem leitura agendada: junta as mensagens seguidas numa resposta só
        _LIVE_PENDING.add(thread_ts)

    def run() -> None:
        try:
            time.sleep(LIVE_REPLY_DELAY_SECONDS)
            with _LIVE_LOCK:
                _LIVE_PENDING.discard(thread_ts)
            _read_report_thread(tenant.channel, thread_ts, "live", _today_brazil(), tenant.key,
                                fila=block.get("fila", ""), register_no_reply=False)
        except Exception:
            logger.exception("[%s] Live thread read failed", tenant.key)

    threading.Thread(target=run, daemon=True).start()



# ---------------------------------------------------------------------------
# Status de hora em hora das EPSs ("Controle de SLA ..." nos canais de indicadores)
# ---------------------------------------------------------------------------
# A EPS posta na thread "Status atual na última hora - HH:MM - CSAT x% - SLA y%" + motivo, próximos passos e
# previsão. A Lyra lê, compara com a meta da fila e cobra só quando há o que cobrar: CSAT (ou SLA) fora da meta,
# justificativa copiada hora após hora, próximo passo genérico ou previsão sem prazo. Causa do lado do iFood
# (cupom, Rosie, sistema) fica só no registro da EPS: não vira cobrança de SLA.
STATUS_CHANNELS = _parse_channel_map(os.getenv("STATUS_CHANNELS", ""))  # "CSU:C0...,AeC:C0..."
_STATUS_CHANNEL_EPS = {channel: eps for eps, channel in STATUS_CHANNELS.items()}
STATUS_REPLY_DELAY_SECONDS = int(os.getenv("STATUS_REPLY_DELAY_SECONDS", "90"))
# Fase de testes: lê o canal real da EPS, mas responde no canal de testes da EPS (EPS_CHANNELS), sem marcar ninguém.
# STATUS_SHADOW=0 = responde na própria thread da EPS, marcando quem postou.
# Também aceita a lista de EPS que seguem em teste (ex.: STATUS_SHADOW=AeC: a AeC fica em teste e a CSU vale).
_SHADOW_RAW = os.getenv("STATUS_SHADOW", "1").strip()
_SHADOW_OFF = {"0", "false", "nao", "não", "off"}
_SHADOW_ON = {"1", "true", "sim", "on"}
STATUS_SHADOW = _SHADOW_RAW.lower() not in _SHADOW_OFF
_SHADOW_EPS = (set() if _SHADOW_RAW.lower() in _SHADOW_OFF | _SHADOW_ON
               else {p.strip().casefold() for p in _SHADOW_RAW.split(",") if p.strip()})
# Canal de testes de cada EPS para o modo teste. Vazio = usa EPS_CHANNELS (que na virada passa a ser o canal oficial).
STATUS_TEST_CHANNELS = _parse_channel_map(os.getenv("STATUS_TEST_CHANNELS", ""))


def _status_shadow(eps: str) -> bool:
    if _SHADOW_EPS:
        return eps.casefold() in _SHADOW_EPS
    return STATUS_SHADOW
STATUS_REPEAT_LIMIT = 3  # mesma justificativa em 3 status seguidos = cobrar atualização real
STATUS_RERAISE_EVERY = 4  # CSAT/SLA fora da meta e justificativa repetida só voltam a ser cobrados a cada 4 status


def _parse_sla_metas(raw: str) -> dict[str, float]:
    """'CX Super Cliente - CSU:90,CX Review - CSU:85' -> {fila: 90.0, ...}"""
    result: dict[str, float] = {}
    for part in raw.split(","):
        if ":" in part:
            fila, value = part.rsplit(":", 1)
            num = _num(value)
            if fila.strip() and num is not None:
                result[fila.strip()] = num
    return result


SLA_METAS = _parse_sla_metas(os.getenv("SLA_METAS", ""))
_STATUS_RE = re.compile(
    r"s?tatus atual na [uú]ltima hora\s*[-–:]?\s*(\d{1,2})(?:[:h](\d{2}))?\s*h?\s*[-–:]?\s*CSAT\s*:?\s*([\d.,]+)\s*%?"
    r"(?:\s*[-–|·/]?\s*SLA\s*:?\s*([\d.,]+)\s*%?)?", re.IGNORECASE)
_STATUS_FIELDS = {
    "motivo": re.compile(r"motivo da n[aã]o entrega\s*(?:\(\s*se\s*:?\s*houver\s*\))?\s*[:\-–]?\s*(.*)", re.IGNORECASE),
    "passos": re.compile(r"pr[oó]ximos passos(?: para resolu[cç][aã]o)?\s*[:\-–]?\s*(.*)", re.IGNORECASE),
    "previsao": re.compile(r"previs[aã]o de conclus[aã]o\s*[:\-–]?\s*(.*)", re.IGNORECASE),
}
_IFOOD_SIDE_RE = re.compile(r"\b(cupom|cupons|rosie|sistema|sistemico|sistemica|instabilidade|ifood|fora do ar|bug|"
                            r"ferramenta|oraculo|salesforce|lentidao|indisponibilidade)\b")
_NO_DATA_STATUS_RE = re.compile(r"\b(sem dados|nrt|dados nao carregad\w*|nao carregou|parcial .* repetida|"
                                r"painel fora|dashboard fora|relatorio fora|base nao atualiz\w*)\b")
STATUS_DROP_PP = 2.0  # CSAT caiu ≥ 2 p.p. de um status pro outro com o mesmo próximo passo
_GENERIC_STEP_RE = re.compile(r"\b(captacao\w*|captar|acompanhamento|monitoramento|em andamento|na meta|"
                              r"tentativas? de reversao|reversao|tratativa|alinhamento|reforco|seguimos|atencao)\b")
_DEADLINE_RE = re.compile(r"\d{1,2}\s*(h\b|:\d{2})|\d{1,2}/\d{1,2}|\bate (as|o fim|amanha)\b")
_STATUS_FILA_ALIASES = (("super cliente", "CX Super Cliente"), ("vip", "CX Super Cliente"), ("review", "CX Review"),
                        ("denuncia", "CX Denúncia"), ("madrugada", "Madrugada"), ("suporte", "CX Suporte"))
_STATUS_PENDING: set[str] = set()


def _status_fila(title: str, eps: str) -> str | None:
    """Fila do título da thread ("Controle de SLA CX VIP - 08/10/26" -> CX Super Cliente - CSU)."""
    celulas = _celulas_of(eps)
    text = _norm_text(title)
    exact = [f for f in celulas if _norm_text(f) in text]
    if exact:
        return max(exact, key=len)
    for alias, base in _STATUS_FILA_ALIASES:
        if re.search(rf"\b{alias}\b", text):
            hit = next((f for f in celulas if _norm_text(f).startswith(_norm_text(base))), None)
            if hit:
                return hit
    return None


# Formato "bloco" (ex.: Controle de SLA/CSAT Review da CSU): "20/10 – 9h:00 / CSAT: 87,05% / Promotoras: 14 /
# Detratoras: 6 / SLA Primeira Resposta: 100% / Volume da fila: 0 / Cenário atual: ... / Ações em andamento: * ..."
_BLOCK_HEAD_RE = re.compile(r"\b\d{1,2}/\d{1,2}\s*[–\-—]\s*(\d{1,2})\s*h\s*:?\s*(\d{2})?", re.IGNORECASE)
_BLOCK_CSAT_RE = re.compile(r"\bCSAT\s*:\s*([\d.,]+)\s*%", re.IGNORECASE)
_BLOCK_FIELDS = {
    "promotoras": re.compile(r"promotor\w*\s*:\s*(\d+)", re.IGNORECASE),
    "detratoras": re.compile(r"detrator\w*\s*:\s*(\d+)", re.IGNORECASE),
    "sla": re.compile(r"SLA\s*(?:de\s*)?(?:primeira|1[ªa°º]?)\s*resposta\s*:\s*([\d.,]+)\s*%", re.IGNORECASE),
    "volume": re.compile(r"volume da fila\s*:\s*(\d+)", re.IGNORECASE),
}
STATUS_CSAT_CHECK_PP = 1.5  # CSAT informado x promotoras ÷ (promotoras + detratoras)


def _looks_like_status(text: str) -> bool:
    raw = str(text or "")
    return bool(_STATUS_RE.search(raw) or (_BLOCK_CSAT_RE.search(raw) and (_BLOCK_HEAD_RE.search(raw)
                                                                          or re.search(r"cen[aá]rio atual", raw, re.I))))


def _parse_block_status(raw: str) -> dict[str, Any] | None:
    cm = _BLOCK_CSAT_RE.search(raw)
    if not cm or not _looks_like_status(raw):
        return None
    hm = _BLOCK_HEAD_RE.search(raw)
    status: dict[str, Any] = {"hora": f"{int(hm.group(1)):02d}:{hm.group(2) or '00'}" if hm else "--:--",
                              "csat": _num(cm.group(1)), "formato": "bloco", "previsao": None}
    for key, rx in _BLOCK_FIELDS.items():
        fm = rx.search(raw)
        status[key] = _num(fm.group(1)) if fm else None
    cen = re.search(r"cen[aá]rio atual\s*:?\s*\n*\s*([^\n]+)", raw, re.IGNORECASE)
    status["motivo"] = cen.group(1).strip() if cen else ""
    acoes = re.split(r"a[cç][oõ]es em andamento\s*:?", raw, maxsplit=1, flags=re.IGNORECASE)
    items = []
    if len(acoes) > 1:
        for line in acoes[1].split("\n"):
            line = re.sub(r"^[\s*•\-–]+", "", line).strip()
            if line and not re.match(r"\d+\s*anexos?$", line, re.IGNORECASE):
                items.append(line)
    status["acoes"] = items
    status["passos"] = " | ".join(items)
    return status


def _parse_status(text: str) -> dict[str, Any] | None:
    raw = re.sub(r"\s*\(editado\)\s*", " ", str(text or ""))
    m = _STATUS_RE.search(raw)
    if not m:
        return _parse_block_status(raw)
    status: dict[str, Any] = {
        "hora": f"{int(m.group(1)):02d}:{m.group(2) or '00'}",
        "csat": _num(m.group(3)), "sla": _num(m.group(4)) if m.group(4) else None,
    }
    for key, rx in _STATUS_FIELDS.items():
        fm = rx.search(raw)
        status[key] = re.sub(r"^[•\-\s]+|[\s.]+$", "", fm.group(1)).strip() if fm else ""
    return status


def _status_issues(statuses: list[dict[str, Any]], fila: str) -> tuple[list[tuple[str, str]], dict[str, Any]]:
    """Pontos a cobrar no status mais recente: [(tipo, texto)]."""
    last = statuses[-1]
    meta, sla_meta = FILA_METAS.get(fila), SLA_METAS.get(fila)
    ifood_side = bool(_IFOOD_SIDE_RE.search(_norm_text(last.get("motivo"))))

    def csat_off(s: dict[str, Any]) -> bool:
        return meta is not None and s.get("csat") is not None and s["csat"] < meta

    def sla_off(s: dict[str, Any]) -> bool:
        return sla_meta is not None and s.get("sla") is not None and s["sla"] < sla_meta

    issues: list[tuple[str, str]] = []
    prom, detr = last.get("promotoras"), last.get("detratoras")
    if last.get("csat") is not None and prom is not None and detr is not None and prom + detr >= 5:
        calc = prom / (prom + detr) * 100
        if abs(calc - last["csat"]) > STATUS_CSAT_CHECK_PP:
            issues.append(("dados", f"O CSAT informado ({_bfmt(last['csat'])}%) não bate com {int(prom)} promotoras e "
                                    f"{int(detr)} detratoras, que dão {_bfmt(calc)}%. Qual é o recorte de cada número?"))
    if not (csat_off(last) or sla_off(last)):
        return issues, last
    if _NO_DATA_STATUS_RE.search(_norm_text(last.get("motivo"))):
        return issues, last  # sem dado novo (NRT fora): número repetido não é cobrança da EPS
    if csat_off(last):
        streak = 0
        for s in reversed(statuses):
            if not csat_off(s):
                break
            streak += 1
        quando = f"há {streak} status seguidos" if streak > 1 else "neste status"
        issues.append(("csat", f"CSAT {_bfmt(last['csat'])}% contra a meta de {_fmt_meta(meta)}% da fila, abaixo {quando}. "
                               "O que está sendo feito pra recuperar e com qual efeito esperado?"))
    if sla_off(last) and not ifood_side:
        issues.append(("sla", f"SLA {_bfmt(last['sla'])}% contra a meta de {_fmt_meta(sla_meta)}%. Qual a causa e a ação?"))
    if len(statuses) > 1:
        prev = statuses[-2]
        if (csat_off(last) and prev.get("csat") is not None and last.get("csat") is not None
                and prev["csat"] - last["csat"] >= STATUS_DROP_PP
                and _norm_text(prev.get("passos")) == _norm_text(last.get("passos")) and last.get("passos")):
            issues.append(("piorou", f"O CSAT caiu de {_bfmt(prev['csat'])}% para {_bfmt(last['csat'])}% e o próximo passo "
                                     "continua o mesmo. O que muda no plano a partir de agora?"))
    key = lambda s: (_norm_text(s.get("motivo")), _norm_text(s.get("passos")))
    repeat = 0
    for s in reversed(statuses):
        if key(s) != key(last) or not key(last)[0]:
            break
        repeat += 1
    if repeat >= STATUS_REPEAT_LIMIT and (repeat - STATUS_REPEAT_LIMIT) % STATUS_RERAISE_EVERY == 0:
        issues.append(("repeat", f"A justificativa e os próximos passos são os mesmos há {repeat} status. "
                                 "O que mudou de uma hora pra outra: o que foi feito e qual o resultado?"))
    if last.get("acoes"):  # formato bloco: avalia ação por ação, cita só o título da vaga
        vagas = [a.split(":", 1)[0].strip() for a in last["acoes"]
                 if _GENERIC_STEP_RE.search(_norm_text(a)) and not re.search(r"\d", a)]
        if vagas:
            issues.append(("passos", "Ações sem dono nem prazo: " + ", ".join(f"\"{v}\"" for v in vagas[:3])
                                     + ". O que exatamente, com quem e até quando?"))
    else:
        passos = _norm_text(last.get("passos"))
        if passos and _GENERIC_STEP_RE.search(passos) and not re.search(r"\d", passos):
            issues.append(("passos", f"\"{last['passos']}\" não diz o que vai ser feito. Qual a ação, com quem e até que horas?"))
    previsao = _norm_text(last.get("previsao"))
    if last.get("previsao") is None and last.get("formato") == "bloco":
        pass  # o formato bloco não tem campo de previsão
    elif not previsao or not _DEADLINE_RE.search(previsao):
        txt = f"\"{last['previsao']}\" não é prazo" if last.get("previsao") else "Sem previsão de conclusão"
        issues.append(("prazo", f"{txt}: até que horas a fila volta pra meta?"))
    return issues, last


def _read_status_thread(channel: str, thread_ts: str, eps: str) -> str | None:
    """Lê a thread de status e, se houver o que cobrar, responde. Devolve o texto postado (ou None)."""
    messages = _fetch_thread(channel, thread_ts)
    if not messages:
        return None
    fila = _status_fila(str(messages[0].get("text") or ""), eps)
    if not fila:
        return None
    bot = _bot_user_id()
    statuses: list[dict[str, Any]] = []
    for m in messages[1:]:
        if m.get("bot_id") or (bot and m.get("user") == bot):
            continue
        for chunk in re.split(r"(?=s?tatus atual na [uú]ltima hora)", str(m.get("text") or ""), flags=re.IGNORECASE):
            parsed = _parse_status(chunk)
            if parsed:
                parsed["user"] = m.get("user")
                statuses.append(parsed)
    if not statuses:
        return None
    issues, last = _status_issues(statuses, fila)
    skey = f"status:{channel}:{thread_ts}"
    try:
        raised = json.loads(STATE.get(skey) or "{}")
    except ValueError:
        raised = {}
    idx = len(statuses)
    if raised.get("_last_idx") == idx:
        return None  # este status já foi lido
    field = {"passos": "passos", "prazo": "previsao"}
    new: list[tuple[str, str]] = []
    for kind, txt in issues:
        prev = raised.get(kind) if isinstance(raised.get(kind), dict) else None
        if kind in field:  # mesmo texto genérico de novo já é coberto pela cobrança de justificativa repetida
            ok = prev is None or prev.get("txt") != _norm_text(last.get(field[kind]))
        elif kind == "repeat":
            ok = True  # _status_issues só devolve no 3º, 7º, 11º status iguais
        else:
            ok = prev is None or idx - int(prev.get("idx", -99)) >= STATUS_RERAISE_EVERY
        if ok:
            new.append((kind, txt))
            raised[kind] = {"idx": idx, "txt": _norm_text(last.get(field.get(kind, ""), ""))}
    raised["_last_idx"] = idx
    STATE.set(skey, json.dumps(raised))
    if not new:
        return None
    head = f"👀 *{fila} · status das {last['hora']}* — CSAT {_bfmt(last['csat'])}%"
    if last.get("sla") is not None:
        head += f" · SLA {_bfmt(last['sla'])}%"
    return _post_status_reply(channel, thread_ts, eps, str(messages[0].get("text") or ""), head,
                              [txt for _, txt in new], last.get("user"))


def _post_status_reply(channel: str, thread_ts: str, eps: str, thread_title: str, head: str,
                       bullets: list[str], user: str | None) -> str:
    """Responde a um controle da EPS. Fase de testes (STATUS_SHADOW): vai para o canal de testes da EPS, numa
    thread por controle, sem marcar ninguém; senão, na própria thread da EPS marcando quem postou."""
    lines = [f"• {b}" for b in bullets]
    test_channel = (STATUS_TEST_CHANNELS.get(eps) or EPS_CHANNELS.get(eps)) if _status_shadow(eps) else None
    if test_channel:
        text = head + ("\n\n" + "\n".join(lines) if lines else "")
        parent_key = f"status_shadow:{channel}:{thread_ts}"
        parent_ts = STATE.get(parent_key)
        if not parent_ts:
            link = _permalink(channel, thread_ts)
            title = re.sub(r"[`*_]", "", thread_title).strip()[:120]
            intro = (f"🧪 *Teste · {eps}* — o que eu cobraria em _{title}_"
                     + (f" (<{link}|thread original>)" if link else "") + ". Nada foi postado no canal da EPS.")
            resp = slack_app.client.chat_postMessage(channel=test_channel, text=intro)
            parent_ts = str(_resp_get(resp, "ts") or "")
            if parent_ts:
                STATE.set(parent_key, parent_ts)
        _post_sections(test_channel, text, thread_ts=parent_ts or None)
        return text
    who = f"<@{user}> " if user else ""
    text = head + "\n\n" + "\n".join(lines + [f"{who}me responde aqui na thread."])
    _post_sections(channel, text, thread_ts=thread_ts)
    return text


# --- Controles em print (Volume, Produtividade): a Lyra lê a imagem e cobra; se não der pra ler, avisa -----------
PRINT_CONTROLS = (("hora", re.compile(r"acompanhamento de fila|hora ?/ ?hora")),
                  ("volume", re.compile(r"\bvolume\b")), ("produtividade", re.compile(r"\bprodutividade\b")))
# Acompanhamento de Fila da AeC (print de hora em hora, P = planejado, R = realizado). Lê só o CONSOLIDADO da hora.
HORA_TMA_RATIO = 1.10   # TMA realizado > 110% do planejado
HORA_SLA_PP = 2.0       # SLA 1ª resposta abaixo da meta (SLA_METAS) ou do P SLA em mais de 2 p.p.
HORA_MIN_CASES = 10     # TMA/SLA só com pelo menos 10 casos tratados/recebidos na hora
_HORA_SEGMENTOS = (("madrugada", "Madrugada"), ("review", "CX Review - AeC"), ("super", "CX Super Cliente - AeC"),
                   ("suporte", "CX Suporte"))
PRINT_MAX_POLLS = 36  # 6 min
# PRINT_READING=0 desliga a leitura dos prints (Volume/Produtividade). Falha técnica nunca vai para a EPS: só alerta.
PRINT_READING = os.getenv("PRINT_READING", "1").strip().lower() not in {"0", "false", "nao", "não", "off"}
VOL_BACKLOG_RATIO = float(os.getenv("VOL_BACKLOG_RATIO", "0.8"))  # produção < 80% da entrada = fila acumulando
HC_SHORT_RATIO = 0.9   # logados < 90% do projetado com volume no projetado ou acima
HC_EXCESS_RATIO = 1.2  # logados ≥ 120% do projetado (e +2 pessoas) com volume abaixo
PROD_TEAM_RATIO = float(os.getenv("PROD_TEAM_RATIO", "0.7"))  # equipe < 70% da média da célula na mesma hora
_PRINT_PENDING: set[str] = set()


def _print_control_kind(title: str) -> str | None:
    text = _norm_text(title)
    return next((kind for kind, rx in PRINT_CONTROLS if rx.search(text)), None)


def _image_files(message: dict[str, Any]) -> list[dict[str, Any]]:
    return [f for f in (message.get("files") or []) if isinstance(f, dict)
            and str(f.get("mimetype") or "").startswith("image/")]


def _print_prompt(kind: str) -> str:
    common = (
        "TAREFA INTERNA DO APP — NÃO é mensagem para o Slack. NÃO chame subagentes e NÃO consulte dados: só LEIA as "
        "imagens anexadas (prints de um controle da EPS) e transcreva os números. Se não conseguir ler TODOS os números "
        "com certeza (imagem pequena, cortada, borrada, ou você não consegue ver imagens), responda SOMENTE "
        "`ILEGIVEL: <motivo curto>`. Nunca chute número. Devolva APENAS JSON válido, números com ponto decimal.\n")
    if kind == "hora":
        return common + (
            "O print é o acompanhamento de fila de hora em hora da EPS: uma linha por Segmento (ex.: IFOOD_CX_REVIEW) ou "
            "por Intervalo, com as colunas P SLA, R SLA 1 Resp, R SLA Res, P Vol, R Vol, Δ, Tratados, P TMA, R TMA, Δ, "
            "P HC Prod, R HC Prod, Δ, P Prod/HC, R Prod/HC (P = planejado, R = realizado). Ignore as colunas Δ. "
            "Percentuais sem o símbolo %, TMA em segundos como está no print.\n"
            '{"tipo": "hora", "linhas": [{"segmento": "IFOOD_CX_REVIEW", "p_sla": 99.43, "r_sla_1resp": 100, '
            '"r_sla_res": null, "p_vol": 52, "r_vol": 5, "tratados": 5, "p_tma": 589.34, "r_tma": 350, "p_hc": 15.48, '
            '"r_hc": 10, "p_prod_hc": 3.35, "r_prod_hc": 0.51}]}\n'
            "Célula vazia: null. Uma linha por segmento ou intervalo que aparecer.")
    if kind == "volume":
        return common + (
            "Os prints trazem, por célula (ex.: CX Review, Super CX), por hora: volume projetado e real (tabela 'Volume por "
            "Hora'), produtividade projetada e real e HC projetado e real (tabela 'Produtividade e HCs por Hora'). Junte "
            "tudo por célula e hora. Só as horas com algum número.\n"
            '{"tipo": "volume", "data": "20/10/2026", "ate": "08:00", "celulas": [{"celula": "CX Review", "horas": '
            '[{"hora": "00:00", "vol_proj": 29, "vol_real": 13, "prod_proj": 29, "prod_real": 13, "hc_proj": 12, "hc_real": 14}], '
            '"total": {"vol_proj": 103, "vol_real": 77, "prod_proj": 103, "prod_real": 48}}]}\n'
            "Campo que não aparece no print: null. total = linha 'Total Geral' de cada tabela.")
    return common + (
        "O print é uma tabela de produtividade por hora: supervisor (linha agrupadora), atendentes (e-mail) e a contagem "
        "por hora, com Total Geral. NÃO devolva e-mails: só a contagem de cada atendente.\n"
        '{"tipo": "produtividade", "celula": "CX Review", "ate": "09:00", "supervisores": [{"supervisor": "NOME DO SUPERVISOR", '
        '"operadores": [{"horas": {"00:00": 2, "01:00": 1}, "total": 3}]}]}\n'
        "Hora sem número: não inclua. total = coluna Total Geral do atendente.")


def _validate_print(kind: str, data: dict[str, Any]) -> str | None:
    """Confere a leitura pelos totais do próprio print. Devolve o problema (ou None se a leitura fecha)."""
    if kind == "hora":
        linhas = data.get("linhas") or []
        if not linhas:
            return "não achei as linhas da tabela"
        if not any(_hora_fila(str(l.get("segmento") or "")) for l in linhas):
            return "não reconheci os segmentos"
        return None
    if kind == "volume":
        cels = data.get("celulas") or []
        if not cels:
            return "não achei as células"
        for cel in cels:
            horas, total = cel.get("horas") or [], cel.get("total") or {}
            for key in ("vol_real", "prod_real"):
                vals = [_num(h.get(key)) for h in horas if h.get(key) is not None]
                if total.get(key) is not None and vals and abs(sum(vals) - (_num(total[key]) or 0)) > 0.5:
                    return f"a soma das horas de {cel.get('celula')} não bate com o total ({key})"
        return None
    sups = data.get("supervisores") or []
    ops = [o for s in sups for o in (s.get("operadores") or [])]
    if not ops:
        return "não achei os atendentes"
    wrong = sum(1 for o in ops if o.get("total") is not None
                and abs(sum(_num(v) or 0 for v in (o.get("horas") or {}).values()) - (_num(o["total"]) or 0)) > 0.5)
    return f"{wrong} atendentes com soma diferente do total" if wrong > 1 else None


def _hh(hora: Any) -> str:
    m = re.match(r"\s*(\d{1,2})", str(hora or ""))
    return f"{int(m.group(1))}h" if m else str(hora)


def _judge_volume(data: dict[str, Any]) -> list[tuple[str, str]]:
    issues: list[tuple[str, str]] = []
    for cel in data.get("celulas") or []:
        name = str(cel.get("celula") or "célula")
        horas = [{k: (_num(v) if k != "hora" else v) for k, v in h.items()} for h in (cel.get("horas") or [])]
        vol = sum(h.get("vol_real") or 0 for h in horas)
        prod_vals = [h.get("prod_real") for h in horas if h.get("prod_real") is not None]
        if vol >= 20 and prod_vals and sum(prod_vals) < VOL_BACKLOG_RATIO * vol:
            issues.append((f"fila:{name}", f"*{name}:* produziram {_fmt_int(sum(prod_vals))} de {_fmt_int(vol)} casos que "
                                            f"entraram ({_bfmt(sum(prod_vals) / vol * 100, 0)}%). Está acumulando fila? Qual o plano?"))
        falta = [h for h in horas if h.get("hc_proj") and h.get("hc_real") is not None and h.get("vol_proj")
                 and h["hc_real"] < HC_SHORT_RATIO * h["hc_proj"] and (h.get("vol_real") or 0) >= h["vol_proj"]]
        if falta:
            ex = falta[0]
            issues.append((f"falta:{name}", f"*{name}:* menos gente que o projetado com volume no projetado ou acima "
                                             f"({', '.join(_hh(h['hora']) for h in falta[:4])}; ex.: {_hh(ex['hora'])}, "
                                             f"{_fmt_int(ex['hc_real'])} logados contra {_fmt_int(ex['hc_proj'])}). Risco de SLA: qual o ajuste?"))
        sobra = [h for h in horas if h.get("hc_proj") is not None and h.get("hc_real") is not None and h.get("vol_proj")
                 and h["hc_real"] >= HC_EXCESS_RATIO * h["hc_proj"] and h["hc_real"] - h["hc_proj"] >= 2
                 and (h.get("vol_real") or 0) < h["vol_proj"]]
        if sobra:
            ex = max(sobra, key=lambda h: h["hc_real"] - h["hc_proj"])
            issues.append((f"sobra:{name}", f"*{name}:* mais gente que o projetado com volume abaixo "
                                             f"({', '.join(_hh(h['hora']) for h in sobra[:4])}; ex.: {_hh(ex['hora'])}, "
                                             f"{_fmt_int(ex['hc_real'])} logados contra {_fmt_int(ex['hc_proj'])} projetados pra "
                                             f"{_fmt_int(ex.get('vol_real') or 0)} entradas). A escala acompanhou o volume?"))
        parado = [h for h in horas if (h.get("hc_real") or 0) >= 2 and (h.get("vol_real") or 0) >= 3 and h.get("prod_real") == 0]
        if parado:
            ex = parado[0]
            issues.append((f"parado:{name}", f"*{name}:* às {_hh(ex['hora'])} tinha {_fmt_int(ex['hc_real'])} pessoas logadas, "
                                              f"{_fmt_int(ex['vol_real'])} entradas e nenhuma produção. O que aconteceu?"))
    return issues


def _hora_fila(segmento: str) -> str | None:
    text = _norm_text(segmento).replace("_", " ")
    return next((fila for key, fila in _HORA_SEGMENTOS if key in text), None)


def _judge_hora(data: dict[str, Any]) -> list[tuple[str, str]]:
    """Acompanhamento de fila hora a hora (AeC): falta de HC, TMA acima do plan e SLA 1ª resposta abaixo da meta."""
    issues: list[tuple[str, str]] = []
    hora = str(data.get("ate") or "").strip()
    quando = f" às {_hh(hora)}" if hora else ""
    for linha in data.get("linhas") or []:
        fila = _hora_fila(str(linha.get("segmento") or ""))
        if not fila:
            continue
        v = {k: _num(linha.get(k)) for k in ("p_sla", "r_sla_1resp", "p_vol", "r_vol", "tratados", "p_tma", "r_tma",
                                              "p_hc", "r_hc")}
        if v["p_hc"] and v["r_hc"] is not None and v["p_vol"] and v["r_vol"] is not None \
                and v["r_hc"] < HC_SHORT_RATIO * v["p_hc"] and v["r_vol"] >= v["p_vol"]:
            issues.append((f"falta:{fila}", f"*{fila}:* HC abaixo do planejado{quando} ({_bfmt(v['r_hc'])} contra "
                                             f"{_bfmt(v['p_hc'])}) com volume no previsto ou acima ({_fmt_int(v['r_vol'])} "
                                             f"contra {_fmt_int(v['p_vol'])}). Qual o ajuste de escala?"))
        if v["p_tma"] and v["r_tma"] and (v["tratados"] or 0) >= HORA_MIN_CASES and v["r_tma"] > HORA_TMA_RATIO * v["p_tma"]:
            issues.append((f"tma:{fila}", f"*{fila}:* TMA de {_fmt_int(v['r_tma'])} s{quando}, "
                                           f"{_bfmt((v['r_tma'] / v['p_tma'] - 1) * 100, 0)}% acima do plan de "
                                           f"{_fmt_int(v['p_tma'])} s. O que está alongando o atendimento?"))
        meta = SLA_METAS.get(fila) or v["p_sla"]
        casos = max(v["r_vol"] or 0, v["tratados"] or 0)
        if meta and v["r_sla_1resp"] is not None and casos >= HORA_MIN_CASES and v["r_sla_1resp"] < meta - HORA_SLA_PP:
            issues.append((f"sla:{fila}", f"*{fila}:* SLA de 1ª resposta em {_bfmt(v['r_sla_1resp'])}%{quando}, abaixo "
                                           f"da meta de {_bfmt(meta)}%. O que está segurando a primeira resposta?"))
    return issues


def _judge_produtividade(data: dict[str, Any]) -> list[tuple[str, str]]:
    """Compara cada equipe com a média da célula NA MESMA HORA (madrugada x dia não se misturam)."""
    cells: list[tuple[str, str, float]] = []
    for sup in data.get("supervisores") or []:
        for op in sup.get("operadores") or []:
            for hora, val in (op.get("horas") or {}).items():
                v = _num(val)
                if v:
                    cells.append((str(sup.get("supervisor") or "?"), _hh(hora), v))
    by_hour: dict[str, list[float]] = {}
    for _, hora, v in cells:
        by_hour.setdefault(hora, []).append(v)
    hour_avg = {h: sum(vs) / len(vs) for h, vs in by_hour.items() if len(vs) >= 3}
    issues: list[tuple[str, str]] = []
    teams: dict[str, list[tuple[float, float]]] = {}
    for sup, hora, v in cells:
        if hora in hour_avg:
            teams.setdefault(sup, []).append((v, hour_avg[hora]))
    for sup, pairs in teams.items():
        if len(pairs) < 4:
            continue
        mine = sum(v for v, _ in pairs) / len(pairs)
        ref = sum(a for _, a in pairs) / len(pairs)
        if ref and mine < PROD_TEAM_RATIO * ref:
            issues.append((f"equipe:{sup}", f"Equipe de {sup.title()}: {_bfmt(mine)} por pessoa/hora contra {_bfmt(ref)} da "
                                            "célula nas mesmas horas. O que está segurando a produção da equipe?"))
    return issues


def _read_print_message(channel: str, thread_ts: str, msg_ts: str, eps: str) -> str | None:
    messages = _fetch_thread(channel, thread_ts)
    if not messages:
        return None
    title = str(messages[0].get("text") or "")
    kind = _print_control_kind(title)
    msg = next((m for m in messages if str(m.get("ts")) == msg_ts), None)
    if not kind or not msg:
        return None
    images = _image_files(msg)
    if not images:
        return None
    hora_txt = ""
    if kind == "hora":
        msg_text = _norm_text(msg.get("text"))
        if "consolidado" not in msg_text:  # as tabelas por célula repetem o consolidado: lê só ele
            return None
        mh = re.search(r"(\d{1,2})\s*[:h]\s*(\d{2})?", str(msg.get("text") or ""))
        hora_txt = f"{int(mh.group(1)):02d}:{mh.group(2) or '00'}" if mh else ""
        images = images[:1]
    user = msg.get("user")
    head_base = f"📷 *{re.sub(r'[`*_]', '', title).strip()[:60]} · {eps}*"
    problem = ""
    tech_problem = ""  # falha nossa (download, Toqan, JSON): só alerta, nunca cobra a EPS
    data: dict[str, Any] | None = None
    try:
        file_ids = []
        for f in images[:4]:
            content = _download_slack_bytes(f)
            if content:
                file_ids.append(toqan_upload_file(content, str(f.get("name") or "print.png"), str(f.get("mimetype") or "image/png")))
        if not file_ids:
            tech_problem = "não consegui baixar a imagem do Slack"
        else:
            answer = ask_toqan_with_files(_print_prompt(kind), file_ids, max_attempts=PRINT_MAX_POLLS)
            if answer.strip().upper().startswith("ILEGIVEL") or answer.strip().upper().startswith("ILEGÍVEL"):
                problem = answer.split(":", 1)[-1].strip()[:150] or "imagem ilegível"
            else:
                data = _block_data_from_json(answer)
                if not data:
                    tech_problem = "a resposta do Toqan não veio como tabela"
                else:
                    problem = _validate_print(kind, data) or ""
    except Exception as exc:
        logger.exception("[%s] Print read failed", eps)
        tech_problem = f"falha ao ler a imagem ({type(exc).__name__})"
    if tech_problem:
        _alert(f"Print de controle ({eps} · {re.sub(r'[`*_]', '', title).strip()[:60]}): {tech_problem}. Nada foi postado para a EPS.")
        return None
    if problem or not data:
        ikey = f"print_ilegivel:{channel}:{thread_ts}"
        if STATE.get(ikey):  # já pediu print legível nessa thread hoje: não repete a cada hora
            _alert(f"Print de controle ({eps}) ilegível de novo: {problem or 'sem dados'}.")
            return None
        STATE.set(ikey, "1")
        return _post_status_reply(channel, thread_ts, eps, title, head_base,
                                  [f"Não consegui ler o print ({problem or 'sem dados'}). Manda em tamanho maior, sem cortar "
                                   "a tabela, ou o Excel do controle."], user)
    if kind == "hora" and hora_txt and not data.get("ate"):
        data["ate"] = hora_txt
    issues = (_judge_hora(data) if kind == "hora" else _judge_volume(data) if kind == "volume"
              else _judge_produtividade(data))
    skey = f"print:{channel}:{thread_ts}"
    try:
        raised = json.loads(STATE.get(skey) or "{}")
    except ValueError:
        raised = {}
    idx = int(raised.get("_n", 0)) + 1
    new = [(k, txt) for k, txt in issues if idx - int(raised.get(k, -99)) >= STATUS_RERAISE_EVERY]
    raised["_n"] = idx
    for k, _ in new:
        raised[k] = idx
    STATE.set(skey, json.dumps(raised))
    if not new:
        return None
    ate = str(data.get("ate") or "").strip()
    return _post_status_reply(channel, thread_ts, eps, title, head_base + (f" — até {ate}" if ate else ""),
                              [txt for _, txt in new], user)


def _maybe_status_reply(event: dict[str, Any]) -> bool:
    channel = str(event.get("channel") or "")
    eps = _STATUS_CHANNEL_EPS.get(channel)
    thread_ts = str(event.get("thread_ts") or "")
    if not eps or not thread_ts or thread_ts == str(event.get("ts") or ""):
        return False
    if event.get("bot_id"):
        return False
    if _image_files(event):
        if not PRINT_READING:
            return False
        msg_ts = str(event.get("ts") or "")
        with _LIVE_LOCK:
            if msg_ts in _PRINT_PENDING:
                return True
            _PRINT_PENDING.add(msg_ts)

        def run_print() -> None:
            try:
                time.sleep(min(STATUS_REPLY_DELAY_SECONDS, 30))
                _read_print_message(channel, thread_ts, msg_ts, eps)
            except Exception:
                logger.exception("[%s] Print control read failed", eps)

        threading.Thread(target=run_print, daemon=True).start()
        return True
    if not _looks_like_status(str(event.get("text") or "")):
        return False
    with _LIVE_LOCK:
        if thread_ts in _STATUS_PENDING:
            return True
        _STATUS_PENDING.add(thread_ts)

    def run() -> None:
        try:
            time.sleep(STATUS_REPLY_DELAY_SECONDS)  # espera edição/complemento antes de ler
            with _LIVE_LOCK:
                _STATUS_PENDING.discard(thread_ts)
            _read_status_thread(channel, thread_ts, eps)
        except Exception:
            logger.exception("[%s] Status thread read failed", eps)

    threading.Thread(target=run, daemon=True).start()
    return True

# ---------------------------------------------------------------------------
# Chamados de moedas: Fluxo Owners CX -> #rewards-engine-support-cx
# ---------------------------------------------------------------------------
# A operação abre pelo "Fluxo Owners CX" (OWNERS_FLOW_CHANNEL). Quando o tópico é Moedas, a Lyra
# confere os dados mínimos, pede o que faltar na thread e abre o chamado no canal do Rewards
# (REWARDS_CHANNEL): pelo webhook do fluxo "Rewards Engine CX Support" (REWARDS_WEBHOOK_URL) ou,
# sem webhook, postando o chamado direto no canal. Depois devolve o retorno do Rewards na thread de origem.
OWNERS_FLOW_CHANNEL = os.getenv("OWNERS_FLOW_CHANNEL", "").strip()
REWARDS_CHANNEL = os.getenv("REWARDS_CHANNEL", "C08M93FKB63").strip()
REWARDS_WEBHOOK_URL = os.getenv("REWARDS_WEBHOOK_URL", "").strip()
REWARDS_PROBLEM_TYPES = [t.strip() for t in os.getenv(
    "REWARDS_PROBLEM_TYPES", "Missão concluída, mas moedas não creditadas").split(";") if t.strip()]
CHAMADOS_WORKSHEET_TITLE = "Chamados Moedas"
CHAMADOS_HEADERS = ["Data", "Ticket SF", "ID do Cliente", "ID do Pedido", "Missão", "Tipo de problema",
                    "Prioridade", "Segmento", "Título", "Solicitante", "Status", "Aberto em", "Finalizado em",
                    "Thread origem", "Thread Rewards"]
_FIELD_LINE_RE = re.compile(r"^\s*\*(?P<label>[^*]+?)\*\s*:?\s*(?P<value>.*)$")
_UUID_RE = re.compile(r"\b[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\b", re.IGNORECASE)
_EMOJI_CODE_RE = re.compile(r":[a-z0-9_+\-]+:")
_CHAMADOS_LOCK = threading.Lock()


def _parse_flow_fields(text: str) -> dict[str, str]:
    """Campos do fluxo: '*Rótulo:* valor' ou '*Rótulo*' com o valor nas linhas seguintes."""
    fields: dict[str, str] = {}
    current = None
    for raw in str(text or "").splitlines():
        line = raw.rstrip()
        m = _FIELD_LINE_RE.match(line)
        if m:
            current = _norm_text(m.group("label").split("(")[0].strip().rstrip(":"))
            fields[current] = m.group("value").strip()
            continue
        if current is not None and line.strip():
            fields[current] = (fields[current] + "\n" + line.strip()).strip()
    return {k: _EMOJI_CODE_RE.sub("", v).strip() for k, v in fields.items()}


def _flow_value(fields: dict[str, str], *labels: str) -> str:
    for label in labels:
        for key, value in fields.items():
            if key.startswith(_norm_text(label)):
                value = value.strip()
                return "" if value in {".", "-", "—"} else value
    return ""


def _is_moedas_flow(fields: dict[str, str], parent_fields: dict[str, str]) -> bool:
    topic = " ".join([_flow_value(fields, "tipo"), _flow_value(parent_fields, "topico"), _flow_value(fields, "topico")])
    return "moeda" in _norm_text(topic)


def _chamado_from_flow(fields: dict[str, str], parent_fields: dict[str, str], parent_text: str) -> dict[str, str]:
    ids = _MENTION_ID_RE.findall(parent_text or "")
    return {
        "ticket": re.sub(r"\D", "", _flow_value(fields, "ticket sf", "ticket", "numero do ticket")),
        "cliente": _flow_value(fields, "id do cliente"),
        "pedido": _flow_value(fields, "id do pedido"),
        "missao": _flow_value(fields, "id ou nome da missao", "missao"),
        "segmento": _flow_value(fields, "qual segmento", "segmento"),
        "prioridade": _flow_value(fields, "prioridade") or _flow_value(parent_fields, "prioridade"),
        "titulo": _flow_value(fields, "titulo") or _flow_value(parent_fields, "titulo"),
        "descricao": _MENTION_RE.sub("", _flow_value(fields, "descricao", "descreva o problema")).strip(),
        "solicitante": ids[0] if ids else "",
    }


def _triage_chamado(chamado: dict[str, str], extra_replies: list[str]) -> dict[str, Any]:
    """Classifica o tipo de problema, acha a missão e lista o que falta (via Toqan, JSON)."""
    prompt = (
        "TAREFA INTERNA DO APP — NÃO é mensagem para o Slack. Responda APENAS com JSON válido.\n\n"
        "Você faz a triagem de um chamado de MOEDAS (Rewards) aberto pela operação de atendimento iFood, "
        "antes de ele ir para o time de Rewards.\n"
        f"Chamado: {json.dumps(chamado, ensure_ascii=False)}\n"
        + (f"Complementos que a operação mandou na thread:\n" + "\n---\n".join(extra_replies) + "\n" if extra_replies else "")
        + f"Tipos de problema aceitos pelo Rewards: {json.dumps(REWARDS_PROBLEM_TYPES, ensure_ascii=False)}\n\n"
        'Formato: {"tipo_problema":"...","missao":"...","faltando":["..."],"descricao_rewards":"...","pergunta":"..."}\n'
        "Regras:\n"
        "- tipo_problema: o tipo da lista que melhor descreve; se nenhum servir, 'Outro: <resumo curto>'.\n"
        "- missao: ID ou nome da missão citado no chamado ou nos complementos; vazio se não houver.\n"
        "- faltando: dados mínimos que o Rewards exige e que NÃO estão no chamado nem nos complementos: "
        "número do ticket, ID do cliente (UUID), o que aconteceu, quando ocorreu, onde ocorreu (tela/fluxo do app), "
        "e a missão quando o problema é de missão. Não peça o que já está lá. Evidência (print) é opcional.\n"
        "- descricao_rewards: descrição limpa para o Rewards (o que, quando, como, onde), juntando chamado e complementos, "
        "sem saudação e sem dados pessoais além dos IDs.\n"
        "- pergunta: se faltar algo, UMA mensagem curta e educada pedindo exatamente o que falta; senão vazio."
    )
    parsed: Any = None
    try:
        parsed = json.loads(_strip_json_code_blocks(ask_toqan(prompt)))
    except Exception:
        logger.warning("Triage was not JSON", exc_info=True)
    if not isinstance(parsed, dict):
        parsed = {}
    faltando = [str(f) for f in (parsed.get("faltando") or []) if str(f).strip()]
    if not chamado.get("ticket"):
        faltando.insert(0, "número do ticket")
    if not _UUID_RE.search(chamado.get("cliente", "")):
        faltando.insert(0, "ID do cliente (UUID)")
    return {
        "tipo_problema": str(parsed.get("tipo_problema") or "").strip() or (REWARDS_PROBLEM_TYPES[0] if REWARDS_PROBLEM_TYPES else ""),
        "missao": str(parsed.get("missao") or chamado.get("missao") or "").strip(),
        "faltando": list(dict.fromkeys(faltando)),
        "descricao": str(parsed.get("descricao_rewards") or chamado.get("descricao") or "").strip(),
        "pergunta": str(parsed.get("pergunta") or "").strip(),
    }


def _download_slack_bytes(file_info: dict[str, Any]) -> bytes | None:
    url = file_info.get("url_private_download") or file_info.get("url_private")
    token = os.getenv("SLACK_BOT_TOKEN")
    if not url or not token:
        return None
    try:
        req = urllib.request.Request(str(url), headers={"Authorization": f"Bearer {token}"})
        with urllib.request.urlopen(req, timeout=60) as response:
            return response.read() or None
    except Exception:
        logger.warning("Could not download evidence file", exc_info=True)
        return None


def _post_rewards_chamado(chamado: dict[str, str], triage: dict[str, Any], origin_link: str) -> str | None:
    """Abre o chamado no Rewards. Devolve o ts da mensagem do chamado no canal (quando acha)."""
    payload = {
        "numero_ticket": chamado["ticket"], "id_cliente": chamado["cliente"], "id_pedido": chamado.get("pedido", ""),
        "missao": triage["missao"], "descricao": triage["descricao"], "tipo_problema": triage["tipo_problema"],
        "solicitante": chamado.get("solicitante", ""), "origem": origin_link,
    }
    if REWARDS_WEBHOOK_URL:
        req = urllib.request.Request(REWARDS_WEBHOOK_URL, data=json.dumps(payload).encode("utf-8"),
                                     headers={"Content-Type": "application/json"}, method="POST")
        with urllib.request.urlopen(req, timeout=30) as response:
            response.read()
        # o fluxo posta sozinho: procura a mensagem pelo número do ticket
        for _ in range(6):
            time.sleep(5)
            try:
                hist = slack_app.client.conversations_history(channel=REWARDS_CHANNEL, limit=20)
                for m in _resp_get(hist, "messages", []) or []:
                    if chamado["ticket"] and chamado["ticket"] in str(m.get("text") or ""):
                        return str(m.get("ts"))
            except Exception:
                logger.warning("Could not read Rewards channel history", exc_info=True)
                break
        return None
    solicitante = f"<@{chamado['solicitante']}>" if chamado.get("solicitante") else "operação"
    text = "\n".join(line for line in [
        "*Chamado de moedas* (aberto pela Lyra a partir do Fluxo Owners CX)",
        f"*Solicitante:* {solicitante}",
        f"*Número do Ticket:* {chamado['ticket']}",
        f"*ID do Cliente:* {chamado['cliente']}",
        f"*ID do Pedido:* {chamado['pedido']}" if chamado.get("pedido") else "",
        f"*Missão:* {triage['missao']}" if triage.get("missao") else "",
        f"*Tipo de problema:* {triage['tipo_problema']}",
        f"*Prioridade:* {chamado['prioridade']}" if chamado.get("prioridade") else "",
        f"*Descreva o problema:* {triage['descricao']}",
        f"<{origin_link}|Chamado de origem>" if origin_link else "",
    ] if line)
    resp = slack_app.client.chat_postMessage(channel=REWARDS_CHANNEL, text=text)
    return str(_resp_get(resp, "ts") or "") or None


def _permalink(channel: str, ts: str) -> str:
    try:
        return str(_resp_get(slack_app.client.chat_getPermalink(channel=channel, message_ts=ts), "permalink") or "")
    except Exception:
        return ""


def _chamados_ws() -> Any | None:
    google_client = _safe_google_client()
    if google_client is None:
        return None
    try:
        return _ensure_worksheet(_open_spreadsheet(google_client), CHAMADOS_WORKSHEET_TITLE, CHAMADOS_HEADERS)
    except Exception:
        logger.warning("Could not open Chamados worksheet", exc_info=True)
        return None


def _handle_owners_flow(channel: str, thread_ts: str) -> None:
    """Lê o chamado do Fluxo Owners CX (mensagem + thread) e, se for de moedas, faz triagem e abre no Rewards."""
    key = f"chamado:{thread_ts}"
    with _CHAMADOS_LOCK:
        state = STATE.get(key)
        if state in {"aberto", "processando", "fora"}:
            return
        STATE.set(key, "processando")
    try:
        messages = _fetch_thread(channel, thread_ts)
        parent = next((m for m in messages if m.get("ts") == thread_ts), {})
        parent_text = str(parent.get("text") or "")
        flow_msgs = [m for m in messages if m.get("bot_id") or m.get("subtype") == "bot_message"]
        details = next((m for m in flow_msgs if m.get("ts") != thread_ts and "descri" in _norm_text(m.get("text"))), None)
        if details is None:
            STATE.set(key, state or "")
            return  # o fluxo ainda não postou os detalhes na thread
        fields = _parse_flow_fields(str(details.get("text") or ""))
        parent_fields = _parse_flow_fields(parent_text)
        if not _is_moedas_flow(fields, parent_fields):
            STATE.set(key, "fora")
            return
        bot = _bot_user_id()
        replies = [str(m.get("text") or "") for m in messages
                   if m.get("ts") != thread_ts and not (m.get("bot_id") or m.get("subtype") == "bot_message")
                   and m.get("user") != bot]
        chamado = _chamado_from_flow(fields, parent_fields, parent_text)
        triage = _triage_chamado(chamado, replies)
        if triage["faltando"]:
            if state != "aguardando":
                quem = f"<@{chamado['solicitante']}> " if chamado.get("solicitante") else ""
                pergunta = triage["pergunta"] or ("Pra eu abrir o chamado no time de Rewards, me manda aqui: "
                                                 + ", ".join(triage["faltando"]) + ".")
                slack_app.client.chat_postMessage(channel=channel, thread_ts=thread_ts, text=f"{quem}{pergunta}")
            STATE.set(key, "aguardando")
            return
        origin = _permalink(channel, thread_ts)
        rewards_ts = _post_rewards_chamado(chamado, triage, origin)
        files = [f for m in messages for f in (m.get("files") or []) if isinstance(f, dict)]
        if rewards_ts and files:
            for f in files[:10]:
                content = _download_slack_bytes(f)
                if content:
                    try:
                        slack_app.client.files_upload_v2(channel=REWARDS_CHANNEL, thread_ts=rewards_ts, content=content,
                                                         filename=f.get("name") or "evidencia.png", title="Evidência")
                    except Exception:
                        logger.warning("Could not upload evidence", exc_info=True)
        link = _permalink(REWARDS_CHANNEL, rewards_ts) if rewards_ts else ""
        slack_app.client.chat_postMessage(
            channel=channel, thread_ts=thread_ts,
            text=("Chamado aberto no time de Rewards ✅" + (f" <{link}|ver chamado>" if link else f" (<#{REWARDS_CHANNEL}>)")
                  + f"\n*Tipo:* {triage['tipo_problema']}" + (f" · *Missão:* {triage['missao']}" if triage["missao"] else "")
                  + "\nTe aviso aqui quando o Rewards responder."))
        STATE.set(key, "aberto")
        if rewards_ts:
            STATE.set(f"rewards:{rewards_ts}", json.dumps({"channel": channel, "ts": thread_ts}))
        ws = _chamados_ws()
        if ws is not None:
            now = _now_brazil()
            ws.append_row([_format_date(now), chamado["ticket"], chamado["cliente"], chamado.get("pedido", ""),
                           triage["missao"], triage["tipo_problema"], chamado.get("prioridade", ""),
                           chamado.get("segmento", ""), chamado.get("titulo", ""), chamado.get("solicitante", ""),
                           "aberto", now.strftime("%d/%m/%Y %H:%M"), "", thread_ts, rewards_ts or ""],
                          value_input_option="RAW")
    except Exception as exc:
        logger.exception("Owners flow handling failed")
        STATE.set(key, "")
        _alert(f"Falha ao abrir chamado de moedas ({thread_ts}): {type(exc).__name__}: {str(exc)[:150]}")


def _handle_rewards_reply(event: dict[str, Any]) -> None:
    """Resposta no chamado do Rewards: devolve na thread de origem (Fluxo Owners CX)."""
    thread_ts = str(event.get("thread_ts") or "")
    origin_raw = STATE.get(f"rewards:{thread_ts}") if thread_ts else ""
    if not origin_raw or event.get("user") == _bot_user_id():
        return
    try:
        origin = json.loads(origin_raw)
    except json.JSONDecodeError:
        return
    text = _MENTION_RE.sub("", str(event.get("text") or "")).strip()
    if not text:
        return
    finished = "solicitacao foi finalizada" in _norm_text(text)
    is_flow = bool(event.get("bot_id") or event.get("subtype") == "bot_message")
    if is_flow and not finished:
        return  # avisos automáticos do fluxo (início de atendimento etc.)
    msg = "✅ O time de Rewards finalizou o chamado." if finished else f"💬 *Retorno do time de Rewards:*\n{text[:1500]}"
    slack_app.client.chat_postMessage(channel=origin["channel"], thread_ts=origin["ts"], text=msg)
    if finished:
        ws = _chamados_ws()
        if ws is not None:
            try:
                for n, row in _worksheet_rows_with_header(ws):
                    if _row_value(row, "Thread Rewards") == thread_ts:
                        ws.update_cell(n, CHAMADOS_HEADERS.index("Status") + 1, "finalizado")
                        ws.update_cell(n, CHAMADOS_HEADERS.index("Finalizado em") + 1, _now_brazil().strftime("%d/%m/%Y %H:%M"))
            except Exception:
                logger.warning("Could not update Chamados row", exc_info=True)


def _route_chamados_event(event: dict[str, Any]) -> bool:
    """Eventos dos canais de chamados. True = já tratado."""
    channel = str(event.get("channel") or "")
    if OWNERS_FLOW_CHANNEL and channel == OWNERS_FLOW_CHANNEL:
        root = str(event.get("thread_ts") or event.get("ts") or "")
        if event.get("user") != _bot_user_id() and root:
            # espera o fluxo postar os detalhes na thread e junta mensagens seguidas
            def run() -> None:
                time.sleep(LIVE_REPLY_DELAY_SECONDS if event.get("thread_ts") and not event.get("bot_id") else 5)
                _handle_owners_flow(channel, root)
            threading.Thread(target=run, daemon=True).start()
        return True
    if REWARDS_CHANNEL and channel == REWARDS_CHANNEL and event.get("thread_ts"):
        threading.Thread(target=_handle_rewards_reply, args=(event,), daemon=True).start()
        return True
    return False


@slack_app.event("message")
def handle_message(event, say, client, ack):
    if _route_chamados_event(event):
        ack()
        return
    # Canal do fluxo "Reportar Oráculo": Excel de MOP é lido sem precisar marcar a Lyra
    # (inclui arquivo postado pela própria mensagem do fluxo).
    if MOP_FLOW_CHANNEL and event.get("channel") == MOP_FLOW_CHANNEL:
        ack()
        if _mop_attachment(event) is not None and event.get("subtype") in (None, "file_share", "bot_message"):
            threading.Thread(target=_handle_mop_flow_event, args=(event, say, client), daemon=True).start()
        return
    # Em canal, a conversa precisa de @menção (tratada em app_mention).
    # Em DM, aceita texto e arquivo (subtype "file_share").
    if event.get("bot_id"):
        ack()
        return
    is_im = event.get("channel_type") == "im"
    if is_im and event.get("subtype") in (None, "file_share"):
        _dispatch(event, say, client, ack)
    else:
        ack()
        if not is_im:
            try:
                if _maybe_status_reply(event):
                    return
                _maybe_live_thread_reply(event)
            except Exception:
                logger.exception("Live thread check failed")


# ---------------------------------------------------------------------------
# HTTP
# ---------------------------------------------------------------------------
@flask_app.get("/")
@flask_app.get("/health")
def health():
    alive = _scheduler_thread is not None and _scheduler_thread.is_alive()
    return jsonify({
        "status": "ok",
        "service": "lyra-slack-bot",
        "scheduler_alive": alive,
        "last_tick": _scheduler_last_tick.isoformat() if _scheduler_last_tick else None,
    }), 200


def _debug_allowed() -> bool:
    return bool(DEBUG_TOKEN) and request.headers.get("X-Debug-Token") == DEBUG_TOKEN


@flask_app.get("/debug/status")
def debug_status():
    """Diagnóstico protegido (header X-Debug-Token). Sem DEBUG_TOKEN, responde 404."""
    if not _debug_allowed():
        return jsonify({"error": "not found"}), 404
    result: dict[str, Any] = {"state_persistent": STATE.persistent}
    try:
        google_client = _get_google_client()
        spreadsheet = _open_spreadsheet(google_client)
        result["sheets_ok"] = True
        result["worksheets"] = [ws.title for ws in spreadsheet.worksheets()]
    except Exception as exc:
        result["sheets_ok"] = False
        result["sheets_error"] = type(exc).__name__
    values = _read_mop_values()
    if values is not None:
        eps_col, email_col = (_find_col(values[0], "EPS"), _find_col(values[0], "Email")) if values else (None, None)
        rows = [r for r in values[1:] if any(str(c).strip() for c in r)]
        result["mop"] = {
            "pessoas": len(rows),
            "cabecalho": values[0] if values else [],
            "por_eps": {eps: len([r for r in rows if _eps_matches(_row_eps(r, eps_col, email_col), eps)]) for eps in EPS_CELULAS},
            "sem_eps_identificada": len([r for r in rows if not _row_eps(r, eps_col, email_col)]),
        }
    today = _now_brazil().date().isoformat()
    result["today"] = {k: v for k, v in STATE.snapshot().items() if today in k and not k.startswith("captured:")}
    return jsonify(result), 200


@flask_app.post("/slack/events")
def slack_events():
    return slack_handler.handle(request)


if __name__ == "__main__":
    # Rode com `python app.py` (Procfile) e UMA réplica: o agendador vive neste processo.
    _start_scheduler()
    flask_app.run(host="0.0.0.0", port=PORT)
